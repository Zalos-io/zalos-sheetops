from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.read_cells import read_cells


def test_read_cells_returns_freeform_range_with_formulas_and_cached_values(
    tmp_path: Path,
) -> None:
    path = tmp_path / "summary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "INPUT FILE SUMMARY"
    sheet["A2"] = "Gross"
    sheet["B2"] = 12.5
    sheet["B2"].number_format = "#,##0.00"
    sheet["C2"] = "=B2*2"
    sheet["A3"] = "Period End"
    sheet["B3"] = datetime(2026, 5, 31)
    sheet["B3"].number_format = "mm/dd/yyyy"
    workbook.save(path)

    result = read_cells(
        path,
        "Summary",
        "A1:D3",
        ["values", "formulas", "cachedValues", "numberFormats", "dataTypes"],
    )

    assert result["sheet"] == "Summary"
    assert result["range"] == "A1:D3"
    assert result["rowCount"] == 3
    assert result["columnCount"] == 4
    assert result["cellCount"] == 12
    assert result["returnedCellCount"] == len(result["cells"])

    cells = {cell["address"]: cell for cell in result["cells"]}
    assert cells["A1"] == {
        "address": "A1",
        "value": "INPUT FILE SUMMARY",
        "cachedValue": "INPUT FILE SUMMARY",
        "numberFormat": "General",
        "dataType": "s",
    }
    assert cells["B2"]["value"] == 12.5
    assert cells["B2"]["numberFormat"] == "#,##0.00"
    assert cells["C2"]["formula"] == "=B2*2"
    assert cells["C2"]["cachedValue"] is None
    assert "D2" not in cells
    assert cells["B3"]["value"] == "2026-05-31T00:00:00"
    assert cells["B3"]["numberFormat"] == "mm/dd/yyyy"


def test_read_cells_xml_engine_returns_bounded_values_and_formulas(
    tmp_path: Path,
) -> None:
    path = tmp_path / "summary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "Metric"
    sheet["B1"] = "Amount"
    sheet["A2"] = "FX Difference"
    sheet["B2"] = "=SUM(Merchant!L:L)"
    sheet["A3"] = "Status"
    sheet["B3"] = "OK"
    workbook.save(path)

    result = read_cells(
        path,
        "Summary",
        "A1:B3",
        ["values", "formulas", "cachedValues", "dataTypes"],
        engine="xml",
    )

    assert result["engine"] == "xml"
    assert result["sheet"] == "Summary"
    assert result["range"] == "A1:B3"
    assert result["cellCount"] == 6
    assert result["returnedCellCount"] == 6
    cells = {cell["address"]: cell for cell in result["cells"]}
    assert cells["A2"]["value"] == "FX Difference"
    assert cells["B2"]["formula"] == "=SUM(Merchant!L:L)"
    assert cells["B2"]["cachedValue"] is None
    assert cells["A2"]["dataType"] == "inlineStr"
    assert cells["B3"]["value"] == "OK"
    assert "row" not in cells["A2"]
    assert "column" not in cells["A2"]
    assert "formula" not in cells["A2"]


def test_read_cells_omits_blank_cells_from_upstream_fixture() -> None:
    path = Path("tests/fixtures/upstream/xlsx2csv/last-column-empty.xlsx")

    result = read_cells(path, "Sheet1", "A1:C5", ["values"])

    assert result["cellCount"] == 15
    assert result["returnedCellCount"] == len(result["cells"])
    assert result["returnedCellCount"] < 15
    cells = {cell["address"]: cell for cell in result["cells"]}
    # Blank cells are omitted; their absence plus cellCount states blankness.
    assert "C2" not in cells
    assert "C4" not in cells
    assert "C5" not in cells
    assert cells["A1"]["value"] is not None


def test_read_cells_xml_engine_translates_shared_formula_followers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "shared.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"
    sheet["A1"] = 1
    sheet["A2"] = 2
    sheet["A3"] = 3
    sheet["B1"] = "=A1*2"
    sheet["B2"] = "=A2*2"
    sheet["B3"] = "=A3*2"
    workbook.save(path)
    from zipfile import ZIP_DEFLATED, ZipFile

    with ZipFile(path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    worksheet_xml = entries["xl/worksheets/sheet1.xml"].decode("utf-8")
    worksheet_xml = worksheet_xml.replace(
        "<f>A1*2</f>",
        '<f t="shared" ref="B1:B3" si="0">A1*2</f>',
    )
    worksheet_xml = worksheet_xml.replace("<f>A2*2</f>", '<f t="shared" si="0"></f>')
    worksheet_xml = worksheet_xml.replace("<f>A3*2</f>", '<f t="shared" si="0"></f>')
    entries["xl/worksheets/sheet1.xml"] = worksheet_xml.encode("utf-8")
    with ZipFile(path, "w", ZIP_DEFLATED) as destination:
        for name, data in entries.items():
            destination.writestr(name, data)

    result = read_cells(path, "Sheet", "B1:B3", ["formulas"], engine="xml")

    cells = {cell["address"]: cell for cell in result["cells"]}
    assert cells["B1"]["formula"] == "=A1*2"
    assert cells["B2"]["formula"] == "=A2*2"
    assert cells["B3"]["formula"] == "=A3*2"


def test_read_cells_xml_translates_followers_when_master_is_outside_range(
    tmp_path: Path,
) -> None:
    path = tmp_path / "shared-out-of-range.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"
    for row in range(1, 6):
        sheet[f"A{row}"] = row
        sheet[f"B{row}"] = f"=A{row}*2"
    workbook.save(path)
    from zipfile import ZIP_DEFLATED, ZipFile

    with ZipFile(path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    worksheet_xml = entries["xl/worksheets/sheet1.xml"].decode("utf-8")
    worksheet_xml = worksheet_xml.replace(
        "<f>A1*2</f>",
        '<f t="shared" ref="B1:B5" si="0">A1*2</f>',
    )
    for row in range(2, 6):
        worksheet_xml = worksheet_xml.replace(
            f"<f>A{row}*2</f>",
            '<f t="shared" si="0"></f>',
        )
    entries["xl/worksheets/sheet1.xml"] = worksheet_xml.encode("utf-8")
    with ZipFile(path, "w", ZIP_DEFLATED) as destination:
        for name, data in entries.items():
            destination.writestr(name, data)

    # The master (B1) sits outside the requested range but must still anchor
    # follower translation.
    result = read_cells(path, "Sheet", "B4:B5", ["formulas"], engine="xml")

    cells = {cell["address"]: cell for cell in result["cells"]}
    assert cells["B4"]["formula"] == "=A4*2"
    assert cells["B5"]["formula"] == "=A5*2"


def test_read_cells_rejects_oversized_ranges(tmp_path: Path) -> None:
    path = tmp_path / "small.xlsx"
    workbook = Workbook()
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error_info:
        read_cells(path, "Sheet", "A1:CV101", ["values"])

    error = error_info.value
    assert error.error_code == "range_too_large"
    assert error.message == "Requested range is too large"
    assert error.details == {
        "range": "A1:CV101",
        "cellCount": 10100,
        "maxCells": 10000,
    }


def test_read_cells_rejects_unknown_include_fields(tmp_path: Path) -> None:
    path = tmp_path / "small.xlsx"
    workbook = Workbook()
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error_info:
        read_cells(path, "Sheet", "A1:A1", ["values", "styles"])

    error = error_info.value
    assert error.error_code == "invalid_field"
    assert error.details == {
        "field": "include",
        "expected": (
            "list[values | formulas | cachedValues | numberFormats | dataTypes]"
        ),
        "invalid": ["styles"],
    }
