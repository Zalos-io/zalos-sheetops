import json
from pathlib import Path

from openpyxl import Workbook

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def test_cells_read_command_returns_requested_range(tmp_path: Path) -> None:
    path = tmp_path / "summary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "INPUT FILE SUMMARY"
    sheet["A2"] = "Gross"
    sheet["B2"] = 12.5
    sheet["C2"] = "=B2*2"
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "cells",
            "read",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Summary",
                    "range": "A1:C2",
                    "include": ["values", "formulas", "cachedValues"],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["command"] == "cells read"
    assert payload["data"]["sheet"] == "Summary"
    assert payload["data"]["range"] == "A1:C2"
    assert payload["data"]["cellCount"] == 6
    cells = {cell["address"]: cell for cell in payload["data"]["cells"]}
    assert cells["A1"]["value"] == "INPUT FILE SUMMARY"
    assert cells["C2"]["formula"] == "=B2*2"
    assert cells["C2"]["cachedValue"] is None


def test_cells_read_command_accepts_xml_engine(tmp_path: Path) -> None:
    path = tmp_path / "summary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "FX Difference"
    sheet["B1"] = "=SUM(Merchant!L:L)"
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "cells",
            "read",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Summary",
                    "range": "A1:B1",
                    "include": ["values", "formulas", "cachedValues"],
                    "engine": "xml",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["data"]["engine"] == "xml"
    cells = {cell["address"]: cell for cell in payload["data"]["cells"]}
    assert cells["A1"]["value"] == "FX Difference"
    assert cells["B1"]["formula"] == "=SUM(Merchant!L:L)"


def test_cells_read_command_requires_sheet_field(tmp_path: Path) -> None:
    path = tmp_path / "summary.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "cells",
            "read",
            "--json",
            json.dumps({"path": str(path), "range": "A1:A1"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "cells read",
        "errorCode": "missing_required_field",
        "message": "Required field is missing or invalid",
        "details": {"field": "sheet"},
        "warnings": [],
        "artifacts": [],
    }


def test_cells_read_command_rejects_unknown_include_fields(tmp_path: Path) -> None:
    path = tmp_path / "summary.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "cells",
            "read",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Sheet",
                    "range": "A1:A1",
                    "include": ["values", "styles"],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "cells read",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {
            "field": "include",
            "expected": (
                "list[values | formulas | cachedValues | numberFormats | "
                "dataTypes | sourceValues]"
            ),
            "invalid": ["styles"],
        },
        "warnings": [],
        "artifacts": [],
    }
