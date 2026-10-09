from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.writer.theme import theme_xml


def test_alignment_matches_independent_cells_and_detects_real_changes() -> None:
    from openpyxl.styles import Alignment

    from zalos_sheetops.benchmarks.public.rubric import style_matches

    a, b = Workbook().active["A1"], Workbook().active["A1"]
    a.alignment = Alignment(horizontal="center", wrap_text=True)
    b.alignment = Alignment(horizontal="center", wrap_text=True)
    assert style_matches(a, a, ["alignment"])
    assert style_matches(a, b, ["alignment"])
    b.alignment = Alignment(horizontal="left", wrap_text=True)
    assert not style_matches(a, b, ["alignment"])


def test_visible_color_and_requested_font_components() -> None:
    from zalos_sheetops.benchmarks.public.rubric import style_matches

    expected, actual = Workbook(), Workbook()
    expected.loaded_theme = theme_xml.encode()
    a, b = expected.active["A1"], actual.active["A1"]
    a.font = Font(bold=True, name="Calibri", charset=134)
    b.font = Font(bold=True, name="Arial")
    a.fill = PatternFill("solid", fgColor=Color(theme=6, tint=0.799981688894314))
    b.fill = PatternFill("solid", fgColor="EBF1DE")
    assert style_matches(a, b, ["bold", "fill"])
    assert not style_matches(a, b, ["fontName"])
    b.fill = PatternFill("solid", fgColor="FFFF00")
    assert not style_matches(a, b, ["fill"])


def test_reviewed_rubric_distinguishes_preserving_styles_from_changing_them() -> None:
    from zalos_sheetops.benchmarks.public.rubric import reviewed_task

    assert reviewed_task({"id": "192-22"})["styleComponents"] == []
    assert reviewed_task({"id": 47827})["styleComponents"] == ["bold", "fill"]
    assert reviewed_task({"id": 50682})["styleComponents"] == []
    assert reviewed_task({"id": "343-20"})["styleComponents"] == ["alignment"]
    assert "id" in reviewed_task({"id": 1})


def test_theme_light_and_dark_indices_are_not_xml_element_positions() -> None:
    from zalos_sheetops.benchmarks.public.rubric import visible_color

    book = Workbook()
    book.loaded_theme = theme_xml.encode()
    assert visible_color(Color(theme=0), book.active["A1"]) == "FFFFFF"
    assert visible_color(Color(theme=1), book.active["A1"]) == "000000"


def test_excel_float_serialization_is_equivalent_without_masking_money_or_ids() -> None:
    from zalos_sheetops.benchmarks.public.rubric import values_equal

    assert values_equal(0.37336964500000003, 0.373369645)
    assert values_equal(0.17977659000000001, 0.17977659)
    assert not values_equal(1000000.01, 1000000.00)
    assert not values_equal(1000000000000001, 1000000000000000)
    assert not values_equal("001", 1)
    assert not values_equal(False, 0)


def test_agent_blank_values_match_without_treating_zero_as_blank() -> None:
    from zalos_sheetops.benchmarks.public.rubric import values_equal

    assert values_equal(None, "")
    assert values_equal("", None)
    assert not values_equal(None, 0)
    assert not values_equal(" ", "")


def test_unresolved_theme_color_remains_distinct_without_crashing() -> None:
    from zalos_sheetops.benchmarks.public.rubric import visible_color

    book = Workbook()
    book.loaded_theme = theme_xml.encode()
    assert visible_color(Color(theme=99), book.active["A1"]) == ("theme", 99, 0.0)


def test_reviewed_oracle_issues_are_not_product_failures() -> None:
    from zalos_sheetops.benchmarks.public.rubric import reviewed_task

    assert (
        reviewed_task({"id": "402-43"})["oracleIssue"]["status"]
        == "ambiguous_task_metadata"
    )
    assert (
        reviewed_task({"id": "395-36"})["oracleIssue"]["status"]
        == "ambiguous_task_metadata"
    )
    assert (
        reviewed_task({"id": "168-17"})["oracleIssue"]["status"]
        == "oracle_preservation_conflict"
    )


def test_reviewed_selection_keeps_unscorable_cases_and_balances_replacements() -> None:
    from zalos_sheetops.benchmarks.public.tasks import reviewed_selection

    pool = [
        {"id": "402-43", "instruction_type": "Sheet-Level Manipulation"},
        {"id": "a", "instruction_type": "Cell-Level Manipulation"},
        {"id": "b", "instruction_type": "Sheet-Level Manipulation"},
        {"id": "c", "instruction_type": "Sheet-Level Manipulation"},
        {"id": "d", "instruction_type": "Cell-Level Manipulation"},
    ]
    selected = [pool[0], pool[1]]
    result = reviewed_selection(pool, selected)
    assert result[:2] == selected
    assert len(result) == 3
    assert result[-1]["instruction_type"] == "Sheet-Level Manipulation"
