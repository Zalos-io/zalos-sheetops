from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.read_table import read_table
from zalos_sheetops.io.xlsx_xml import XlsxWorkbook


def test_read_table_detects_header_and_returns_records(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Report title", None])
    sheet.append([" Date ", " Amount "])
    sheet.append(["2026-05-01", "1,234.50"])
    workbook.save(path)

    result = read_table(path, "Data", "detect", None)

    assert result["sheet"] == "Data"
    assert result["headerRow"] == 2
    assert result["headers"] == ["date", "amount"]
    assert result["rows"] == [["2026-05-01", "1,234.50"]]


def test_read_table_detects_header_without_openpyxl_materialization(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Report title", None])
    sheet.append([" Date ", " Amount "])
    sheet.append(["2026-05-01", "1,234.50"])
    workbook.save(path)

    def fail_if_openpyxl_loads_workbook(*args: object, **kwargs: object) -> None:
        raise AssertionError("read-table should use XML streaming")

    monkeypatch.setattr(
        "zalos_sheetops.io.read_table.load_workbook",
        fail_if_openpyxl_loads_workbook,
        raising=False,
    )

    result = read_table(path, "Data", "detect", None)

    assert result["sheet"] == "Data"
    assert result["headerRow"] == 2
    assert result["headers"] == ["date", "amount"]
    assert result["rows"] == [["2026-05-01", "1,234.50"]]


def test_read_table_uses_explicit_header_row_selection(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["ignored", "ignored"])
    sheet.append(["Date", "Amount"])
    sheet.append(["2026-05-01", "25"])
    workbook.save(path)

    result = read_table(path, "Data", "row", 2)

    assert result["headerRow"] == 2
    assert result["rows"] == [["2026-05-01", "25"]]


def test_read_table_explicit_header_without_openpyxl_materialization(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["ignored", "ignored"])
    sheet.append(["Date", "Amount"])
    sheet.append(["2026-05-01", "25"])
    workbook.save(path)

    def fail_if_openpyxl_loads_workbook(*args: object, **kwargs: object) -> None:
        raise AssertionError("read-table should use XML streaming")

    monkeypatch.setattr(
        "zalos_sheetops.io.read_table.load_workbook",
        fail_if_openpyxl_loads_workbook,
        raising=False,
    )

    result = read_table(path, "Data", "row", 2)

    assert result["headerRow"] == 2
    assert result["rows"] == [["2026-05-01", "25"]]


def test_read_table_limits_returned_rows_and_stops_counting(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    sheet.append(["2026-05-01", "10"])
    sheet.append(["2026-05-02", "20"])
    sheet.append(["2026-05-03", "30"])
    workbook.save(path)

    result = read_table(path, "Data", "row", 1, max_rows=2)

    assert result["rowCount"] is None
    assert result["returnedRowCount"] == 2
    assert result["truncated"] is True
    assert result["rows"] == [
        ["2026-05-01", "10"],
        ["2026-05-02", "20"],
    ]


def test_read_table_reports_exact_row_count_when_max_rows_reaches_end(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    sheet.append(["2026-05-01", "10"])
    sheet.append(["2026-05-02", "20"])
    workbook.save(path)

    result = read_table(path, "Data", "row", 1, max_rows=5)

    assert result["rowCount"] == 2
    assert result["returnedRowCount"] == 2
    assert result["truncated"] is False


def test_read_table_row_strategy_stops_streaming_at_max_rows(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    for index in range(1, 51):
        sheet.append([f"2026-05-{index:02d}", str(index)])
    workbook.save(path)
    consumed = {"rows": 0}
    original_iter_rows = XlsxWorkbook.iter_rows

    def counting_iter_rows(self: XlsxWorkbook, sheet_name: str) -> object:
        for row in original_iter_rows(self, sheet_name):
            consumed["rows"] += 1
            yield row

    monkeypatch.setattr(XlsxWorkbook, "iter_rows", counting_iter_rows)

    result = read_table(path, "Data", "row", 1, max_rows=3)

    assert result["rowCount"] is None
    assert result["returnedRowCount"] == 3
    assert result["truncated"] is True
    assert consumed["rows"] <= 5


def test_read_table_rejects_unknown_sheet_name(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        read_table(path, "Missing", "detect", None)

    assert error.value.error_code == "invalid_sheet"
    assert error.value.details == {"sheet": "Missing"}


def test_read_table_rejects_out_of_range_header_row(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        read_table(path, "Data", "row", 3)

    assert error.value.error_code == "invalid_header_row"
    assert error.value.details == {"headerRow": 3, "sheet": "Data", "maxRow": 1}


def test_read_table_rejects_non_positive_header_row(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        read_table(path, "Data", "row", 0)

    assert error.value.error_code == "invalid_header_row"
    assert error.value.details == {"headerRow": 0, "sheet": "Data", "maxRow": 1}


def test_read_table_rejects_duplicate_canonical_headers(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Amount", " amount "])
    sheet.append(["10", "20"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        read_table(path, "Data", "row", 1)

    assert error.value.error_code == "invalid_headers"
    assert error.value.details == {"duplicates": ["amount"], "blankColumns": []}


def test_read_table_rejects_blank_canonical_headers(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "   "])
    sheet.append(["2026-05-01", "10"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        read_table(path, "Data", "row", 1)

    assert error.value.error_code == "invalid_headers"
    assert error.value.details == {"duplicates": [], "blankColumns": [2]}


def test_read_table_rejects_detect_mode_when_no_header_can_be_found(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append([None, None])
    sheet.append(["   ", None])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        read_table(path, "Data", "detect", None)

    assert error.value.error_code == "header_not_found"
    assert error.value.details == {"sheet": "Data", "headerStrategy": "detect"}


def test_read_table_rejects_invalid_header_strategy(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        read_table(path, "Data", "auto", None)

    assert error.value.error_code == "invalid_header_strategy"
    assert error.value.details == {"headerStrategy": "auto"}


def test_read_table_rejects_ambiguous_detect_mode_when_metadata_row_is_denser(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Quarterly", "Revenue", "Draft", "Review", "FY26", "Final"])
    sheet.append(["Date", "Amount", "Status", "Region", "Owner", "Source"])
    sheet.append(["2026-05-01", "10", "posted", "emea", "ops", "erp"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        read_table(path, "Data", "detect", None)

    assert error.value.error_code == "ambiguous_header_row"
    assert error.value.details == {
        "sheet": "Data",
        "headerStrategy": "detect",
        "headerRows": [1, 2],
    }


def test_read_table_row_strategy_still_reads_hostile_fixture(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Quarterly", "Revenue", "Draft", "Review", "FY26", "Final"])
    sheet.append(["Date", "Amount", "Status", "Region", "Owner", "Source"])
    sheet.append(["2026-05-01", "10", "posted", "emea", "ops", "erp"])
    workbook.save(path)

    result = read_table(path, "Data", "row", 2)

    assert result["headerRow"] == 2
    assert result["headers"] == [
        "date",
        "amount",
        "status",
        "region",
        "owner",
        "source",
    ]
    assert result["rows"] == [["2026-05-01", "10", "posted", "emea", "ops", "erp"]]


def test_read_table_rejects_invalid_workbook_content(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("not really an xlsx", encoding="utf-8")

    with pytest.raises(SheetOpsError) as error:
        read_table(path, None, "detect", None)

    assert error.value.error_code == "invalid_workbook"
    assert error.value.details == {"path": str(path)}


def test_read_table_rejects_malformed_zip_valid_workbook(tmp_path: Path) -> None:
    path = tmp_path / "broken-xml.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types")

    with pytest.raises(SheetOpsError) as error:
        read_table(path, None, "detect", None)

    assert error.value.error_code == "invalid_workbook"
    assert error.value.details == {"path": str(path)}
