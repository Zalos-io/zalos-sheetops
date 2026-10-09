import json
from pathlib import Path

from openpyxl import Workbook

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def test_formulas_scan_command_returns_nonzero_when_error_tokens_exist(
    tmp_path: Path,
) -> None:
    path = tmp_path / "formulas.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "#VALUE!"
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        ["formulas", "scan", "--json", json.dumps({"path": str(path)}), "--raw-output"],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "formula_errors_found",
        "message": "Formula error tokens were found",
        "details": {
            "engine": "xml",
            "include": ["formulas", "errors"],
            "formulaCount": 0,
            "errorTokenCount": 1,
            "returnedFormulaCount": 0,
            "returnedErrorTokenCount": 1,
            "truncated": False,
            "stoppedEarly": False,
            "scope": {"sheet": None, "range": None},
            "formulas": [],
            "errorTokens": [{"sheet": "Sheet", "cell": "A1", "value": "#VALUE!"}],
        },
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_accepts_sheet_range_and_max_results() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "formulas",
            "scan",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Sheet1",
                    "range": "A1:A4",
                    "maxResults": 2,
                    "engine": "xml",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["command"] == "formulas scan"
    assert payload["data"]["engine"] == "xml"
    assert payload["data"]["include"] == ["formulas", "errors"]
    assert payload["data"]["stoppedEarly"] is False
    assert payload["data"]["scope"] == {"sheet": "Sheet1", "range": "A1:A4"}
    assert payload["data"]["formulaCount"] == 4
    assert payload["data"]["returnedFormulaCount"] == 2
    assert payload["data"]["errorTokenCount"] == 0
    assert payload["data"]["truncated"] is True
    assert [entry["cell"] for entry in payload["data"]["formulas"]] == ["A1", "A2"]


def test_formulas_scan_command_can_force_openpyxl_engine() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "formulas",
            "scan",
            "--json",
            json.dumps({"path": str(path), "engine": "openpyxl", "maxResults": 1}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["errorCode"] == "formula_errors_found"
    assert payload["details"]["engine"] == "openpyxl"
    assert payload["details"]["formulaCount"] == 12
    assert payload["details"]["returnedFormulaCount"] == 1


def test_formulas_scan_command_accepts_include_and_fail_fast() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "formulas",
            "scan",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "include": ["errors"],
                    "failFast": True,
                    "engine": "xml",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["errorCode"] == "formula_errors_found"
    assert payload["details"]["engine"] == "xml"
    assert payload["details"]["include"] == ["errors"]
    assert payload["details"]["stoppedEarly"] is True
    assert payload["details"]["formulaCount"] == 0
    assert payload["details"]["errorTokenCount"] == 1
    assert payload["details"]["errorTokens"] == [
        {"sheet": "Sheet1", "cell": "A5", "value": "#DIV/0!"},
    ]


def test_formulas_scan_command_rejects_invalid_engine(tmp_path: Path) -> None:
    path = tmp_path / "formulas.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "formulas",
            "scan",
            "--json",
            json.dumps({"path": str(path), "engine": "fast"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "engine", "expected": "auto | xml | openpyxl"},
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_rejects_invalid_include(tmp_path: Path) -> None:
    path = tmp_path / "formulas.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "formulas",
            "scan",
            "--json",
            json.dumps({"path": str(path), "include": ["formulas", "values"]}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {
            "field": "include",
            "expected": "non-empty list[formulas | errors]",
        },
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_rejects_invalid_fail_fast(tmp_path: Path) -> None:
    path = tmp_path / "formulas.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "formulas",
            "scan",
            "--json",
            json.dumps({"path": str(path), "failFast": "yes"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "failFast", "expected": "boolean"},
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_rejects_invalid_max_results(tmp_path: Path) -> None:
    path = tmp_path / "formulas.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "formulas",
            "scan",
            "--json",
            json.dumps({"path": str(path), "maxResults": 0}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "maxResults", "expected": "positive integer | null"},
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_rejects_malformed_json() -> None:
    runner = CliRunner()

    result = runner.invoke(
        app,
        ["formulas", "scan", "--json", "{", "--raw-output"],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "invalid_json",
        "message": "Request payload is not valid JSON",
        "details": {"position": 1},
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_requires_path_field() -> None:
    runner = CliRunner()

    result = runner.invoke(
        app,
        ["formulas", "scan", "--json", json.dumps({}), "--raw-output"],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "missing_required_field",
        "message": "Required field is missing or invalid",
        "details": {"field": "path"},
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_rejects_invalid_path() -> None:
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "formulas",
            "scan",
            "--json",
            json.dumps({"path": "../secrets.xlsx"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "invalid_path",
        "message": "Path contains disallowed characters or traversal",
        "details": {"path": "../secrets.xlsx"},
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_rejects_missing_file(tmp_path: Path) -> None:
    path = tmp_path / "missing.xlsx"
    runner = CliRunner()

    result = runner.invoke(
        app,
        ["formulas", "scan", "--json", json.dumps({"path": str(path)}), "--raw-output"],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "file_not_found",
        "message": "Workbook file does not exist",
        "details": {"path": str(path)},
        "warnings": [],
        "artifacts": [],
    }


def test_formulas_scan_command_rejects_invalid_workbook_content(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("not really an xlsx", encoding="utf-8")
    runner = CliRunner()

    result = runner.invoke(
        app,
        ["formulas", "scan", "--json", json.dumps({"path": str(path)}), "--raw-output"],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "formulas scan",
        "errorCode": "invalid_workbook",
        "message": "Workbook could not be opened",
        "details": {"path": str(path)},
        "warnings": [],
        "artifacts": [],
    }
