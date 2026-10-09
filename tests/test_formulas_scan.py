from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.formulas import xml_scan
from zalos_sheetops.formulas.scan import scan_formulas


def test_scan_formulas_reports_formula_cells_and_error_literals(tmp_path: Path) -> None:
    path = tmp_path / "formulas.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recon"
    sheet["A1"] = "Amount"
    sheet["A2"] = 10
    sheet["B2"] = "=A2*2"
    sheet["C2"] = "#REF!"
    workbook.save(path)

    result = scan_formulas(path)

    assert result["include"] == ["formulas", "errors"]
    assert result["stoppedEarly"] is False
    assert result["formulaCount"] == 1
    assert result["returnedFormulaCount"] == 1
    assert result["errorTokenCount"] == 1
    assert result["returnedErrorTokenCount"] == 1
    assert result["truncated"] is False
    assert result["formulas"] == [
        {"sheet": "Recon", "cell": "B2", "formula": "=A2*2"},
    ]
    assert result["errorTokens"] == [
        {"sheet": "Recon", "cell": "C2", "value": "#REF!"},
    ]


def test_scan_formulas_reports_cached_formula_result_errors() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    result = scan_formulas(path)

    assert result["engine"] == "xml"
    assert result["include"] == ["formulas", "errors"]
    assert result["stoppedEarly"] is False
    assert result["formulaCount"] == 12
    assert result["errorTokenCount"] == 8
    assert result["errorTokens"] == [
        {"sheet": "Sheet1", "cell": "A5", "value": "#DIV/0!"},
        {"sheet": "Sheet1", "cell": "A6", "value": "#N/A"},
        {"sheet": "Sheet1", "cell": "A7", "value": "#NAME?"},
        {"sheet": "Sheet1", "cell": "A8", "value": "#NULL!"},
        {"sheet": "Sheet1", "cell": "A9", "value": "#NUM!"},
        {"sheet": "Sheet1", "cell": "A10", "value": "#REF!"},
        {"sheet": "Sheet1", "cell": "A11", "value": "#VALUE!"},
        {"sheet": "Sheet1", "cell": "A12", "value": "#DIV/0!"},
    ]


def test_scan_formulas_can_scope_to_a_sheet_range() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    result = scan_formulas(
        path,
        sheet_name="Sheet1",
        cell_range="A1:A4",
        engine="xml",
    )

    assert result["engine"] == "xml"
    assert result["scope"] == {"sheet": "Sheet1", "range": "A1:A4"}
    assert result["formulaCount"] == 4
    assert result["errorTokenCount"] == 0
    assert [entry["cell"] for entry in result["formulas"]] == ["A1", "A2", "A3", "A4"]


def test_scan_formulas_applies_max_results_without_hiding_total_counts() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    result = scan_formulas(path, max_results=2, engine="xml")

    assert result["formulaCount"] == 12
    assert result["returnedFormulaCount"] == 2
    assert result["errorTokenCount"] == 8
    assert result["returnedErrorTokenCount"] == 2
    assert result["truncated"] is True
    assert [entry["cell"] for entry in result["formulas"]] == ["A1", "A2"]
    assert [entry["cell"] for entry in result["errorTokens"]] == ["A5", "A6"]


def test_scan_formulas_xml_engine_counts_newer_excel_error_tokens(
    tmp_path: Path,
) -> None:
    path = tmp_path / "newer-errors.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Errors"
    sheet["A1"] = "#SPILL!"
    sheet["A2"] = "#CALC!"
    workbook.save(path)

    result = scan_formulas(path, engine="xml")

    assert result["engine"] == "xml"
    assert result["errorTokenCount"] == 2
    assert result["errorTokens"] == [
        {"sheet": "Errors", "cell": "A1", "value": "#SPILL!"},
        {"sheet": "Errors", "cell": "A2", "value": "#CALC!"},
    ]


def test_scan_formulas_xml_engine_translates_shared_formula_followers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "shared-formulas.xlsx"
    _write_shared_formula_fixture(path)

    result = scan_formulas(path, engine="xml")

    assert result["engine"] == "xml"
    assert result["formulaCount"] == 3
    assert result["formulas"] == [
        {"sheet": "Sheet", "cell": "B1", "formula": "=A1*2"},
        {"sheet": "Sheet", "cell": "B2", "formula": "=A2*2"},
        {"sheet": "Sheet", "cell": "B3", "formula": "=A3*2"},
    ]


def test_scan_formulas_xml_translation_matches_openpyxl_engine(
    tmp_path: Path,
) -> None:
    path = tmp_path / "shared-formulas.xlsx"
    _write_shared_formula_fixture(path)

    xml_result = scan_formulas(path, engine="xml", include=["formulas"])
    openpyxl_result = scan_formulas(path, engine="openpyxl", include=["formulas"])

    assert xml_result["formulas"] == openpyxl_result["formulas"]


def test_scan_formulas_xml_engine_rejects_orphan_shared_followers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "orphan-shared-formulas.xlsx"
    _write_orphan_shared_formula_fixture(path)

    with pytest.raises(SheetOpsError) as error_info:
        scan_formulas(path, engine="xml")

    error = error_info.value
    assert error.error_code == "unsupported_formula_scan_engine"
    assert error.message == "XML formula scan cannot handle this workbook feature"
    assert error.details == {
        "engine": "xml",
        "feature": "shared_formula_without_text",
        "sheet": "Sheet",
        "cell": "B2",
    }


def test_scan_formulas_can_force_openpyxl_engine() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    result = scan_formulas(path, engine="openpyxl")

    assert result["engine"] == "openpyxl"
    assert result["formulaCount"] == 12
    assert result["errorTokenCount"] == 8


def test_scan_formulas_can_include_only_formula_cells() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    result = scan_formulas(path, include=["formulas"], engine="xml")

    assert result["engine"] == "xml"
    assert result["include"] == ["formulas"]
    assert result["formulaCount"] == 12
    assert result["errorTokenCount"] == 0
    assert result["returnedErrorTokenCount"] == 0
    assert result["errorTokens"] == []


def test_scan_formulas_can_include_only_error_tokens() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    result = scan_formulas(path, include=["errors"], engine="xml")

    assert result["engine"] == "xml"
    assert result["include"] == ["errors"]
    assert result["formulaCount"] == 0
    assert result["formulas"] == []
    assert result["errorTokenCount"] == 8
    assert [entry["cell"] for entry in result["errorTokens"]] == [
        "A5",
        "A6",
        "A7",
        "A8",
        "A9",
        "A10",
        "A11",
        "A12",
    ]


def test_scan_formulas_fail_fast_stops_after_first_error_token() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    result = scan_formulas(
        path,
        include=["errors"],
        engine="xml",
        fail_fast=True,
    )

    assert result["engine"] == "xml"
    assert result["include"] == ["errors"]
    assert result["stoppedEarly"] is True
    assert result["formulaCount"] == 0
    assert result["errorTokenCount"] == 1
    assert result["returnedErrorTokenCount"] == 1
    assert result["truncated"] is False
    assert result["errorTokens"] == [
        {"sheet": "Sheet1", "cell": "A5", "value": "#DIV/0!"},
    ]


def test_scan_formulas_openpyxl_engine_honors_fail_fast() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    result = scan_formulas(
        path,
        include=["errors"],
        engine="openpyxl",
        fail_fast=True,
    )

    assert result["engine"] == "openpyxl"
    assert result["stoppedEarly"] is True
    assert result["errorTokenCount"] == 1
    assert result["errorTokens"] == [
        {"sheet": "Sheet1", "cell": "A5", "value": "#DIV/0!"},
    ]


def test_scan_formulas_errors_only_xml_scan_does_not_fallback_on_shared_formulas(
    tmp_path: Path,
) -> None:
    path = tmp_path / "shared-formulas.xlsx"
    _write_shared_formula_fixture(path)

    result = scan_formulas(path, include=["errors"], engine="xml")

    assert result["engine"] == "xml"
    assert result["formulaCount"] == 0
    assert result["errorTokenCount"] == 0


def test_scan_formulas_xml_engine_skips_sheets_without_formula_or_error_markers(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "skip-clean-sheets.xlsx"
    workbook = Workbook()
    clean = workbook.active
    clean.title = "Clean"
    clean["A1"] = "No formulas here"
    clean["A2"] = 10
    formula = workbook.create_sheet("Formula")
    formula["A1"] = 5
    formula["B1"] = "=A1*2"
    workbook.save(path)

    scanned_sheets: list[str] = []
    original_scan_worksheet = xml_scan._scan_worksheet

    def recording_scan_worksheet(*args: object) -> dict[str, object]:
        worksheet = args[1]
        scanned_sheets.append(worksheet.name)
        return original_scan_worksheet(*args)

    monkeypatch.setattr(xml_scan, "_scan_worksheet", recording_scan_worksheet)

    result = scan_formulas(path, engine="xml")

    assert result["formulaCount"] == 1
    assert scanned_sheets == ["Formula"]


def test_scan_formulas_max_results_caps_items_across_sheets(
    tmp_path: Path,
) -> None:
    path = tmp_path / "multi-sheet.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    for row in range(1, 4):
        first[f"A{row}"] = f"=1*{row}"
    second = workbook.create_sheet("Second")
    for row in range(1, 4):
        second[f"A{row}"] = f"=2*{row}"
    workbook.save(path)

    result = scan_formulas(path, max_results=4, engine="xml")

    assert result["formulaCount"] == 6
    assert result["returnedFormulaCount"] == 4
    assert result["truncated"] is True
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    with pytest.raises(SheetOpsError) as error_info:
        scan_formulas(path, engine="fast")

    error = error_info.value
    assert error.error_code == "invalid_field"
    assert error.message == "Field has invalid type"
    assert error.details == {
        "field": "engine",
        "expected": "auto | xml | openpyxl",
    }


def test_scan_formulas_rejects_invalid_include() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx")

    with pytest.raises(SheetOpsError) as error_info:
        scan_formulas(path, include=["formulas", "values"])

    error = error_info.value
    assert error.error_code == "invalid_field"
    assert error.message == "Field has invalid type"
    assert error.details == {
        "field": "include",
        "expected": "non-empty list[formulas | errors]",
    }


def test_scan_formulas_rejects_invalid_workbook_content(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("not really an xlsx", encoding="utf-8")

    with pytest.raises(SheetOpsError) as error_info:
        scan_formulas(path)

    error = error_info.value
    assert error.error_code == "invalid_workbook"
    assert error.message == "Workbook could not be opened"
    assert error.details == {"path": str(path)}


def test_scan_formulas_xml_engine_rejects_incomplete_xlsx_with_stable_error(
    tmp_path: Path,
) -> None:
    path = tmp_path / "incomplete.xlsx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types></Types>")

    with pytest.raises(SheetOpsError) as error_info:
        scan_formulas(path, engine="xml")

    error = error_info.value
    assert error.error_code == "invalid_workbook"
    assert error.message == "Workbook could not be opened"
    assert error.details == {"path": str(path)}


def _write_shared_formula_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"
    sheet["A1"] = 1
    sheet["A2"] = 2
    sheet["A3"] = 3
    sheet["B1"] = "=A1*2"
    sheet["B2"] = "=A2*2"
    sheet["B3"] = "=A3*2"
    workbook.save(path)

    with ZipFile(path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}

    worksheet_xml = entries["xl/worksheets/sheet1.xml"].decode("utf-8")
    worksheet_xml = worksheet_xml.replace(
        "<f>A1*2</f>",
        '<f t="shared" ref="B1:B3" si="0">A1*2</f>',
    )
    worksheet_xml = worksheet_xml.replace(
        "<f>A2*2</f>",
        '<f t="shared" si="0"></f>',
    )
    worksheet_xml = worksheet_xml.replace(
        "<f>A3*2</f>",
        '<f t="shared" si="0"></f>',
    )
    entries["xl/worksheets/sheet1.xml"] = worksheet_xml.encode("utf-8")

    with ZipFile(path, "w", ZIP_DEFLATED) as destination:
        for name, data in entries.items():
            destination.writestr(name, data)


def _write_orphan_shared_formula_fixture(path: Path) -> None:
    """Shared-formula followers whose master carries no formula text."""
    _write_shared_formula_fixture(path)

    with ZipFile(path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}

    worksheet_xml = entries["xl/worksheets/sheet1.xml"].decode("utf-8")
    worksheet_xml = worksheet_xml.replace(
        '<f t="shared" ref="B1:B3" si="0">A1*2</f>',
        "",
    )
    entries["xl/worksheets/sheet1.xml"] = worksheet_xml.encode("utf-8")

    with ZipFile(path, "w", ZIP_DEFLATED) as destination:
        for name, data in entries.items():
            destination.writestr(name, data)
