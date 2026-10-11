import json
from pathlib import Path

from openpyxl import Workbook

from zalos_sheetops.hermes_plugin import SpreadsheetDelivery


def test_cancelled_prepared_policy_is_closed_and_not_resumed(
    tmp_path: Path,
    monkeypatch: object,
) -> None:
    import sys
    from types import ModuleType

    from zalos_sheetops.hermes_plugin import register

    config = {"cwd": str(tmp_path), "env_type": "local"}
    for name, attributes in {
        "agent": {},
        "tools": {},
        "agent.message_content": {"flatten_message_text": str},
        "tools.terminal_tool": {"_get_env_config": lambda: config},
        "tools.terminal_tool_lifecycle": {"ensure_task_env": lambda _: None},
    }.items():
        module = ModuleType(name)
        vars(module).update(attributes)
        monkeypatch.setitem(sys.modules, name, module)

    class Context:
        def __init__(self) -> None:
            self.hooks = {}

        def register_hook(self, name: str, callback: object) -> None:
            self.hooks[name] = callback

        def get_config(self, _name: str) -> str:
            return "unavailable-test-engine"

    closed = []
    close = SpreadsheetDelivery.close

    def tracked_close(runtime: SpreadsheetDelivery) -> None:
        roots = [
            state["policy"]._root for state in runtime.turns.values() if state["policy"]
        ]
        close(runtime)
        closed.extend(roots)
        assert all(not root.exists() for root in roots)

    monkeypatch.setattr(SpreadsheetDelivery, "close", tracked_close)
    context = Context()
    register(context)
    source = tmp_path / "input.xlsx"
    Workbook().save(source)
    context.hooks["prepare_delivery"](
        task_id="task", turn_id="one", user_message=str(source), session_id="session"
    )
    # The per-turn filename is deterministic and appears in the prepare guidance.
    import hashlib

    contract = tmp_path / (
        ".sheetops-delivery-" + hashlib.sha256(b"one").hexdigest()[:20] + ".json"
    )
    contract.write_text(
        json.dumps(
            {
                "source": str(source),
                "output": "output.xlsx",
                "allowedRanges": {"Sheet": ["A1"]},
            }
        )
    )
    context.hooks["pre_tool_call"](turn_id="one", tool_name="terminal")
    context.hooks["finish_delivery"](turn_id="one", interrupted=True)
    assert len(closed) == 1
    assert context.hooks["validate_delivery"](turn_id="one")["action"] == "block"
    prepared = context.hooks["prepare_delivery"](
        task_id="task", turn_id="two", user_message="continue", session_id="session"
    )
    assert prepared["message"] != "Delivery retains the prior frozen contract."
    context.hooks["finish_delivery"](turn_id="two", interrupted=True)


def test_skipped_guard_and_late_scope_cannot_release_output(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    Workbook().save(source)
    runtime = SpreadsheetDelivery(cwd=tmp_path, engine=None)
    assert runtime.prepare("one", str(source))["active"]
    output = tmp_path / "output.xlsx"
    output.write_bytes(source.read_bytes())
    assert runtime.validate("one", "output.xlsx")["action"] == "block"
    (tmp_path / ".sheetops-delivery.json").write_text(
        json.dumps({"output": "output.xlsx", "allowedRanges": {"Sheet": ["A1"]}})
    )
    assert runtime.validate("one", "output.xlsx")["action"] == "block"
    runtime.close()


def test_scope_freezes_before_edit_and_approval_binds_served_bytes(
    tmp_path: Path,
) -> None:
    source = tmp_path / "input.xlsx"
    Workbook().save(source)
    runtime = SpreadsheetDelivery(cwd=tmp_path, engine=None)
    runtime.prepare("one", str(source))
    contract = tmp_path / ".sheetops-delivery.json"
    contract.write_text(
        json.dumps({"output": "output.xlsx", "allowedRanges": {"Sheet": ["A1"]}})
    )
    runtime.observe("one")
    output = tmp_path / "output.xlsx"
    output.write_bytes(source.read_bytes())
    assert runtime.validate("one")["action"] == "allow"
    assert runtime.authorize(output, output.read_bytes())["action"] == "allow"
    contract.write_text(
        json.dumps({"output": "output.xlsx", "allowedRanges": {"Sheet": ["A1:Z100"]}})
    )
    assert runtime.validate("one", "output.xlsx")["action"] == "block"
    assert runtime.authorize(output, output.read_bytes())["action"] == "block"
    runtime.close()


def test_pathless_creation_still_requires_a_frozen_contract(tmp_path: Path) -> None:
    runtime = SpreadsheetDelivery(cwd=tmp_path, engine=None)
    assert runtime.prepare("one", "Create a spreadsheet")["active"]
    assert runtime.validate("one")["action"] == "block"  # creation requires a contract
    output = tmp_path / "output.xlsx"
    Workbook().save(output)
    assert runtime.validate("one", "output.xlsx")["action"] == "block"
    runtime.close()


def test_creation_contract_freezes_at_the_next_tool_boundary(tmp_path: Path) -> None:
    runtime = SpreadsheetDelivery(cwd=tmp_path, engine=None)
    runtime.prepare("one", "Create a spreadsheet")
    contract = tmp_path / ".sheetops-delivery.json"
    contract.write_text(
        json.dumps({"output": "output.xlsx", "allowedRanges": {"Sheet": ["A1"]}})
    )
    assert runtime.before_tool("one")["action"] == "allow"
    book = Workbook()
    book.active["A1"] = 5
    book.save(tmp_path / "output.xlsx")
    assert runtime.validate("one")["action"] == "allow"
    contract.write_text(
        json.dumps({"output": "output.xlsx", "allowedRanges": {"Sheet": ["A1:Z100"]}})
    )
    assert runtime.before_tool("one")["action"] == "block"
    assert runtime.validate("one", "output.xlsx")["action"] == "block"
    runtime.close()


def test_unrelated_workbooks_do_not_block_ordinary_text(tmp_path: Path) -> None:
    (tmp_path / "unrelated.xlsx").write_bytes(b"not a readable workbook")
    runtime = SpreadsheetDelivery(cwd=tmp_path, engine=None)
    assert runtime.prepare("one", "What is 2+2?")["active"]
    assert runtime.validate("one", "4")["action"] == "allow"
    assert (
        runtime.authorize(tmp_path / "unrelated.xlsx", b"anything")["action"] == "block"
    )
    runtime.close()


def test_missing_engine_retry_keeps_scope_and_evaluates_when_engine_appears(
    tmp_path: Path,
) -> None:
    import pytest
    from openpyxl import load_workbook

    engine = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if not engine.is_file():
        pytest.skip("LibreOffice integration requires an installed engine")
    source = tmp_path / "input.xlsx"
    book = Workbook()
    book.active["A1"] = 1
    book.active["B1"] = "=A1+1"
    book.save(source)
    runtime = SpreadsheetDelivery(cwd=tmp_path, engine=None)
    runtime.prepare("one", f"Edit {source}")
    (tmp_path / ".sheetops-delivery.json").write_text(
        json.dumps(
            {
                "source": str(source),
                "output": "output.xlsx",
                "allowedRanges": {"Sheet": ["A1"]},
            }
        )
    )
    runtime.before_tool("one")
    output = tmp_path / "output.xlsx"
    output.write_bytes(source.read_bytes())
    rejected = runtime.validate("one", str(output))
    assert rejected["errorCode"] == "calculation_engine_unavailable"
    assert "no calculation engine is available" in rejected["publicMessage"]
    runtime.refresh_engine(str(engine))
    assert runtime.validate("one", str(output))["action"] == "allow"
    assert load_workbook(output, data_only=True).active["B1"].value == 2
    (tmp_path / ".sheetops-delivery.json").write_text(
        json.dumps(
            {
                "source": str(source),
                "output": "output.xlsx",
                "allowedRanges": {"Sheet": ["A1:B1"]},
            }
        )
    )
    assert runtime.validate("one", str(output))["action"] == "block"
    runtime.close()
