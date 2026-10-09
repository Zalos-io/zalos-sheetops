import json
from pathlib import Path

from openpyxl import Workbook

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def test_read_table_command_returns_json_records(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Title", None])
    sheet.append(["Date", "Amount"])
    sheet.append(["2026-05-01", "10"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "detect",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["command"] == "read-table"
    assert payload["data"]["headers"] == ["date", "amount"]
    assert payload["data"]["rows"] == [["2026-05-01", "10"]]


def test_read_table_command_limits_returned_rows(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    sheet.append(["2026-05-01", "10"])
    sheet.append(["2026-05-02", "20"])
    sheet.append(["2026-05-03", "30"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "row",
                    "headerRow": 1,
                    "maxRows": 2,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["data"]["rowCount"] is None
    assert payload["data"]["returnedRowCount"] == 2
    assert payload["data"]["truncated"] is True
    assert payload["data"]["rows"] == [
        ["2026-05-01", "10"],
        ["2026-05-02", "20"],
    ]


def test_read_table_command_rejects_invalid_max_rows(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "row",
                    "headerRow": 1,
                    "maxRows": 0,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "maxRows", "expected": "positive integer | null"},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_invalid_sheet_name(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Missing",
                    "headerStrategy": "detect",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_sheet",
        "message": "Requested sheet does not exist",
        "details": {"sheet": "Missing"},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_out_of_range_header_row(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "row",
                    "headerRow": 3,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_header_row",
        "message": "Header row is outside the sheet bounds",
        "details": {"headerRow": 3, "sheet": "Data", "maxRow": 1},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_non_positive_header_row(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "row",
                    "headerRow": 0,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_header_row",
        "message": "Header row is outside the sheet bounds",
        "details": {"headerRow": 0, "sheet": "Data", "maxRow": 1},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_boolean_header_row(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "row",
                    "headerRow": True,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "headerRow", "expected": "integer | null"},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_duplicate_headers(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Amount", " amount "])
    sheet.append(["10", "20"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "row",
                    "headerRow": 1,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_headers",
        "message": "Header row contains duplicate or blank columns",
        "details": {"duplicates": ["amount"], "blankColumns": []},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_blank_headers(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "   "])
    sheet.append(["2026-05-01", "10"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "row",
                    "headerRow": 1,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_headers",
        "message": "Header row contains duplicate or blank columns",
        "details": {"duplicates": [], "blankColumns": [2]},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_detect_mode_when_no_header_can_be_found(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append([None, None])
    sheet.append(["   ", None])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "detect",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "header_not_found",
        "message": "Unable to detect header row",
        "details": {"sheet": "Data", "headerStrategy": "detect"},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_invalid_header_strategy(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "auto",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_header_strategy",
        "message": "Header strategy must be 'detect' or 'row'",
        "details": {"headerStrategy": "auto"},
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_ambiguous_detect_mode_when_metadata_row_is_denser(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Quarterly", "Revenue", "Draft", "Review", "FY26", "Final"])
    sheet.append(["Date", "Amount", "Status", "Region", "Owner", "Source"])
    sheet.append(["2026-05-01", "10", "posted", "emea", "ops", "erp"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Data",
                    "headerStrategy": "detect",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "ambiguous_header_row",
        "message": "Multiple header rows are plausible in detect mode",
        "details": {
            "sheet": "Data",
            "headerStrategy": "detect",
            "headerRows": [1, 2],
        },
        "warnings": [],
        "artifacts": [],
    }


def test_read_table_command_rejects_invalid_workbook_content(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("not really an xlsx", encoding="utf-8")
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "read-table",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "headerStrategy": "detect",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload == {
        "ok": False,
        "command": "read-table",
        "errorCode": "invalid_workbook",
        "message": "Workbook could not be opened",
        "details": {"path": str(path)},
        "warnings": [],
        "artifacts": [],
    }
