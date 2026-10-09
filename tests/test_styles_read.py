from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Color,
    Font,
    PatternFill,
    Protection,
    Side,
)

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.styles.read import read_styles

UPSTREAM_COMPLEX_STYLES_FIXTURE = Path(
    "tests/fixtures/upstream/openpyxl/complex-styles.xlsx"
)


def test_read_styles_extracts_upstream_font_and_fill_metadata() -> None:
    result = read_styles(UPSTREAM_COMPLEX_STYLES_FIXTURE, "Sheet1", "A1:I3")

    cells = {cell["address"]: cell for cell in result["cells"]}
    assert cells["A3"]["style"]["font"] == {
        "name": "Arial",
        "size": 12.0,
        "bold": True,
        "italic": False,
        "underline": None,
        "strike": False,
        "color": {"type": "theme", "theme": 1, "resolvedRgb": "#000000"},
    }
    assert cells["I1"]["style"]["fill"] == {
        "type": "solid",
        "fgColor": {
            "type": "rgb",
            "rgb": "#006600",
            "resolvedRgb": "#006600",
        },
        "bgColor": {
            "type": "indexed",
            "indexed": 64,
            "resolvedRgb": None,
        },
    }
    assert result["sheet"] == "Sheet1"
    assert result["range"] == "A1:I3"
    assert result["cellCount"] == 27


def test_read_styles_extracts_alignment_border_and_protection(
    tmp_path: Path,
) -> None:
    path = tmp_path / "styled.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["B2"] = "Manual override"
    sheet["B2"].font = Font(name="Aptos", size=11, italic=True, color="FF112233")
    sheet["B2"].fill = PatternFill("solid", fgColor="FFD9EAF7")
    sheet["B2"].alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
        shrink_to_fit=True,
    )
    sheet["B2"].border = Border(bottom=Side(style="thin", color="FF445566"))
    sheet["B2"].protection = Protection(locked=False, hidden=True)
    workbook.save(path)

    result = read_styles(path, "Report", "B2:B2")

    assert result == {
        "sheet": "Report",
        "range": "B2:B2",
        "rowCount": 1,
        "columnCount": 1,
        "cellCount": 1,
        "cells": [
            {
                "address": "B2",
                "style": {
                    "hasStyle": True,
                    "numberFormat": "General",
                    "font": {
                        "name": "Aptos",
                        "size": 11.0,
                        "bold": False,
                        "italic": True,
                        "underline": None,
                        "strike": False,
                        "color": {
                            "type": "rgb",
                            "rgb": "#112233",
                            "resolvedRgb": "#112233",
                        },
                    },
                    "fill": {
                        "type": "solid",
                        "fgColor": {
                            "type": "rgb",
                            "rgb": "#D9EAF7",
                            "resolvedRgb": "#D9EAF7",
                        },
                        "bgColor": None,
                    },
                    "alignment": {
                        "horizontal": "center",
                        "vertical": "center",
                        "wrapText": True,
                        "shrinkToFit": True,
                        "textRotation": 0,
                    },
                    "border": {
                        "left": {"style": None, "color": None},
                        "right": {"style": None, "color": None},
                        "top": {"style": None, "color": None},
                        "bottom": {
                            "style": "thin",
                            "color": {
                                "type": "rgb",
                                "rgb": "#445566",
                                "resolvedRgb": "#445566",
                            },
                        },
                    },
                    "protection": {"locked": False, "hidden": True},
                },
            }
        ],
    }


def test_read_styles_resolves_theme_and_indexed_colors(tmp_path: Path) -> None:
    theme_result = read_styles(
        UPSTREAM_COMPLEX_STYLES_FIXTURE,
        "Sheet1",
        "A8:A8",
    )

    theme_fill = theme_result["cells"][0]["style"]["fill"]["fgColor"]
    assert theme_fill == {
        "type": "theme",
        "theme": 8,
        "resolvedRgb": "#4BACC6",
    }

    path = tmp_path / "indexed-color.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Indexed"
    sheet["A1"].font = Font(color=Color(indexed=3))
    workbook.save(path)

    indexed_result = read_styles(path, "Sheet", "A1:A1")

    indexed_color = indexed_result["cells"][0]["style"]["font"]["color"]
    assert indexed_color == {
        "type": "indexed",
        "indexed": 3,
        "resolvedRgb": "#00FF00",
    }


def test_read_styles_resolves_theme_tint_in_hls_luminance(tmp_path: Path) -> None:
    path = tmp_path / "theme-tint.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Tinted"
    sheet["A1"].fill = PatternFill(
        "solid",
        fgColor=Color(theme=8, tint=0.4),
    )
    workbook.save(path)

    result = read_styles(path, "Sheet", "A1:A1")

    assert result["cells"][0]["style"]["fill"]["fgColor"] == {
        "type": "theme",
        "theme": 8,
        "tint": 0.4,
        "resolvedRgb": "#93CDDD",
    }


def test_read_styles_rejects_ranges_above_max_cells(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    workbook = Workbook()
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error_info:
        read_styles(path, "Sheet", "A1:C3", max_cells=8)

    assert error_info.value.error_code == "range_too_large"
    assert error_info.value.details == {
        "range": "A1:C3",
        "cellCount": 9,
        "maxCells": 8,
    }
