from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

from zalos_sheetops.benchmarks.public.grading import grade_workbook


def test_correct_answer_with_unrequested_edit_fails(tmp_path: Path) -> None:
    initial, golden, output = [
        tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")
    ]
    book = Workbook()
    book.active["A1"] = "protected text"
    book.active["B1"] = 2
    book.save(initial)
    book.active["B1"] = 4
    book.save(golden)
    book.active["A1"] = "silently damaged"
    book.save(output)
    result = grade_workbook(initial, golden, output, {"answer_position": "B1"})
    assert result["status"] == "incorrect"
    assert result["preservationFailures"][0]["cell"] == "A1"


def test_missing_output_stays_a_failed_attempt(tmp_path: Path) -> None:
    result = grade_workbook(
        tmp_path / "input.xlsx", tmp_path / "golden.xlsx", tmp_path / "missing.xlsx", {}
    )
    assert result["status"] == "missing_output"


def test_unknown_formula_cache_is_not_scored_as_correct_blank(tmp_path: Path) -> None:
    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.active["A1"] = "=1+2"
    for path in paths:
        book.save(path)
    result = grade_workbook(*paths, {"answer_position": "A1"})
    assert result["status"] == "oracle_cache_missing"


def test_correct_public_answer_passes(tmp_path: Path) -> None:
    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.active["A1"] = 1
    book.save(paths[0])
    book.active["A1"] = 2
    book.save(paths[1])
    book.save(paths[2])
    assert grade_workbook(*paths, {"answer_position": "A1"})["status"] == "pass"
    loaded = load_workbook(paths[0])
    assert loaded.active["A1"].value == 1


def test_correct_values_without_required_structure_fail(tmp_path: Path) -> None:
    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.active["A1"] = 1
    book.save(paths[0])
    book.defined_names.add(DefinedName("Required", attr_text="'Sheet'!$A$1"))
    book.active.sheet_view.showGridLines = False
    book.active["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    book.save(paths[1])
    book.defined_names.clear()
    book.active.sheet_view.showGridLines = True
    book.active["A1"].fill = PatternFill()
    book.save(paths[2])
    result = grade_workbook(
        *paths,
        {"answer_position": "A1", "instruction": "Fill A1 yellow and hide gridlines"},
    )
    assert result["status"] == "incorrect"
    assert {d["kind"] for d in result["structureDifferences"]} >= {
        "defined_names",
        "sheet_view",
    }
    assert any(d.get("kind") == "required_style" for d in result["valueDifferences"])


def test_corrupt_output_is_a_failed_attempt(tmp_path: Path) -> None:
    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.save(paths[0])
    book.save(paths[1])
    paths[2].write_text("not an XLSX")
    assert (
        grade_workbook(*paths, {"answer_position": "A1"})["status"] == "invalid_output"
    )


def test_evaluated_empty_string_is_not_a_missing_formula_cache(tmp_path: Path) -> None:
    from zipfile import ZipFile

    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    for path in paths:
        book = Workbook()
        book.active["A1"] = '=IF(TRUE,"",1)'
        book.save(path)
        # LibreOffice marks an evaluated empty-string cache as t="str".
        with ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        member = "xl/worksheets/sheet1.xml"
        members[member] = members[member].replace(b'<c r="A1">', b'<c r="A1" t="str">')
        with ZipFile(path, "w") as archive:
            for name, data in members.items():
                archive.writestr(name, data)
    assert grade_workbook(*paths, {"answer_position": "A1"})["status"] == "pass"
    raw = tmp_path / "raw.xlsx"
    book = Workbook()
    book.active["A1"] = '=IF(TRUE,"",1)'
    book.save(raw)
    assert (
        grade_workbook(
            raw, raw, raw, {"answer_position": "A1"}, evaluated=(paths[1], paths[2])
        )["status"]
        == "pass"
    )


def test_answer_region_ignores_incidental_golden_reference_edits(
    tmp_path: Path,
) -> None:
    from openpyxl.styles import Font

    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.active["C1"] = "unrelated reference"
    book.save(paths[0])
    book.active["B1"] = "answer"
    book.active["B1"].font = Font(name="Arial")
    book.active["C1"] = None
    book.save(paths[1])
    book.active["C1"] = "unrelated reference"
    book.active["B1"].font = Font(name="Calibri")
    book.save(paths[2])
    assert grade_workbook(*paths, {"answer_position": "B1"})["status"] == "pass"


def test_unchanged_array_formulas_must_be_preserved(tmp_path: Path) -> None:
    from openpyxl.worksheet.formula import ArrayFormula

    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.active["A1"] = ArrayFormula(ref="A1", text="=SUM(1,2)")
    book.save(paths[0])
    book.active["B1"] = 2
    book.save(paths[1])
    book.active["A1"] = "=42"
    book.save(paths[2])
    result = grade_workbook(*paths, {"answer_position": "B1"})
    assert result["status"] == "incorrect"
    assert result["preservationFailures"] == [{"sheet": "Sheet", "cell": "A1"}]


def test_formula_request_cannot_pass_with_only_computed_values(tmp_path: Path) -> None:
    initial, golden, actual, cached = [
        tmp_path / name
        for name in ("input.xlsx", "golden.xlsx", "output.xlsx", "cached.xlsx")
    ]
    book = Workbook()
    book.active["A1"] = 0
    book.save(initial)
    book.active["B1"] = "=1+2"
    book.save(golden)
    book.active["B1"] = 3
    book.save(actual)
    book.save(cached)
    result = grade_workbook(
        initial,
        golden,
        actual,
        {"answer_position": "B1", "instruction": "Write a formula in B1"},
        evaluated=(cached, actual),
    )
    assert result["status"] == "incorrect"
    assert any(d.get("kind") == "missing_formula" for d in result["valueDifferences"])


def test_checks_count_failures_beyond_diagnostic_sample_limit(tmp_path: Path) -> None:
    from openpyxl.styles import PatternFill

    initial, golden, actual = [
        tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")
    ]
    book = Workbook()
    for row in range(1, 51):
        book.active.cell(row, 1).value = row
    book.save(initial)
    book.save(actual)
    for row in range(1, 51):
        book.active.cell(row, 1).fill = PatternFill("solid", fgColor="00FF00")
    book.save(golden)
    book = load_workbook(actual)
    book.active["A50"] = 999
    book.save(actual)
    result = grade_workbook(
        initial,
        golden,
        actual,
        {"answer_position": "A1:A50", "instruction": "Fill all answer cells green"},
    )
    assert result["checks"]["answerValues"] is False
    assert result["failureCounts"]["values"] == 1
    assert result["failureCounts"]["styles"] == 50
    assert len(result["valueDifferences"]) <= 20


def test_reviewed_preservation_request_does_not_require_incidental_golden_style(
    tmp_path: Path,
) -> None:
    from openpyxl.styles import Font

    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.active["A1"] = 1
    book.save(paths[0])
    book.active["B1"] = 2
    book.save(paths[2])
    book.active["B1"].font = Font(name="Arial")
    book.save(paths[1])
    task = {
        "answer_position": "B1",
        "instruction": "Preserve existing data and formatting",
        "styleComponents": [],
    }
    assert grade_workbook(*paths, task)["status"] == "pass"


def test_reviewed_style_uses_requested_components_and_visible_theme_color(
    tmp_path: Path,
) -> None:
    from openpyxl.styles import Font
    from openpyxl.styles.colors import Color
    from openpyxl.writer.theme import theme_xml

    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.loaded_theme = theme_xml.encode()
    book.save(paths[0])
    book.active["A1"] = 2
    book.active["A1"].font = Font(bold=True, charset=134)
    book.active["A1"].fill = PatternFill(
        "solid", fgColor=Color(theme=6, tint=0.799981688894314)
    )
    book.save(paths[1])
    book.active["A1"].font = Font(bold=True)
    book.active["A1"].fill = PatternFill("solid", fgColor="EBF1DE")
    book.save(paths[2])
    task = {
        "answer_position": "A1",
        "instruction": "Bold and fill A1 #EBF1DE",
        "styleComponents": ["bold", "fill"],
    }
    assert grade_workbook(*paths, task)["status"] == "pass"
    book.active["A1"].font = Font(bold=False)
    book.save(paths[2])
    assert grade_workbook(*paths, task)["status"] == "incorrect"


def test_excel_serialization_preserves_numbers_but_money_changes_fail(
    tmp_path: Path,
) -> None:
    from zipfile import ZipFile

    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    book.active["A1"] = 0.373369645
    book.active["B1"] = 1000000.01
    book.active["C1"] = 2
    for path in paths:
        book.save(path)
    for path in paths[:2]:
        with ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        members["xl/worksheets/sheet1.xml"] = members[
            "xl/worksheets/sheet1.xml"
        ].replace(b">0.373369645<", b">0.37336964500000003<")
        with ZipFile(path, "w") as archive:
            for name, data in members.items():
                archive.writestr(name, data)
    assert grade_workbook(*paths, {"answer_position": "C1"})["status"] == "pass"
    book.active["B1"] = 1000000.00
    book.save(paths[2])
    assert grade_workbook(*paths, {"answer_position": "C1"})["status"] == "incorrect"


def test_reviewed_oracle_issue_is_unscorable_even_if_output_matches_gold(
    tmp_path: Path,
) -> None:
    paths = [tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")]
    book = Workbook()
    for path in paths:
        book.save(path)
    task = {
        "answer_position": "A1",
        "oracleIssue": {
            "status": "ambiguous_task_metadata",
            "reason": "Unannounced answer sheet",
        },
    }
    assert grade_workbook(*paths, task)["status"] == "ambiguous_task_metadata"


def test_uncalculated_output_formula_cannot_match_a_verified_blank_cache(
    tmp_path: Path,
) -> None:
    from zipfile import ZipFile

    initial, golden, actual = [
        tmp_path / name for name in ("input.xlsx", "golden.xlsx", "output.xlsx")
    ]
    book = Workbook()
    book.save(initial)
    book.active["A1"] = '=IF(TRUE,"",1)'
    book.save(golden)
    book.save(actual)
    with ZipFile(golden) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b'<c r="A1">', b'<c r="A1" t="str">'
    )
    with ZipFile(golden, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    assert (
        grade_workbook(initial, golden, actual, {"answer_position": "A1"})["status"]
        == "output_cache_missing"
    )


def test_table_ids_are_local_but_range_changes_fail_preservation() -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table

    from zalos_sheetops.benchmarks.public.grading import structure_differences

    expected, actual = Workbook(), Workbook()
    for book, identity in [(expected, 2), (actual, 1)]:
        book.active.append(["Name", "Value"])
        book.active.append(["Example", 1])
        book.active.add_table(Table(id=identity, displayName="Example", ref="A1:B2"))
    assert structure_differences(expected, actual) == []
    actual.active.tables["Example"].ref = "A1:B3"
    assert structure_differences(expected, actual) == [
        {"sheet": "Sheet", "kind": "tables"}
    ]
