import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from zalos_sheetops.benchmarks.workflows import runner
from zalos_sheetops.benchmarks.workflows.cases import build_cases
from zalos_sheetops.benchmarks.workflows.grade import grade_case


def test_workflow_oracles_accept_correct_outputs_and_reject_changed_id(
    tmp_path: Path,
) -> None:
    cases = build_cases(tmp_path, detail_rows=100)
    csv_case = next(case for case in cases if case["id"] == "csv_ingestion")
    workspace = tmp_path / "candidate"
    shutil.copytree(tmp_path / "csv_ingestion" / "golden", workspace)
    assert grade_case(csv_case, workspace)["correct"] is True
    target = workspace / "output.csv"
    target.write_text(target.read_text().replace("000000000000000001", "1"))
    assert grade_case(csv_case, workspace)["correct"] is False


def test_formula_case_requires_evaluation_and_preserved_workbook(
    tmp_path: Path,
) -> None:
    cases = build_cases(tmp_path, detail_rows=100)
    case = next(case for case in cases if case["id"] == "targeted_edit")
    golden = tmp_path / "targeted_edit" / "golden"
    grade = grade_case(case, golden)
    assert grade["correct"] is False
    assert grade["formulaEvaluation"] == "unavailable"


def test_report_rejects_replacing_detail_formulas_with_values(tmp_path: Path) -> None:
    cases = build_cases(tmp_path, detail_rows=100)
    case = next(case for case in cases if case["id"] == "formula_report")
    workspace = tmp_path / "report-candidate"
    shutil.copytree(Path(case["golden"]), workspace)
    book = load_workbook(workspace / "output.xlsx")
    book["Data"]["D2"] = 15
    book.save(workspace / "output.xlsx")
    grade = grade_case(case, workspace)
    assert "missing_formula:Data!D2" in grade["failures"]


def test_report_rejects_wrong_labels_extra_rows_and_date_display(
    tmp_path: Path,
) -> None:
    cases = build_cases(tmp_path, detail_rows=100)
    case = next(case for case in cases if case["id"] == "formula_report")
    workspace = tmp_path / "report-candidate"
    shutil.copytree(Path(case["golden"]), workspace)
    book = load_workbook(workspace / "output.xlsx")
    book["Summary"]["A2"] = "Incorrect label"
    book["Data"].append(["extra", None, 0, 0])
    book["Data"]["B2"].number_format = "mm/dd/yyyy"
    book.save(workspace / "output.xlsx")
    failures = grade_case(case, workspace)["failures"]
    assert "data:Summary!A2" in failures
    assert "dimensions:Data" in failures
    assert "date_format:Data!B2" in failures


def test_formula_benchmark_rejects_images_without_calculation_engine(
    monkeypatch,
) -> None:
    def probe(*args, **kwargs) -> str:
        return '{"engine": null}'

    monkeypatch.setattr(runner.subprocess, "check_output", probe)
    with pytest.raises(ValueError, match="calculation engine"):
        runner.probe_images({"baseline": "test-image"})


def test_report_rejects_constant_formulas(tmp_path: Path) -> None:
    cases = build_cases(tmp_path, detail_rows=10)
    case = next(case for case in cases if case["id"] == "formula_report")
    workspace = tmp_path / "candidate"
    shutil.copytree(Path(case["golden"]), workspace)
    book = load_workbook(workspace / "output.xlsx")
    book["Data"]["D2"] = "=15"
    book["Summary"]["B2"] = "=130000"
    book.save(workspace / "output.xlsx")
    failures = grade_case(case, workspace)["failures"]
    assert "formula_relationship:Data!D2" in failures
    assert "formula_relationship:Summary!B2" in failures


def test_edit_rejects_removed_validation_and_layout(tmp_path: Path) -> None:
    cases = build_cases(tmp_path, detail_rows=10)
    case = next(case for case in cases if case["id"] == "targeted_edit")
    workspace = tmp_path / "candidate"
    shutil.copytree(Path(case["golden"]), workspace)
    book = load_workbook(workspace / "output.xlsx")
    book["Inputs"].data_validations.dataValidation.clear()
    book["Inputs"].freeze_panes = None
    book["Inputs"].column_dimensions["D"].width = 50
    book.save(workspace / "output.xlsx")
    assert "sheet_features:Inputs" in grade_case(case, workspace)["failures"]
