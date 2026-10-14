import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def _parse_output(output: str) -> dict[str, object]:
    return json.loads(output)


def _write_workbook(path: Path, value: str = "=A1*2") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 1
    sheet["B1"] = value
    workbook.save(path)


def _mock_detect_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.detect_libreoffice",
        lambda requested: {
            "available": True,
            "executable": sys.executable,
            "source": "explicit",
            "platform": "test",
        },
    )


def test_schema_marks_recalc_available_and_includes_executable() -> None:
    runner = CliRunner()

    result = runner.invoke(app, ["schema", "--raw-output"])

    assert result.exit_code == 0
    payload = _parse_output(result.output)
    commands = {command["name"]: command for command in payload["data"]["commands"]}
    assert commands["recalc"]["status"] == "available"
    assert commands["recalc"]["input"] == {
        "path": "string",
        "outputDir": "string",
        "filename": "string | null",
        "executable": "string | null",
    }


def test_recalc_command_surfaces_preflight_missing_error(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    _write_workbook(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "recalc",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "outputDir": str(tmp_path / "out"),
                    "executable": "/missing/libreoffice",
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = _parse_output(result.output)
    assert payload["ok"] is False
    assert payload["command"] == "recalc"
    assert payload["errorCode"] == "recalc_executable_missing"
    assert payload["details"]["executable"] == "/missing/libreoffice"


def test_recalc_command_returns_error_when_formula_errors_are_found(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    _write_workbook(path)
    runner = CliRunner()

    def fake_run(
        command: list[str],
        check: bool,
        capture_output: bool,
        text: bool,
    ) -> subprocess.CompletedProcess[str]:
        converted_path = Path(command[6]) / Path(command[4]).name
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = 1
        sheet["B1"] = "#VALUE!"
        workbook.save(converted_path)

        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr("zalos_sheetops.recalc.libreoffice.subprocess.run", fake_run)
    _mock_detect_libreoffice(monkeypatch)

    result = runner.invoke(
        app,
        [
            "recalc",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "outputDir": str(tmp_path / "out"),
                    "executable": sys.executable,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = _parse_output(result.output)
    assert payload["ok"] is False
    assert payload["command"] == "recalc"
    assert payload["errorCode"] == "formula_errors_found"
    assert payload["details"]["errorTokenCount"] == 1


def test_recalc_command_returns_success_with_artifact(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    _write_workbook(path)
    runner = CliRunner()

    def fake_run(
        command: list[str],
        check: bool,
        capture_output: bool,
        text: bool,
    ) -> subprocess.CompletedProcess[str]:
        converted_path = Path(command[6]) / Path(command[4]).name
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = 1
        sheet["B1"] = "=A1*2"
        workbook.save(converted_path)

        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr("zalos_sheetops.recalc.libreoffice.subprocess.run", fake_run)
    _mock_detect_libreoffice(monkeypatch)

    result = runner.invoke(
        app,
        [
            "recalc",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "outputDir": str(tmp_path / "out"),
                    "executable": sys.executable,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = _parse_output(result.output)
    assert payload["ok"] is True
    assert payload["command"] == "recalc"
    assert payload["data"]["formulaEvaluation"] == "fresh"
    assert payload["data"]["recalculatedPath"] == str(
        tmp_path / "out" / "source-recalculated.xlsx"
    )
    assert payload["artifacts"] == [
        {
            "kind": "xlsx",
            "path": str(tmp_path / "out" / "source-recalculated.xlsx"),
        }
    ]


def test_recalc_command_rejects_invalid_payload_fields() -> None:
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "recalc",
            "--json",
            json.dumps({"path": 1, "outputDir": "out"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = _parse_output(result.output)
    assert payload == {
        "ok": False,
        "command": "recalc",
        "errorCode": "missing_required_field",
        "message": "Required field is missing or invalid",
        "details": {"field": "path"},
        "warnings": [],
        "artifacts": [],
    }


def test_recalc_command_rejects_missing_path_payload_field() -> None:
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "recalc",
            "--json",
            json.dumps({"outputDir": "out"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert _parse_output(result.output) == {
        "ok": False,
        "command": "recalc",
        "errorCode": "missing_required_field",
        "message": "Required field is missing or invalid",
        "details": {"field": "path"},
        "warnings": [],
        "artifacts": [],
    }


@pytest.mark.parametrize(
    ("payload", "error_code", "message", "details"),
    [
        (
            {"path": "__PATH__"},
            "missing_required_field",
            "Required field is missing or invalid",
            {"field": "outputDir"},
        ),
        (
            {"path": "__PATH__", "outputDir": 1},
            "missing_required_field",
            "Required field is missing or invalid",
            {"field": "outputDir"},
        ),
        (
            {"path": "__PATH__", "outputDir": "out", "filename": 1},
            "invalid_field",
            "Field has invalid type",
            {"field": "filename", "expected": "string | null"},
        ),
        (
            {"path": "__PATH__", "outputDir": "out", "executable": 1},
            "invalid_field",
            "Field has invalid type",
            {"field": "executable", "expected": "string | null"},
        ),
    ],
)
def test_recalc_command_rejects_public_contract_field_errors(
    tmp_path: Path,
    payload: dict[str, object],
    error_code: str,
    message: str,
    details: dict[str, object],
) -> None:
    path = tmp_path / "source.xlsx"
    _write_workbook(path)
    payload["path"] = str(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        ["recalc", "--json", json.dumps(payload), "--raw-output"],
    )

    assert result.exit_code == 1
    assert _parse_output(result.output) == {
        "ok": False,
        "command": "recalc",
        "errorCode": error_code,
        "message": message,
        "details": details,
        "warnings": [],
        "artifacts": [],
    }
