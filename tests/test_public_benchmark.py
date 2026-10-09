import json
from pathlib import Path

from zalos_sheetops.benchmarks.public.scoring import score_outputs, summarize
from zalos_sheetops.benchmarks.public.tasks import select_tasks


def test_wrong_cell_never_receives_speed_ranking() -> None:
    expected = [[True, 1, "001", None]]
    result = score_outputs(expected, [[1, True, 1, ""]])
    assert result["correct"] is False
    assert result["mismatchCount"] == 4
    summary = summarize(
        [
            {"case": "a", "adapter": "bad", "status": "incorrect", "seconds": 0.01},
            {"case": "a", "adapter": "good", "status": "pass", "seconds": 1.0},
            {"case": "a", "adapter": "good", "status": "timeout", "seconds": 2.0},
        ]
    )
    assert summary[0]["medianSeconds"] is None
    assert summary[1]["passRate"] == 0.5
    assert summary[1]["medianSeconds"] is None


def test_score_checks_missing_rows_and_records_first_difference() -> None:
    result = score_outputs([[1], [2]], [[1]])
    assert result["correct"] is False
    assert result["firstDifferences"][0]["row"] == 2
    assert score_outputs([[1.25]], [[1.25]])["correct"] is True


def test_task_selection_is_seeded_and_does_not_mutate_source(tmp_path: Path) -> None:
    tasks = [
        {"id": str(n), "instruction_type": "Cell-Level Manipulation"} for n in range(40)
    ]
    source = tmp_path / "dataset.json"
    source.write_text(json.dumps(tasks), encoding="utf-8")
    first = select_tasks(tasks, count=30, seed=20261009)
    assert first == select_tasks(tasks, count=30, seed=20261009)
    assert len({task["id"] for task in first}) == 30
    assert json.loads(source.read_text(encoding="utf-8")) == tasks


def test_archive_verification_rejects_wrong_public_source_bytes(tmp_path) -> None:
    import hashlib

    import pytest

    from zalos_sheetops.benchmarks.public.fixtures import verify_archive

    archive = tmp_path / "archive.tar.gz"
    archive.write_bytes(b"verified example")
    digest = hashlib.sha256(b"verified example").hexdigest()
    verify_archive(archive, digest)
    archive.write_bytes(b"changed example")
    with pytest.raises(ValueError, match="pinned dataset"):
        verify_archive(archive, digest)


def test_interrupted_download_cannot_be_reused_as_complete_input(
    tmp_path, monkeypatch
) -> None:
    import pytest

    from zalos_sheetops.benchmarks.public.fixtures import download

    class InterruptedSource:
        def __init__(self) -> None:
            self.reads = 0

        def __enter__(self) -> object:
            return self

        def __exit__(self, *args) -> None:
            pass

        def read(self, size: int) -> bytes:
            self.reads += 1
            if self.reads == 1:
                return b"partial workbook"
            raise OSError("simulated transport interruption")

    def open_source(*args, **kwargs) -> InterruptedSource:
        return InterruptedSource()

    monkeypatch.setattr("urllib.request.urlopen", open_source)
    destination = tmp_path / "input.xlsx"
    with pytest.raises(OSError, match="interruption"):
        download("https://example.test/pinned.xlsx", destination)
    assert not destination.exists()
    assert not list(tmp_path.iterdir())


def test_values_only_baseline_reads_workbook_once(tmp_path, monkeypatch) -> None:
    import openpyxl

    from zalos_sheetops.benchmarks.public.adapters import openpyxl_cells

    path = tmp_path / "input.xlsx"
    book = openpyxl.Workbook()
    book.active["A1"] = 12.5
    book.save(path)
    original = openpyxl.load_workbook
    reads = []

    def traced_load(*args, **kwargs) -> object:
        reads.append(kwargs)
        return original(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", traced_load)
    assert openpyxl_cells({"path": str(path), "sheet": "Sheet", "range": "A1"}) == [
        [12.5]
    ]
    assert len(reads) == 1
    assert reads[0]["data_only"] is True


def test_replay_outputs_json_with_oracle_correctness_result(tmp_path: Path) -> None:
    import subprocess
    import sys
    from datetime import datetime

    from openpyxl import Workbook

    root = tmp_path / "data"
    (root / "generated").mkdir(parents=True)
    (root / "excelbench").mkdir()
    (root / "excelbench/manifest.json").write_text('{"files":[]}')
    book = Workbook()
    book.active.title = "data"
    book.active["A1"] = "created"
    book.active["A2"] = datetime(2026, 1, 2)
    book.save(root / "generated/readbench-3.xlsx")
    process = subprocess.run(
        [
            sys.executable,
            "-m",
            "zalos_sheetops.benchmarks.public.runner",
            "replay",
            "--data",
            str(root),
            "--output",
            str(tmp_path / "result"),
            "--sizes",
            "3",
            "--case",
            "3-cells-start",
            "--adapter",
            "sheetops-api",
        ],
        capture_output=True,
        text=True,
    )
    assert process.returncode == 0, process.stderr
    result = json.loads(process.stdout)
    assert result["case"] == "3-cells-start"
    assert result["status"] == "incorrect"
    assert (
        result["sampleComparison"]["firstDifferences"][0]["expected"]
        == "2026-01-02T00:00:00"
    )
    assert len(result["inputSha256"]) == 64
