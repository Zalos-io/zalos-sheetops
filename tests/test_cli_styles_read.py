import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def test_styles_read_command_returns_style_payload(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "Header"
    sheet["A1"].font = Font(bold=True)
    sheet["A1"].fill = PatternFill("solid", fgColor="FFD9EAF7")
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "read",
            "--json",
            json.dumps({"path": str(path), "sheet": "Report", "range": "A1:A1"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["command"] == "styles read"
    assert payload["data"]["cells"][0]["style"]["font"]["bold"] is True
    assert payload["data"]["cells"][0]["style"]["fill"]["fgColor"] == {
        "type": "rgb",
        "rgb": "#D9EAF7",
        "resolvedRgb": "#D9EAF7",
    }
    assert payload["warnings"] == []
    assert payload["artifacts"] == []


def test_styles_read_command_rejects_invalid_range(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "read",
            "--json",
            json.dumps({"path": str(path), "sheet": "Sheet", "range": "A:A"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "styles read",
        "errorCode": "invalid_range",
        "message": "Cell range must be a bounded A1-style range",
        "details": {"range": "A:A"},
        "warnings": [],
        "artifacts": [],
    }


def test_styles_read_command_rejects_invalid_max_cells(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "read",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Sheet",
                    "range": "A1:A1",
                    "maxCells": 0,
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "styles read",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "maxCells", "expected": "positive integer | null"},
        "warnings": [],
        "artifacts": [],
    }
