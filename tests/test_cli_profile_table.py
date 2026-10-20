import hashlib
import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def invoke(payload: dict) -> tuple[int, dict]:
    result = CliRunner().invoke(app, ["profile-table", "--json", json.dumps(payload)])
    return result.exit_code, json.loads(result.output)


def workbook(path: Path, *, epoch_1904: bool = False) -> None:
    book = Workbook()
    if epoch_1904:
        book.epoch = CALENDAR_MAC_1904
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Reference", "Amount", "Date"])
    sheet.append(
        ["00002", "10000.12345678901234567890123456789012345", datetime(2026, 6, 30)]
    )
    sheet.append(["00001", "-10000", datetime(2026, 7, 1)])
    book.save(path)


@pytest.mark.parametrize("epoch_1904", [False, True])
def test_profile_command_reads_and_validates_complete_typed_population(
    tmp_path: Path, epoch_1904: bool
) -> None:
    path = tmp_path / "input.xlsx"
    workbook(path, epoch_1904=epoch_1904)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    code, result = invoke(
        {
            "path": str(path),
            "sheet": "Data",
            "headerStrategy": "row",
            "headerRow": 1,
            "valueMode": "typed",
            "keyColumns": ["reference"],
            "amountColumns": ["amount"],
        }
    )
    assert code == 0 and result["ok"] is True and result["command"] == "profile-table"
    data = result["data"]
    assert data["source"]["rowCount"] == data["source"]["returnedRowCount"] == 2
    assert data["source"]["truncated"] is False
    assert "rows" not in data["source"] and "sourceRows" not in data["source"]
    profile = data["profile"]
    assert profile["validated"] is True
    assert profile["columns"]["reference"] == {
        "first": "00002",
        "last": "00001",
        "min": "00001",
        "max": "00002",
        "blankCount": 0,
    }
    assert profile["totals"]["amount"] == "0.12345678901234567890123456789012345"
    assert (
        datetime.fromisoformat(profile["columns"]["date"]["first"]).date().isoformat()
        == "2026-06-30"
    )
    assert profile["checks"] == {
        "completePopulation": True,
        "rectangularRows": True,
        "nonblankTextKeys": ["reference"],
        "finiteDecimalAmounts": ["amount"],
    }
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_profile_command_csv_preserves_multiline_precision_and_source_order(
    tmp_path: Path,
) -> None:
    path = tmp_path / "input.csv"
    path.write_text(
        'reference;amount;note\n00002;0.12345678901234567890123456789012345;"memo\ncafé"\n00001;1;last\n',
        encoding="utf-16",
    )
    code, result = invoke(
        {
            "path": str(path),
            "encoding": "auto",
            "delimiter": ";",
            "keyColumns": ["reference"],
            "amountColumns": ["amount"],
        }
    )
    assert code == 0
    source, profile = result["data"]["source"], result["data"]["profile"]
    assert source["encoding"] == "utf-16" and source["delimiter"] == ";"
    assert "sourceRows" not in source and "rows" not in source
    assert profile["columns"]["note"]["first"] == "memo\ncafé"
    assert profile["columns"]["reference"]["first"] == "00002"
    assert profile["totals"]["amount"] == "1.12345678901234567890123456789012345"


def test_profile_command_counts_all_invalid_records_including_unmatched(
    tmp_path: Path,
) -> None:
    path = tmp_path / "input.csv"
    path.write_text("reference,amount\nmatched,1\nunmatched,NaN\n,2\nother,Infinity\n")
    code, result = invoke(
        {"path": str(path), "keyColumns": ["reference"], "amountColumns": ["amount"]}
    )
    assert code == 1 and result["ok"] is False
    assert result["errorCode"] == "invalid_table_values"
    assert result["details"]["invalidRecordCount"] == 3
    assert [i["sourceRow"] for i in result["details"]["issues"]] == [3, 4, 5]


def test_profile_command_rejects_unheaded_zero_beyond_preview_window(
    tmp_path: Path,
) -> None:
    path = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["reference", "amount"])
    sheet.append(["0001", "1"])
    sheet.cell(1103, 3, 0)
    book.save(path)
    code, result = invoke(
        {"path": str(path), "sheet": "Data", "headerStrategy": "row", "headerRow": 1}
    )
    assert code == 1 and result["errorCode"] == "table_width_mismatch"
    assert result["details"] == {"sheet": "Data", "row": 1103, "headerColumns": 2}


def test_profile_command_rejects_row_limits_before_reading(tmp_path: Path) -> None:
    code, result = invoke({"path": str(tmp_path / "absent.xlsx"), "maxRows": 1})
    assert code == 1 and result["errorCode"] == "incomplete_population"
    assert result["details"] == {"field": "maxRows"}


@pytest.mark.parametrize(
    "fields,error",
    [
        ({"keyColumns": ["missing"]}, "missing_column"),
        ({"amountColumns": "amount"}, "invalid_field"),
    ],
)
def test_profile_command_preserves_column_validation(
    tmp_path: Path, fields: dict, error: str
) -> None:
    path = tmp_path / "input.csv"
    path.write_text("reference,amount\n0001,1\n")
    code, result = invoke({"path": str(path), **fields})
    assert code == 1 and result["errorCode"] == error


def test_schema_exposes_profile_table_contract() -> None:
    result = CliRunner().invoke(app, ["schema"])
    payload = json.loads(result.output)
    contract = next(
        c for c in payload["data"]["commands"] if c["name"] == "profile-table"
    )
    assert contract["input"]["path"] == "string"


@pytest.mark.parametrize("command", [["profile-table"], ["cells", "read"]])
def test_command_help_exposes_only_its_json_contract(command: list[str]) -> None:
    result = CliRunner().invoke(app, [*command, "--help"])
    assert result.exit_code == 0
    contract = json.loads(result.output.split("JSON contract:\n", 1)[1])
    assert contract["name"] == " ".join(command)
    assert contract["input"]["path"] == "string"


def test_missing_profile_column_lists_available_normalized_headers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "input.csv"
    path.write_text("Reference,Amount\n0001,1\n")
    code, result = invoke(
        {"path": str(path), "keyColumns": ["Reference"], "amountColumns": ["amount"]}
    )
    assert code == 1 and result["errorCode"] == "missing_column"
    assert result["details"] == {
        "columns": ["Reference"],
        "availableColumns": ["reference", "amount"],
    }
