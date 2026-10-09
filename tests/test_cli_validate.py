import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def test_validate_command_returns_nonzero_for_invalid_workbook(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["Date"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "validate",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "requiredSheets": ["Data"],
                    "requiredColumns": {"Data": ["Date", "Amount"]},
                    "uniqueKeys": {},
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["ok"] is False
    assert payload["errorCode"] == "validation_failed"
    assert payload["details"]["issues"][0]["missing"] == ["Amount"]


def test_validate_command_rejects_invalid_required_sheets_type(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "validate",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "requiredSheets": "Data",
                    "requiredColumns": {},
                    "uniqueKeys": {},
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "validate",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "requiredSheets", "expected": "list[string]"},
        "warnings": [],
        "artifacts": [],
    }


def test_validate_command_returns_duplicate_unique_key_issue(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Reference"])
    sheet.append(["2026-06-01", "INV-1"])
    sheet.append(["2026-06-01", "INV-1"])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "validate",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "requiredSheets": ["Data"],
                    "requiredColumns": {"Data": ["Date", "Reference"]},
                    "uniqueKeys": {"Data": ["Date", "Reference"]},
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["ok"] is False
    assert payload["errorCode"] == "validation_failed"
    assert payload["details"]["issues"] == [
        {
            "code": "duplicate_unique_key",
            "sheet": "Data",
            "keyColumns": ["Date", "Reference"],
            "duplicates": [
                {
                    "key": {"Date": "2026-06-01", "Reference": "INV-1"},
                    "rows": [2, 3],
                }
            ],
        }
    ]


def test_validate_command_accepts_mixed_layout_sheet_rules(tmp_path: Path) -> None:
    path = tmp_path / "recon.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Chase Payment Tech Summary"
    summary["B4"] = "=-'Merchant'!L1"
    merchant = workbook.create_sheet("Merchant")
    merchant.append(["Merchant Order Number", "Deposit Amount in USD"])
    merchant.append(["P-001", 10])
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "validate",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheets": [
                        {
                            "name": "Summary",
                            "kind": "summary",
                            "requiredCells": [
                                {
                                    "cell": "A1",
                                    "equals": "Chase Payment Tech Summary",
                                }
                            ],
                            "requiredFormulas": [
                                {"cell": "B4", "contains": "Merchant"}
                            ],
                        },
                        {
                            "name": "Merchant",
                            "kind": "table",
                            "requiredColumns": [
                                "Merchant Order Number",
                                "Deposit Amount in USD",
                            ],
                            "uniqueKeys": ["Merchant Order Number"],
                        },
                    ],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["data"]["sheetsChecked"] == ["Summary", "Merchant"]


def test_validate_command_reports_mixed_layout_formula_failure(
    tmp_path: Path,
) -> None:
    path = tmp_path / "recon.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["B4"] = 25
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "validate",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheets": [
                        {
                            "name": "Summary",
                            "kind": "summary",
                            "requiredFormulas": [
                                {"cell": "B4", "contains": "Merchant"}
                            ],
                        }
                    ],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["ok"] is False
    assert payload["errorCode"] == "validation_failed"
    assert payload["details"]["issues"] == [
        {
            "code": "missing_required_formula",
            "sheet": "Summary",
            "cell": "B4",
            "value": "25",
        }
    ]


@pytest.mark.parametrize(
    ("field_name", "payload"),
    [
        (
            "requiredSheets",
            {
                "path": "__PATH__",
                "requiredColumns": {},
                "uniqueKeys": {},
            },
        ),
        (
            "requiredColumns",
            {
                "path": "__PATH__",
                "requiredSheets": ["Data"],
                "uniqueKeys": {},
            },
        ),
        (
            "uniqueKeys",
            {
                "path": "__PATH__",
                "requiredSheets": ["Data"],
                "requiredColumns": {},
            },
        ),
    ],
)
def test_validate_command_requires_public_contract_fields(
    tmp_path: Path,
    field_name: str,
    payload: dict[str, object],
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.save(path)
    payload["path"] = str(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        ["validate", "--json", json.dumps(payload), "--raw-output"],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "validate",
        "errorCode": "missing_required_field",
        "message": "Required field is missing or invalid",
        "details": {"field": field_name},
        "warnings": [],
        "artifacts": [],
    }


def test_validate_command_rejects_invalid_workbook_content(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("not really an xlsx", encoding="utf-8")
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "validate",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "requiredSheets": ["Data"],
                    "requiredColumns": {"Data": ["Date"]},
                    "uniqueKeys": {},
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "validate",
        "errorCode": "invalid_workbook",
        "message": "Workbook could not be opened",
        "details": {"path": str(path)},
        "warnings": [],
        "artifacts": [],
    }
