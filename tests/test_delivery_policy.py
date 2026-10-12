from pathlib import Path

import pytest
from openpyxl import Workbook

from zalos_sheetops.delivery import DeliveryPolicy


def book_bytes(
    tmp_path: Path, *, changed: bool = False, formula: bool = False
) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 3 if changed else 1
    sheet["B1"] = "=A1+1" if formula else "protected"
    path = tmp_path / "book.xlsx"
    book.save(path)
    return path.read_bytes()


def test_delivery_requires_frozen_scope_and_checks_output_itself(
    tmp_path: Path,
) -> None:
    original = book_bytes(tmp_path)
    changed = book_bytes(tmp_path, changed=True)
    with DeliveryPolicy(original) as policy:
        assert policy.validate(changed)["errorCode"] == "missing_contract"
        policy.declare({"Sheet": ["B1"]})
        assert policy.validate(changed)["errorCode"] == "invalid_output"
        with pytest.raises(ValueError, match="frozen"):
            policy.declare({"Sheet": ["A1:B1"]})


def test_changed_bytes_invalidate_previous_approval(tmp_path: Path) -> None:
    original = book_bytes(tmp_path)
    changed = book_bytes(tmp_path, changed=True)
    with DeliveryPolicy(original) as policy:
        policy.declare({"Sheet": ["B1"]})
        receipt = policy.validate(original)
        assert receipt["action"] == "allow"
        assert policy.is_approved(original)
        assert not policy.is_approved(changed)
        assert policy.validate(changed)["action"] == "block"
        assert not policy.is_approved(original)


def test_formula_workbook_blocks_when_engine_is_missing(tmp_path: Path) -> None:
    original = book_bytes(tmp_path, formula=True)
    with DeliveryPolicy(original) as policy:
        policy.declare({"Sheet": ["A1"]})
        result = policy.validate(original)
        assert result["action"] == "block"
        assert result["errorCode"] == "calculation_engine_unavailable"
        assert not policy.is_approved(original)


def test_engine_refreshes_delivered_caches_and_preserves_other_zip_parts(
    tmp_path: Path,
) -> None:
    from io import BytesIO
    from zipfile import ZipFile

    from openpyxl import load_workbook

    engine = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if not engine.is_file():
        pytest.skip("LibreOffice integration requires an installed engine")
    original = book_bytes(tmp_path, formula=True)
    changed = book_bytes(tmp_path, changed=True, formula=True)
    with DeliveryPolicy(original, engine=str(engine)) as policy:
        policy.declare({"Sheet": ["A1"]})
        assert policy.validate(changed)["action"] == "allow"
        delivered = policy.validated_bytes
        assert policy.is_approved(delivered)
        assert not policy.is_approved(changed)
        cached = load_workbook(BytesIO(delivered), data_only=True)
        assert cached.active["B1"].value == 4
        cached.close()
        with ZipFile(BytesIO(changed)) as before, ZipFile(BytesIO(delivered)) as after:
            for name in before.namelist():
                if not name.startswith("xl/worksheets/"):
                    assert before.read(name) == after.read(name)
