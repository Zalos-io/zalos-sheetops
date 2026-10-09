from pathlib import Path

import pytest
from openpyxl import Workbook

from zalos_sheetops.compare.workbooks import (
    compare_labeled_cells,
    compare_workbook_structure,
)
from zalos_sheetops.errors import SheetOpsError


def test_compare_workbook_structure_reports_added_sheets_and_dimensions(
    tmp_path: Path,
) -> None:
    expected = tmp_path / "expected.xlsx"
    actual = tmp_path / "actual.xlsx"
    _write_structure_workbook(expected, include_controls=False, merchant_rows=3)
    _write_structure_workbook(actual, include_controls=True, merchant_rows=4)

    result = compare_workbook_structure(expected, actual)

    assert result["same"] is False
    assert result["summary"] == {"added": 1, "removed": 0, "changed": 1}
    assert result["addedSheets"] == [
        {
            "name": "Input Controls",
            "position": 3,
            "dimensions": "A1:A1",
            "hidden": False,
            "formulaCount": 0,
        }
    ]
    assert result["removedSheets"] == []
    assert result["changedSheets"] == [
        {
            "name": "Merchant",
            "changes": [
                {"field": "dimensions", "expected": "A1:B3", "actual": "A1:B4"},
                {"field": "maxRow", "expected": 3, "actual": 4},
            ],
        }
    ]


def test_compare_labeled_cells_reports_values_formulas_and_missing_labels(
    tmp_path: Path,
) -> None:
    expected = tmp_path / "expected.xlsx"
    actual = tmp_path / "actual.xlsx"
    _write_summary_workbook(
        expected,
        fx_formula="=-Merchant!L1",
        status="OK",
        include_bridge=True,
    )
    _write_summary_workbook(
        actual,
        fx_formula=(
            '=SUMIFS(Merchant!L:L,Merchant!J:J,"<>In-transit",Merchant!J:J,"<>One-off")'
        ),
        status="CHECK",
        include_bridge=False,
    )

    result = compare_labeled_cells(
        expected,
        actual,
        sheet_name="Summary",
        label_column="A",
        value_column="B",
        labels=["FX Difference", "Status", "Bridge Check"],
    )

    assert result["same"] is False
    assert result["summary"] == {"matched": 0, "changed": 2, "missing": 1}
    comparisons = {item["label"]: item for item in result["comparisons"]}
    assert comparisons["FX Difference"] == {
        "label": "FX Difference",
        "status": "changed",
        "expected": {
            "cell": "B2",
            "value": None,
            "formula": "=-Merchant!L1",
        },
        "actual": {
            "cell": "B2",
            "value": None,
            "formula": (
                '=SUMIFS(Merchant!L:L,Merchant!J:J,"<>In-transit",'
                'Merchant!J:J,"<>One-off")'
            ),
        },
        "differences": [
            {
                "field": "formula",
                "expected": "=-Merchant!L1",
                "actual": (
                    '=SUMIFS(Merchant!L:L,Merchant!J:J,"<>In-transit",'
                    'Merchant!J:J,"<>One-off")'
                ),
            }
        ],
    }
    assert comparisons["Status"]["status"] == "changed"
    assert comparisons["Bridge Check"] == {
        "label": "Bridge Check",
        "status": "missing",
        "expected": {"cell": "B4", "value": 0, "formula": None},
        "actual": None,
        "differences": [{"field": "actual", "expected": "present"}],
    }


def test_compare_labeled_cells_rejects_missing_sheet(tmp_path: Path) -> None:
    expected = tmp_path / "expected.xlsx"
    actual = tmp_path / "actual.xlsx"
    Workbook().save(expected)
    Workbook().save(actual)

    with pytest.raises(SheetOpsError) as error_info:
        compare_labeled_cells(
            expected,
            actual,
            sheet_name="Summary",
            label_column="A",
            value_column="B",
            labels=["Status"],
        )

    assert error_info.value.error_code == "invalid_sheet"
    assert error_info.value.details == {"sheet": "Summary"}


def test_compare_labeled_cells_ignores_duplicate_unrequested_labels(
    tmp_path: Path,
) -> None:
    expected = tmp_path / "expected.xlsx"
    actual = tmp_path / "actual.xlsx"
    _write_summary_workbook(
        expected,
        fx_formula="=1+1",
        status="OK",
        include_bridge=False,
    )
    _write_summary_workbook(
        actual,
        fx_formula="=1+1",
        status="OK",
        include_bridge=False,
    )
    for path in [expected, actual]:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary["A1"] = "Account"
        summary["B1"] = "1000"
        summary["A2"] = "Status"
        summary["B2"] = "OK"
        summary["A3"] = "Account"
        summary["B3"] = "2000"
        workbook.save(path)

    result = compare_labeled_cells(
        expected,
        actual,
        sheet_name="Summary",
        label_column="A",
        value_column="B",
        labels=["Status"],
    )

    assert result["same"] is True
    assert result["summary"] == {"matched": 1, "changed": 0, "missing": 0}


def _write_structure_workbook(
    path: Path,
    *,
    include_controls: bool,
    merchant_rows: int,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Status"
    merchant = workbook.create_sheet("Merchant")
    merchant["A1"] = "Reference"
    merchant["B1"] = "Amount"
    for row in range(2, merchant_rows + 1):
        merchant.cell(row=row, column=1).value = f"P-{row}"
        merchant.cell(row=row, column=2).value = row
    if include_controls:
        controls = workbook.create_sheet("Input Controls")
        controls["A1"] = "Controls"
    workbook.save(path)


def _write_summary_workbook(
    path: Path,
    *,
    fx_formula: str,
    status: str,
    include_bridge: bool,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Metric"
    summary["B1"] = "Value"
    summary["A2"] = "FX Difference"
    summary["B2"] = fx_formula
    summary["A3"] = "Status"
    summary["B3"] = status
    if include_bridge:
        summary["A4"] = "Bridge Check"
        summary["B4"] = 0
    workbook.save(path)
