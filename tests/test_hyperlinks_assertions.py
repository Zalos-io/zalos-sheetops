from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.hyperlink import Hyperlink

from zalos_sheetops.hyperlinks.assertions import assert_hyperlinks

UPSTREAM_HYPERLINK_FIXTURE = Path("tests/fixtures/upstream/xlsxwriter/hyperlink01.xlsx")


def _write_hyperlink_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Links"
    sheet["B2"] = "Customer portal"
    sheet["B2"].hyperlink = Hyperlink(
        ref="B2",
        target="https://example.com/customer?id=42",
        tooltip="Open customer portal",
        display="Customer portal",
    )
    sheet["C3"] = "Summary"
    sheet["C3"].hyperlink = "#Summary!A1"
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Total"
    workbook.save(path)


def test_assert_hyperlinks_uses_xml_for_target_display_and_tooltip(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "links.xlsx"
    _write_hyperlink_workbook(path)

    def fail_if_openpyxl_loads_workbook(*args: object, **kwargs: object) -> None:
        raise AssertionError("hyperlinks assert should not load the full workbook")

    monkeypatch.setattr(
        "openpyxl.reader.excel.load_workbook",
        fail_if_openpyxl_loads_workbook,
    )

    result = assert_hyperlinks(
        path,
        "Links",
        [
            {
                "cell": "B2",
                "target": "https://example.com/customer?id=42",
                "display": "Customer portal",
                "tooltip": "Open customer portal",
            },
            {"cell": "C3", "target": "#Summary!A1"},
        ],
    )

    assert result == {
        "valid": True,
        "sheet": "Links",
        "checkedCount": 2,
        "failureCount": 0,
        "failures": [],
    }


def test_assert_hyperlinks_reports_missing_and_mismatched_fields(
    tmp_path: Path,
) -> None:
    path = tmp_path / "links.xlsx"
    _write_hyperlink_workbook(path)

    result = assert_hyperlinks(
        path,
        "Links",
        [
            {"cell": "A1", "target": "https://example.com/missing"},
            {"cell": "B2", "target": "https://example.com/wrong"},
            {"cell": "B2", "display": "Wrong label"},
            {"cell": "B2", "tooltip": "Wrong tooltip"},
        ],
    )

    assert result == {
        "valid": False,
        "sheet": "Links",
        "checkedCount": 4,
        "failureCount": 4,
        "failures": [
            {
                "cell": "A1",
                "code": "missing_hyperlink",
                "expected": "https://example.com/missing",
                "actual": None,
            },
            {
                "cell": "B2",
                "code": "target_mismatch",
                "expected": "https://example.com/wrong",
                "actual": "https://example.com/customer?id=42",
            },
            {
                "cell": "B2",
                "code": "display_mismatch",
                "expected": "Wrong label",
                "actual": "Customer portal",
            },
            {
                "cell": "B2",
                "code": "tooltip_mismatch",
                "expected": "Wrong tooltip",
                "actual": "Open customer portal",
            },
        ],
    }


def test_assert_hyperlinks_accepts_upstream_xlsxwriter_fixture() -> None:
    result = assert_hyperlinks(
        UPSTREAM_HYPERLINK_FIXTURE,
        "Sheet1",
        [
            {
                "cell": "A1",
                "target": "http://www.perl.org/",
                "display": "http://www.perl.org/",
            }
        ],
    )

    assert result == {
        "valid": True,
        "sheet": "Sheet1",
        "checkedCount": 1,
        "failureCount": 0,
        "failures": [],
    }
