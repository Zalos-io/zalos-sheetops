import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner

UPSTREAM_STYLE_FIXTURE = Path("tests/fixtures/upstream/openpyxl/sample.xlsx")


def _write_style_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Merchant"
    sheet["E2"] = "C Date"
    sheet["E2"].fill = PatternFill("solid", fgColor="D9EAF7")
    sheet["G3"] = "=C3*F3"
    sheet["G3"].fill = PatternFill("solid", fgColor="D9EAF7")
    sheet["H2"] = "MOP Code"
    sheet["H2"].comment = Comment(
        "MOP Code means Method of Payment Code from Chase ACT10.",
        "SheetOps",
    )
    sheet["I2"] = "Manual override"
    sheet["I2"].font = Font(italic=True)
    workbook.save(path)


def test_styles_assert_command_accepts_fill_comment_and_formula_checks(
    tmp_path: Path,
) -> None:
    path = tmp_path / "styled.xlsx"
    _write_style_fixture(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "assert",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Merchant",
                    "checks": [
                        {"cell": "E2", "fill": "#D9EAF7"},
                        {"cell": "G3", "fill": "FFD9EAF7", "formula": True},
                        {"cell": "H2", "commentContains": "Method of Payment Code"},
                    ],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    assert json.loads(result.output) == {
        "ok": True,
        "command": "styles assert",
        "data": {
            "valid": True,
            "sheet": "Merchant",
            "checkedCount": 3,
            "failureCount": 0,
            "failures": [],
        },
        "warnings": [],
        "artifacts": [],
    }


def test_styles_assert_command_reports_fill_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    _write_style_fixture(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "assert",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Merchant",
                    "checks": [{"cell": "E2", "fill": "#FFC7CE"}],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["ok"] is False
    assert payload["command"] == "styles assert"
    assert payload["errorCode"] == "style_assertion_failed"
    assert payload["message"] == "Style assertions failed"
    assert payload["details"] == {
        "valid": False,
        "sheet": "Merchant",
        "checkedCount": 1,
        "failureCount": 1,
        "failures": [
            {
                "cell": "E2",
                "code": "fill_mismatch",
                "expected": "#FFC7CE",
                "actual": "#D9EAF7",
            }
        ],
    }


def test_styles_assert_command_accepts_bold_only_check_from_upstream_fixture() -> None:
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "assert",
            "--json",
            json.dumps(
                {
                    "path": str(UPSTREAM_STYLE_FIXTURE),
                    "sheet": "OXM",
                    "checks": [{"cell": "A1", "bold": True}],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["data"]["failureCount"] == 0


def test_styles_assert_command_accepts_italic_only_check(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    _write_style_fixture(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "assert",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Merchant",
                    "checks": [{"cell": "I2", "italic": True}],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["data"]["failureCount"] == 0


def test_styles_assert_command_reports_italic_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    _write_style_fixture(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "assert",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Merchant",
                    "checks": [{"cell": "E2", "italic": True}],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["details"]["failures"] == [
        {
            "cell": "E2",
            "code": "italic_mismatch",
            "expected": True,
            "actual": False,
        }
    ]


def test_styles_assert_command_reports_missing_comment_text(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    _write_style_fixture(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "assert",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Merchant",
                    "checks": [{"cell": "H2", "commentContains": "not present"}],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["details"]["failures"] == [
        {
            "cell": "H2",
            "code": "comment_mismatch",
            "expected": "not present",
            "actual": "MOP Code means Method of Payment Code from Chase ACT10.",
        }
    ]


def test_styles_assert_command_reports_missing_formula(tmp_path: Path) -> None:
    path = tmp_path / "styled.xlsx"
    _write_style_fixture(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "assert",
            "--json",
            json.dumps(
                {
                    "path": str(path),
                    "sheet": "Merchant",
                    "checks": [{"cell": "E2", "formula": True}],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["details"]["failures"] == [
        {
            "cell": "E2",
            "code": "formula_mismatch",
            "expected": True,
            "actual": False,
        }
    ]


def test_styles_assert_command_rejects_invalid_checks_field(
    tmp_path: Path,
) -> None:
    path = tmp_path / "styled.xlsx"
    _write_style_fixture(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "styles",
            "assert",
            "--json",
            json.dumps({"path": str(path), "sheet": "Merchant", "checks": []}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "styles assert",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "checks", "expected": "non-empty list[styleCheck]"},
        "warnings": [],
        "artifacts": [],
    }
