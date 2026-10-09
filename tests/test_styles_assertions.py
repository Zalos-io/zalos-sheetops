from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.styles.assertions import assert_styles

UPSTREAM_STYLE_FIXTURE = Path("tests/fixtures/upstream/openpyxl/sample.xlsx")


def _write_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Merchant"
    sheet["A1"] = "Added"
    sheet["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    sheet["A1"].comment = Comment("Agent-added column.", "SheetOps")
    sheet["A2"] = "=B2*C2"
    sheet["A3"] = "Manual adjustment"
    sheet["A3"].font = Font(italic=True)
    sheet["F69"] = "=B69*C69"
    sheet["F69"].fill = PatternFill("solid", fgColor="D9EAF7")
    sheet["F69"].comment = Comment(
        "Net CM as a percentage of total invoice charge amount",
        "SheetOps",
    )
    workbook.save(path)


def test_assert_styles_returns_failures_without_raising(tmp_path: Path) -> None:
    path = tmp_path / "style.xlsx"
    _write_workbook(path)

    result = assert_styles(
        path,
        "Merchant",
        [
            {"cell": "A1", "fill": "#FFC7CE"},
            {"cell": "A1", "commentContains": "not there"},
            {"cell": "A1", "formula": True},
        ],
    )

    assert result == {
        "valid": False,
        "sheet": "Merchant",
        "checkedCount": 3,
        "failureCount": 3,
        "failures": [
            {
                "cell": "A1",
                "code": "fill_mismatch",
                "expected": "#FFC7CE",
                "actual": "#D9EAF7",
            },
            {
                "cell": "A1",
                "code": "comment_mismatch",
                "expected": "not there",
                "actual": "Agent-added column.",
            },
            {
                "cell": "A1",
                "code": "formula_mismatch",
                "expected": True,
                "actual": False,
            },
        ],
    }


def test_assert_styles_accepts_argb_color_equivalent(tmp_path: Path) -> None:
    path = tmp_path / "style.xlsx"
    _write_workbook(path)

    result = assert_styles(path, "Merchant", [{"cell": "A1", "fill": "FFD9EAF7"}])

    assert result["valid"] is True
    assert result["failureCount"] == 0


def test_assert_styles_uses_xml_fast_path_for_formula_fill_and_comment(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "style.xlsx"
    _write_workbook(path)

    def fail_if_openpyxl_loads_workbook(*args: object, **kwargs: object) -> None:
        raise AssertionError("styles assert should not load the full workbook")

    monkeypatch.setattr(
        "openpyxl.load_workbook",
        fail_if_openpyxl_loads_workbook,
    )

    result = assert_styles(
        path,
        "Merchant",
        [
            {
                "cell": "F69",
                "formula": True,
                "fill": "#D9EAF7",
                "commentContains": (
                    "Net CM as a percentage of total invoice charge amount"
                ),
            }
        ],
    )

    assert result == {
        "valid": True,
        "sheet": "Merchant",
        "checkedCount": 1,
        "failureCount": 0,
        "failures": [],
    }


def test_assert_styles_accepts_bold_only_check_from_upstream_fixture() -> None:
    result = assert_styles(
        UPSTREAM_STYLE_FIXTURE,
        "OXM",
        [{"cell": "A1", "bold": True}],
    )

    assert result == {
        "valid": True,
        "sheet": "OXM",
        "checkedCount": 1,
        "failureCount": 0,
        "failures": [],
    }


def test_assert_styles_reports_bold_mismatch_from_upstream_fixture() -> None:
    result = assert_styles(
        UPSTREAM_STYLE_FIXTURE,
        "OXM",
        [{"cell": "A1", "bold": False}],
    )

    assert result["valid"] is False
    assert result["failures"] == [
        {
            "cell": "A1",
            "code": "bold_mismatch",
            "expected": False,
            "actual": True,
        }
    ]


def test_assert_styles_accepts_italic_only_check(tmp_path: Path) -> None:
    path = tmp_path / "style.xlsx"
    _write_workbook(path)

    result = assert_styles(path, "Merchant", [{"cell": "A3", "italic": True}])

    assert result == {
        "valid": True,
        "sheet": "Merchant",
        "checkedCount": 1,
        "failureCount": 0,
        "failures": [],
    }


def test_assert_styles_reports_italic_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "style.xlsx"
    _write_workbook(path)

    result = assert_styles(path, "Merchant", [{"cell": "A1", "italic": True}])

    assert result["valid"] is False
    assert result["failures"] == [
        {
            "cell": "A1",
            "code": "italic_mismatch",
            "expected": True,
            "actual": False,
        }
    ]


def test_assert_styles_rejects_invalid_color(tmp_path: Path) -> None:
    path = tmp_path / "style.xlsx"
    _write_workbook(path)

    with pytest.raises(SheetOpsError) as error_info:
        assert_styles(path, "Merchant", [{"cell": "A1", "fill": "blue"}])

    assert error_info.value.error_code == "invalid_field"
    assert error_info.value.details == {
        "field": "checks[0].fill",
        "expected": "hex color string",
    }
