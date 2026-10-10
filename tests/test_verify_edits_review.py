from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Color, Font
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from zalos_sheetops.compare.edits import verify_edits


def test_workbook_theme_changes_cannot_restyle_unrelated_cells(tmp_path: Path) -> None:
    from zipfile import ZIP_DEFLATED, ZipFile

    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    book = Workbook()
    book.active.title = "Data"
    book.active["A1"] = "Keep"
    book.active["A1"].font = Font(color=Color(theme=4))
    book.save(original)
    with ZipFile(original) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    assert b"4F81BD" in entries["xl/theme/theme1.xml"]
    entries["xl/theme/theme1.xml"] = entries["xl/theme/theme1.xml"].replace(
        b"4F81BD", b"FF0000"
    )
    with ZipFile(actual, "w", ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    result = verify_edits(original, actual, {})
    assert not result["verified"], result


def test_conditional_formatting_preserves_resolved_differential_styles(
    tmp_path: Path,
) -> None:
    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    for path, font in [(original, Font(bold=True)), (actual, Font(italic=True))]:
        book = Workbook()
        book.active.title = "Data"
        book.active.conditional_formatting.add(
            "A1:A3", CellIsRule(operator="greaterThan", formula=["0"], font=font)
        )
        book.save(path)
    result = verify_edits(original, actual, {})
    assert not result["verified"], result
    assert "sheet_structure:Data" in result["structureChanges"]


@pytest.mark.parametrize(
    "change", ["local_name", "print_rows", "print_columns", "column_font", "row_font"]
)
def test_structure_preserves_local_names_print_titles_and_dimension_styles(
    tmp_path: Path, change: str
) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.defined_names.add(DefinedName("Local", attr_text="'Data'!$A$1"))
    sheet.print_title_rows = "1:1"
    sheet.print_title_cols = "A:A"
    sheet.column_dimensions["A"].font = Font(bold=True)
    sheet.row_dimensions[1].font = Font(bold=True)
    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    book.save(original)
    assert verify_edits(original, original, {})["verified"]
    # Independent workbooks may reuse the same style index for different fonts.
    if change in {"column_font", "row_font"}:
        book = Workbook()
        sheet = book.active
        sheet.title = "Data"
        sheet.defined_names.add(DefinedName("Local", attr_text="'Data'!$A$1"))
        sheet.print_title_rows = "1:1"
        sheet.print_title_cols = "A:A"
        sheet.column_dimensions["A"].font = Font(
            italic=change == "column_font", bold=change != "column_font"
        )
        sheet.row_dimensions[1].font = Font(
            italic=change == "row_font", bold=change != "row_font"
        )
    elif change == "local_name":
        sheet.defined_names["Local"].attr_text = "'Data'!$B$1"
    elif change == "print_rows":
        sheet.print_title_rows = "2:2"
    else:
        sheet.print_title_cols = "B:B"
    book.save(actual)
    result = verify_edits(original, actual, {})
    assert not result["verified"], result
    assert "sheet_structure:Data" in result["structureChanges"]


def test_array_footprint_and_evaluated_storage_must_match_scope(tmp_path: Path) -> None:
    original, actual, evaluated = [
        tmp_path / name for name in ("original.xlsx", "actual.xlsx", "evaluated.xlsx")
    ]
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    book.save(original)
    sheet["B2"] = ArrayFormula(ref="B2:B2000", text="=ROW(A1:A1999)")
    book.save(actual)
    sheet["B2"] = ArrayFormula(ref="B2", text="=ROW(A1:A1999)")
    book.save(evaluated)
    result = verify_edits(original, actual, {"Data": ["B2"]}, evaluated_path=evaluated)
    reasons = {issue["reason"] for issue in result["formulaIssues"]}
    assert "formula_range_outside_scope" in reasons
    assert "formula_storage_mismatch" in reasons
    # A storage-range-only change must be checked even if its text is unchanged.
    result = verify_edits(evaluated, actual, {"Data": ["B2"]}, evaluated_path=evaluated)
    assert "formula_range_outside_scope" in {
        issue["reason"] for issue in result["formulaIssues"]
    }
    sheet["B2"] = None
    book.save(evaluated)
    result = verify_edits(actual, evaluated, {"Data": ["B2"]})
    assert "formula_range_outside_scope" in {
        issue["reason"] for issue in result["formulaIssues"]
    }


@pytest.mark.parametrize(
    "formula",
    [
        "=IFS(TRUE,1)",
        "=IFNA(A1,0)",
        "=SWITCH(A1,1,2)",
        "=MAXIFS(A1:A3,A1:A3,1)",
        "=_xlfn.FILTER(A1:A3,A1:A3>0)",
        "=_xlws.FILTER(A1:A3,A1:A3>0)",
    ],
)
def test_modern_functions_require_complete_storage_prefix(
    tmp_path: Path, formula: str
) -> None:
    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    book = Workbook()
    book.active.title = "Data"
    book.save(original)
    book.active["B2"] = ArrayFormula(ref="B2", text=formula)
    book.save(actual)
    result = verify_edits(original, actual, {"Data": ["B2"]}, evaluated_path=actual)
    assert "missing_function_prefix" in {
        issue["reason"] for issue in result["formulaIssues"]
    }


def test_existing_data_table_formula_is_preserved_and_changed_table_fails_stably(
    tmp_path: Path,
) -> None:
    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    book = Workbook()
    book.active.title = "Data"
    book.active["B2"] = DataTableFormula(ref="B2:B3", r1="A1")
    book.save(original)
    assert verify_edits(original, original, {})["verified"]
    book.active["B2"] = DataTableFormula(ref="B2:B3", r1="A2")
    book.save(actual)
    result = verify_edits(original, actual, {"Data": ["B2:B3"]})
    assert not result["verified"]
    assert "unsupported_data_table_validation" in {
        issue["reason"] for issue in result["formulaIssues"]
    }


def test_array_cached_errors_beyond_anchor_are_rejected(tmp_path: Path) -> None:
    import xlsxwriter

    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    with xlsxwriter.Workbook(original) as book:
        book.add_worksheet("Data")
    with xlsxwriter.Workbook(actual) as book:
        sheet = book.add_worksheet("Data")
        sheet.write_array_formula("B2:B3", "=A2:A3/0", None, 0)
        sheet.write_number("B3", 0)
    from xml.etree import ElementTree as ET
    from zipfile import ZIP_DEFLATED, ZipFile

    with ZipFile(actual) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cell = root.find(".//s:c[@r='B3']", ns)
    cell.set("t", "e")
    cell.find("s:v", ns).text = "#DIV/0!"
    entries["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with ZipFile(actual, "w", ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    result = verify_edits(original, actual, {"Data": ["B2:B3"]}, evaluated_path=actual)
    assert any(
        issue["cell"] == "B3" and issue["reason"] == "formula_error"
        for issue in result["formulaIssues"]
    )


def test_array_cached_results_cover_entire_declared_output(tmp_path: Path) -> None:
    import xlsxwriter

    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    with xlsxwriter.Workbook(original) as book:
        book.add_worksheet("Data")
    with xlsxwriter.Workbook(actual) as book:
        sheet = book.add_worksheet("Data")
        sheet.write_array_formula("B2:B3", "={1;2}", None, 1)
    from xml.etree import ElementTree as ET
    from zipfile import ZIP_DEFLATED, ZipFile

    with ZipFile(actual) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cell = root.find(".//s:c[@r='B3']", ns)
    cell.remove(cell.find("s:v", ns))
    entries["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with ZipFile(actual, "w", ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    result = verify_edits(original, actual, {"Data": ["B2:B3"]}, evaluated_path=actual)
    assert not result["verified"], result
    assert "missing_cached_result" in {
        issue["reason"] for issue in result["formulaIssues"]
    }
    # Adjacent declared ranges can jointly cover the entire array footprint.
    with xlsxwriter.Workbook(actual) as book:
        sheet = book.add_worksheet("Data")
        sheet.write_array_formula("B2:B3", "={1;2}", None, 1)
        sheet.write_number("B3", 2)
    assert verify_edits(
        original, actual, {"Data": ["B2", "B3"]}, evaluated_path=actual
    )["verified"]
