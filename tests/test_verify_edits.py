import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def verify(original: Path, actual: Path, **options: object) -> dict:
    request = {
        "originalPath": str(original),
        "path": str(actual),
        "allowedRanges": {"Data": ["B2:B3"]},
        **options,
    }
    result = CliRunner().invoke(app, ["verify-edits", "--json", json.dumps(request)])
    return json.loads(result.output)


def test_verify_edits_catches_added_unrelated_cells_and_style_changes(
    tmp_path: Path,
) -> None:
    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet["A1"] = "Keep this header"
    sheet["B2"] = 1
    book.save(original)
    sheet["B2"] = 2
    book.save(actual)
    assert verify(original, actual)["ok"] is True
    sheet["BN1"] = "Unrequested title"
    sheet["A1"].font = Font(bold=True)
    book.save(actual)
    response = verify(original, actual)
    assert response["ok"] is False
    assert response["errorCode"] == "verification_failed"
    assert {
        (item["sheet"], item["cell"])
        for item in response["details"]["unrelatedChanges"]
    } == {("Data", "A1"), ("Data", "BN1")}


def test_verify_edits_requires_evaluated_matching_formulas_and_rejects_errors(
    tmp_path: Path,
) -> None:
    original, actual, evaluated = [
        tmp_path / name for name in ["original.xlsx", "actual.xlsx", "evaluated.xlsx"]
    ]
    book = Workbook()
    book.active.title = "Data"
    book.active["A1"] = 2
    book.save(original)
    book.active["B2"] = "=A1*2"
    book.save(actual)
    missing = verify(original, actual)
    assert missing["ok"] is False
    assert missing["details"]["formulaIssues"][0]["reason"] == "evaluation_required"
    # A freshly serialized formula with no result cannot count as evaluated.
    book.save(evaluated)
    blank = verify(original, actual, evaluatedPath=str(evaluated))
    assert blank["ok"] is False
    assert blank["details"]["formulaIssues"][0]["reason"] == "missing_cached_result"
    from zipfile import ZIP_DEFLATED, ZipFile

    with ZipFile(evaluated) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    from xml.etree import ElementTree as ET

    root = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root.find(".//s:c[@r='B2']/s:v", ns).text = "4"
    entries["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with ZipFile(evaluated, "w", ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    assert verify(original, actual, evaluatedPath=str(evaluated))["ok"] is True
    # A numeric cache alone does not prove it is the same formula.
    changed = load_workbook(evaluated)
    changed["Data"]["B2"] = "=1/0"
    changed.save(evaluated)
    mismatch = verify(original, actual, evaluatedPath=str(evaluated))
    assert mismatch["ok"] is False
    assert mismatch["details"]["formulaIssues"][0]["reason"] == "formula_mismatch"


def test_verify_edits_rejects_unprefixed_modern_formulas_even_with_a_saved_cache(
    tmp_path: Path,
) -> None:
    original, actual = tmp_path / "original.xlsx", tmp_path / "actual.xlsx"
    book = Workbook()
    book.active.title = "Data"
    book.save(original)
    book.active["B2"] = '=TEXTJOIN(",",TRUE,FILTER(A1:A3,A1:A3<>""))'
    book.save(actual)
    response = verify(original, actual, evaluatedPath=str(actual))
    assert response["ok"] is False
    assert "missing_function_prefix" in {
        issue["reason"] for issue in response["details"]["formulaIssues"]
    }


@pytest.mark.parametrize("cache,valid", [("", True), (0, True), ("#DIV/0!", False)])
def test_evaluated_formula_cache_distinguishes_blank_zero_and_error(
    tmp_path: Path,
    cache: object,
    valid: bool,
) -> None:
    import xlsxwriter

    original, actual, evaluated = [
        tmp_path / name for name in ["original.xlsx", "actual.xlsx", "evaluated.xlsx"]
    ]
    with xlsxwriter.Workbook(original) as book:
        book.add_worksheet("Data")
    for path in [actual, evaluated]:
        with xlsxwriter.Workbook(path) as book:
            book.add_worksheet("Data").write_formula(
                "B2", '=IF(TRUE,"",1)', None, cache
            )
        if cache == "":
            # An evaluated empty string has explicit string storage. An
            # untyped empty <v> remains an unavailable cache.
            from zipfile import ZIP_DEFLATED, ZipFile

            with ZipFile(path) as archive:
                entries = {name: archive.read(name) for name in archive.namelist()}
            entries["xl/worksheets/sheet1.xml"] = entries[
                "xl/worksheets/sheet1.xml"
            ].replace(b'<c r="B2">', b'<c r="B2" t="str">')
            with ZipFile(path, "w", ZIP_DEFLATED) as archive:
                for name, content in entries.items():
                    archive.writestr(name, content)
    response = verify(original, actual, evaluatedPath=str(evaluated))
    assert response["ok"] is valid, response
    if not valid:
        assert response["details"]["formulaIssues"][0]["reason"] == "formula_error"


def test_dynamic_function_prefix_alone_does_not_validate_array_storage(
    tmp_path: Path,
) -> None:
    import xlsxwriter

    original, actual, evaluated = [
        tmp_path / name for name in ["original.xlsx", "actual.xlsx", "evaluated.xlsx"]
    ]
    book = Workbook()
    book.active.title = "Data"
    book.save(original)
    book.active["B2"] = '=_xlfn._xlws.FILTER(A1:A3,A1:A3<>"","")'
    book.save(actual)
    with xlsxwriter.Workbook(evaluated) as result:
        result.add_worksheet("Data").write_dynamic_array_formula(
            "B2", '=_xlfn._xlws.FILTER(A1:A3,A1:A3<>"","")', None, 0
        )
    response = verify(original, actual, evaluatedPath=str(evaluated))
    assert response["ok"] is False
    assert "array_storage_required" in {
        issue["reason"] for issue in response["details"]["formulaIssues"]
    }


@pytest.mark.parametrize("unsupported_field", ["originalPath", "path", "evaluatedPath"])
def test_verify_edits_returns_json_for_unsupported_workbook_format(
    tmp_path: Path, unsupported_field: str
) -> None:
    supported = tmp_path / "supported.xlsx"
    unsupported = tmp_path / "legacy.xls"
    book = Workbook()
    book.active.title = "Data"
    book.save(supported)
    unsupported.write_bytes(supported.read_bytes())
    request = {
        "originalPath": str(supported),
        "path": str(supported),
        "evaluatedPath": str(supported),
        "allowedRanges": {},
        unsupported_field: str(unsupported),
    }
    result = CliRunner().invoke(app, ["verify-edits", "--json", json.dumps(request)])
    assert result.exit_code == 1
    response = json.loads(result.output)
    assert response["ok"] is False
    assert response["command"] == "verify-edits"
    assert response["errorCode"] == "invalid_workbook"
    assert "Traceback" not in result.output
