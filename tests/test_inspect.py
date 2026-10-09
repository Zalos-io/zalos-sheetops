from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.inspect import inspect_workbook

CONTENT_TYPES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
    'package.relationships+xml"/>'
    '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
    'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/'
    'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
    "</Types>"
)
ROOT_RELS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
    'relationships"><Relationship Id="rId1" Type="http://schemas.'
    'openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="xl/workbook.xml"/></Relationships>'
)
WORKBOOK_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/'
    'main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
    'relationships"><sheets><sheet name="Data" sheetId="1" r:id="rId1"/>'
    "</sheets></workbook>"
)
WORKBOOK_RELS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
    'relationships"><Relationship Id="rId1" Type="http://schemas.'
    'openxmlformats.org/officeDocument/2006/relationships/worksheet" '
    'Target="worksheets/sheet1.xml"/></Relationships>'
)


def _write_boundary_formula_fixture(path: Path) -> None:
    """Place a formula tag exactly astride the 1MB tag-count chunk boundary.

    The counting scan reads 1MB chunks and holds back a 128-byte tail, so the
    first boundary cut falls at offset 1048448. A ``<f>`` starting two bytes
    earlier cannot match inside the scanned segment and its start is not part
    of the held-back tail.
    """
    header = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
        '2006/main"><dimension ref="A1:B2"/><sheetData>'
    )
    padded_row_prefix = '<row r="1"><c r="A1" t="inlineStr"><is><t>'
    padded_row_suffix = "</t></is></c></row>"
    formula_cell_prefix = '<row r="2"><c r="B2">'
    formula_offset = 1048448 - 2
    padding = formula_offset - (
        len(header)
        + len(padded_row_prefix)
        + len(padded_row_suffix)
        + len(formula_cell_prefix)
    )
    trailing_row = (
        '<row r="3"><c r="A3" t="inlineStr"><is><t>'
        + "y" * 4096
        + "</t></is></c></row>"
    )
    sheet_xml = (
        header
        + padded_row_prefix
        + "x" * padding
        + padded_row_suffix
        + formula_cell_prefix
        + "<f>1+1</f><v>2</v></c></row>"
        + trailing_row
        + "</sheetData></worksheet>"
    )
    encoded = sheet_xml.encode("utf-8")
    assert encoded.index(b"<f>") == formula_offset
    # The sheet must span more than one scan chunk so the formula sits at an
    # interior boundary instead of being rescued by the final tail flush.
    assert len(encoded) > 1048576

    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", CONTENT_TYPES_XML)
        archive.writestr("_rels/.rels", ROOT_RELS_XML)
        archive.writestr("xl/workbook.xml", WORKBOOK_XML)
        archive.writestr("xl/_rels/workbook.xml.rels", WORKBOOK_RELS_XML)
        archive.writestr("xl/worksheets/sheet1.xml", sheet_xml)


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELATIONSHIP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_RELATIONSHIP_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _write_image_workbook(
    path: Path,
    *,
    include_image_relationship: bool = True,
    image_count: int = 1,
    wrap_in_alternate_content: bool = False,
    include_alternate_fallback: bool = False,
) -> None:
    anchors = []
    for index in range(image_count):
        anchors.append(
            f"""
            <xdr:oneCellAnchor>
              <xdr:from>
                <xdr:col>{index + 2}</xdr:col>
                <xdr:colOff>0</xdr:colOff>
                <xdr:row>{index + 4}</xdr:row>
                <xdr:rowOff>0</xdr:rowOff>
              </xdr:from>
              <xdr:ext cx="1905000" cy="952500"/>
              <xdr:pic>
                <xdr:nvPicPr>
                  <xdr:cNvPr id="{index + 2}" name="Image {index + 1}"
                    descr="Evidence {index + 1}"/>
                  <xdr:cNvPicPr/>
                </xdr:nvPicPr>
                <xdr:blipFill>
                  <a:blip r:embed="rIdImage1"/>
                  <a:stretch><a:fillRect/></a:stretch>
                </xdr:blipFill>
                <xdr:spPr/>
              </xdr:pic>
              <xdr:clientData/>
            </xdr:oneCellAnchor>
            """
        )
    image_relationship = (
        """
        <Relationship Id="rIdImage1"
          Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
          Target="../media/image1.png"/>
        """
        if include_image_relationship
        else ""
    )
    drawing_content = "".join(anchors)
    if wrap_in_alternate_content:
        fallback = (
            f"<mc:Fallback>{drawing_content}</mc:Fallback>"
            if include_alternate_fallback
            else ""
        )
        drawing_content = (
            '<mc:AlternateContent><mc:Choice Requires="xdr">'
            f"{drawing_content}"
            f"</mc:Choice>{fallback}</mc:AlternateContent>"
        )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "xl/workbook.xml",
            f"""
            <workbook xmlns="{MAIN_NS}"
              xmlns:r="{RELATIONSHIP_NS}">
              <sheets>
                <sheet name="Evidence" sheetId="1" r:id="rIdSheet1"/>
              </sheets>
            </workbook>
            """,
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            f"""
            <Relationships xmlns="{PACKAGE_RELATIONSHIP_NS}">
              <Relationship Id="rIdSheet1"
                Type="{RELATIONSHIP_NS}/worksheet"
                Target="worksheets/sheet1.xml"/>
            </Relationships>
            """,
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            f"""
            <worksheet xmlns="{MAIN_NS}" xmlns:r="{RELATIONSHIP_NS}">
              <dimension ref="A1"/>
              <sheetData/>
              <drawing r:id="rIdDrawing1"/>
            </worksheet>
            """,
        )
        archive.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            f"""
            <Relationships xmlns="{PACKAGE_RELATIONSHIP_NS}">
              <Relationship Id="rIdDrawing1"
                Type="{RELATIONSHIP_NS}/drawing"
                Target="../drawings/drawing1.xml"/>
            </Relationships>
            """,
        )
        archive.writestr(
            "xl/drawings/drawing1.xml",
            f"""
            <xdr:wsDr xmlns:xdr="{DRAWING_NS}"
              xmlns:a="{DRAWINGML_NS}" xmlns:r="{RELATIONSHIP_NS}"
              xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
              {drawing_content}
            </xdr:wsDr>
            """,
        )
        archive.writestr(
            "xl/drawings/_rels/drawing1.xml.rels",
            f"""
            <Relationships xmlns="{PACKAGE_RELATIONSHIP_NS}">
              {image_relationship}
            </Relationships>
            """,
        )
        archive.writestr("xl/media/image1.png", b"not-decoded-by-inspection")


@pytest.mark.parametrize("dimension", ["", '<dimension ref="invalid"/>'])
@pytest.mark.parametrize("implicit", [False, True])
def test_inspect_derives_missing_dimensions_from_cells(
    tmp_path: Path, dimension: str, implicit: bool
) -> None:
    path = tmp_path / "missing-dimensions.xlsx"
    sheet_xml = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        + dimension
        + '<sheetData><row r="1"><c r="A1" t="inlineStr">'
        "<is><t>Title</t></is></c></row>"
        '<row r="44"><c r="B44" t="inlineStr"><is><t>OK</t></is></c></row>'
        '<row r="45"><c r="E45"><v>716.40</v></c></row>'
        + ('<row><c r="E46"><v>1</v></c><c><v>2</v></c></row>' if implicit else "")
        + "</sheetData></worksheet>"
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", CONTENT_TYPES_XML)
        archive.writestr("_rels/.rels", ROOT_RELS_XML)
        archive.writestr("xl/workbook.xml", WORKBOOK_XML)
        archive.writestr("xl/_rels/workbook.xml.rels", WORKBOOK_RELS_XML)
        archive.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    result = inspect_workbook(path)
    sheet = result["sheets"][0]
    assert sheet["dimensions"] == ("A1:F46" if implicit else "A1:E45")
    assert sheet["maxRow"] == (46 if implicit else 45)
    assert sheet["maxColumn"] == (6 if implicit else 5)


def test_inspect_defaults_to_xml_engine_for_workbook_summary(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "inspection.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Review"
    sheet["A1"] = "Amount"
    sheet["B1"] = "Formula"
    sheet["A2"] = 100
    sheet["B2"] = "=A2*2"
    sheet["C2"] = "Open"
    sheet["C2"].hyperlink = "https://example.com/txn/1"
    sheet.merge_cells("D1:E1")
    workbook.save(path)

    def fail_if_openpyxl_loads_workbook(*args: object, **kwargs: object) -> None:
        raise AssertionError("inspect should default to XML summary")

    monkeypatch.setattr(
        "openpyxl.load_workbook",
        fail_if_openpyxl_loads_workbook,
    )

    result = inspect_workbook(path)

    assert result["workbook"]["engine"] == "xml"
    review = result["sheets"][0]
    assert review["dimensions"] == "A1:E2"
    assert review["mergedRangeCount"] == 1
    assert review["mergedRanges"] == ["D1:E1"]
    assert review["mergedRangesTruncated"] is False
    assert review["hyperlinkCount"] == 1
    assert review["formulaCells"] == ["B2"]
    assert review["state"] == "visible"
    assert review["graphicCount"] == 0
    assert review["graphics"] == []
    assert review["graphicsTruncated"] is False


def test_inspect_xml_engine_reports_structure_without_openpyxl_materialization(
    tmp_path: Path,
) -> None:
    path = tmp_path / "inspection.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Review"
    sheet["A1"] = "Amount"
    sheet["B1"] = "Formula"
    sheet["A2"] = 100
    sheet["B2"] = "=A2*2"
    sheet["C2"] = "Open"
    sheet["C2"].hyperlink = "https://example.com/txn/1"
    sheet.merge_cells("D1:E1")
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Internal"
    workbook.save(path)

    result = inspect_workbook(path, engine="xml")

    assert result["workbook"] == {
        "path": str(path),
        "sheetCount": 2,
        "formulaCount": 1,
        "engine": "xml",
        "definedNameCount": 0,
        "definedNames": [],
        "definedNamesTruncated": False,
    }
    assert result["sheets"] == [
        {
            "name": "Review",
            "maxRow": 2,
            "maxColumn": 5,
            "hidden": False,
            "state": "visible",
            "formulaCount": 1,
            "formulaCells": ["B2"],
            "formulaCellsTruncated": False,
            "dimensions": "A1:E2",
            "mergedRangeCount": 1,
            "mergedRanges": ["D1:E1"],
            "mergedRangesTruncated": False,
            "hyperlinkCount": 1,
            "graphicCount": 0,
            "graphics": [],
            "graphicsTruncated": False,
        },
        {
            "name": "Hidden",
            "maxRow": 1,
            "maxColumn": 1,
            "hidden": True,
            "state": "hidden",
            "formulaCount": 0,
            "formulaCells": [],
            "formulaCellsTruncated": False,
            "dimensions": "A1:A1",
            "mergedRangeCount": 0,
            "mergedRanges": [],
            "mergedRangesTruncated": False,
            "hyperlinkCount": 0,
            "graphicCount": 0,
            "graphics": [],
            "graphicsTruncated": False,
        },
    ]


def test_inspect_xml_reports_defined_names_and_exact_sheet_states() -> None:
    path = Path("tests/fixtures/upstream/openpyxl/hidden_sheets.xlsx")

    result = inspect_workbook(path, engine="xml", max_defined_names=3)

    assert result["workbook"]["definedNameCount"] == 4
    assert result["workbook"]["definedNamesTruncated"] is True
    assert result["workbook"]["definedNames"] == [
        {
            "name": "name",
            "value": '"charlie"',
            "scope": "workbook",
            "hidden": False,
            "builtIn": False,
        },
        {
            "name": "pi",
            "value": "41699",
            "scope": "workbook",
            "hidden": False,
            "builtIn": False,
        },
        {
            "name": "_xlnm.Print_Area",
            "value": "Sheet!$A$1:$D$5,Sheet!$B$9:$F$14",
            "scope": "Sheet",
            "hidden": False,
            "builtIn": True,
        },
    ]
    assert [sheet["state"] for sheet in result["sheets"]] == [
        "visible",
        "hidden",
        "veryHidden",
    ]


def test_inspect_xml_validates_defined_names_past_truncation_boundary(
    tmp_path: Path,
) -> None:
    path = tmp_path / "malformed-truncated-defined-name.xlsx"
    workbook_xml = WORKBOOK_XML.replace(
        "</sheets>",
        (
            "</sheets><definedNames>"
            '<definedName name="Visible">Data!$A$1</definedName>'
            "<definedName>Data!$A$2</definedName>"
            "</definedNames>"
        ),
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", CONTENT_TYPES_XML)
        archive.writestr("_rels/.rels", ROOT_RELS_XML)
        archive.writestr("xl/workbook.xml", workbook_xml)
        archive.writestr("xl/_rels/workbook.xml.rels", WORKBOOK_RELS_XML)
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            f'<worksheet xmlns="{MAIN_NS}"><dimension ref="A1"/><sheetData/>'
            "</worksheet>",
        )

    with pytest.raises(SheetOpsError) as error_info:
        inspect_workbook(path, engine="xml", max_defined_names=1)

    assert error_info.value.error_code == "invalid_workbook"
    assert error_info.value.details == {"path": str(path)}


def test_inspect_xml_reports_chart_anchors_without_guessing_dimensions(
    tmp_path: Path,
) -> None:
    path = tmp_path / "chart.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dashboard"
    sheet.append(["Month", "Revenue"])
    sheet.append(["Jan", 10])
    sheet.append(["Feb", 20])
    chart = BarChart()
    chart.title = "Revenue"
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True
    )
    chart.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=3, row=1),
        to=AnchorMarker(col=7, row=15),
    )
    sheet.add_chart(chart)
    workbook.save(path)

    result = inspect_workbook(path, engine="xml")

    dashboard = result["sheets"][0]
    assert dashboard["graphicCount"] == 1
    assert dashboard["graphicsTruncated"] is False
    assert dashboard["graphics"] == [
        {
            "type": "chart",
            "title": "Revenue",
            "anchor": {
                "type": "twoCell",
                "from": "D2",
                "to": "H16",
                "range": "D2:H16",
            },
        },
    ]


def test_inspect_xml_reports_image_anchors_and_graphics_truncation(
    tmp_path: Path,
) -> None:
    path = tmp_path / "images.xlsx"
    _write_image_workbook(path, image_count=2)

    result = inspect_workbook(path, engine="xml", max_graphics=1)

    evidence = result["sheets"][0]
    assert evidence["graphicCount"] == 2
    assert evidence["graphicsTruncated"] is True
    assert evidence["graphics"] == [
        {
            "type": "image",
            "title": "Evidence 1",
            "anchor": {
                "type": "oneCell",
                "from": "C5",
                "extentEmu": {"width": 1905000, "height": 952500},
            },
        }
    ]


def test_inspect_xml_reads_image_anchor_inside_alternate_content(
    tmp_path: Path,
) -> None:
    path = tmp_path / "alternate-content-image.xlsx"
    _write_image_workbook(path, wrap_in_alternate_content=True)

    result = inspect_workbook(path, engine="xml")

    evidence = result["sheets"][0]
    assert evidence["graphicCount"] == 1
    assert evidence["graphics"][0]["type"] == "image"
    assert evidence["graphics"][0]["anchor"]["from"] == "C5"


def test_inspect_xml_does_not_double_count_alternate_content_fallback(
    tmp_path: Path,
) -> None:
    path = tmp_path / "alternate-content-fallback.xlsx"
    _write_image_workbook(
        path,
        wrap_in_alternate_content=True,
        include_alternate_fallback=True,
    )

    result = inspect_workbook(path, engine="xml")

    evidence = result["sheets"][0]
    assert evidence["graphicCount"] == 1
    assert len(evidence["graphics"]) == 1


def test_inspect_xml_truncates_merged_ranges_without_hiding_count(
    tmp_path: Path,
) -> None:
    path = tmp_path / "many-merged-ranges.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:B1")
    sheet.merge_cells("A2:B2")
    sheet.merge_cells("A3:B3")
    workbook.save(path)

    result = inspect_workbook(path, engine="xml", max_merged_ranges=2)

    summary = result["sheets"][0]
    assert summary["mergedRangeCount"] == 3
    assert summary["mergedRanges"] == ["A3:B3", "A2:B2"]
    assert summary["mergedRangesTruncated"] is True


def test_inspect_xml_rejects_unresolved_image_relationship(tmp_path: Path) -> None:
    path = tmp_path / "broken-image-relationship.xlsx"
    _write_image_workbook(path, include_image_relationship=False)

    with pytest.raises(SheetOpsError) as error_info:
        inspect_workbook(path, engine="xml")

    error = error_info.value
    assert error.error_code == "invalid_workbook"
    assert error.message == "Workbook could not be opened"
    assert error.details == {"path": str(path)}


def test_inspect_xml_engine_truncates_formula_cells_without_hiding_count(
    tmp_path: Path,
) -> None:
    path = tmp_path / "many-formulas.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = 1
    sheet["B1"] = "=A1*2"
    sheet["B2"] = "=A1*3"
    sheet["B3"] = "=A1*4"
    workbook.save(path)

    result = inspect_workbook(path, engine="xml", max_formula_cells=2)

    data = result["sheets"][0]
    assert data["formulaCount"] == 3
    assert data["formulaCells"] == ["B1", "B2"]
    assert data["formulaCellsTruncated"] is True


def test_inspect_default_formula_cell_cap_stays_token_bounded(
    tmp_path: Path,
) -> None:
    path = tmp_path / "many-formulas.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = 1
    for row in range(1, 26):
        sheet[f"B{row}"] = f"=A1*{row}"
    workbook.save(path)

    result = inspect_workbook(path)

    data = result["sheets"][0]
    assert data["formulaCount"] == 25
    assert len(data["formulaCells"]) == 20
    assert data["formulaCellsTruncated"] is True


def test_inspect_counts_formula_tags_that_straddle_scan_chunk_boundaries(
    tmp_path: Path,
) -> None:
    path = tmp_path / "boundary.xlsx"
    _write_boundary_formula_fixture(path)

    result = inspect_workbook(path, engine="xml")

    data = result["sheets"][0]
    assert data["formulaCount"] == 1
    assert data["formulaCells"] == ["B2"]


def test_inspect_rejects_invalid_workbook_content(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("not really an xlsx", encoding="utf-8")

    with pytest.raises(SheetOpsError) as error_info:
        inspect_workbook(path)

    error = error_info.value
    assert error.error_code == "invalid_workbook"
    assert error.message == "Workbook could not be opened"
    assert error.details == {"path": str(path)}
