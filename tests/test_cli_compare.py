import json
from pathlib import Path

from openpyxl import Workbook

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def test_compare_structure_command_returns_difference_error(tmp_path: Path) -> None:
    expected = tmp_path / "expected.xlsx"
    actual = tmp_path / "actual.xlsx"
    _write_structure_workbook(expected, include_controls=False)
    _write_structure_workbook(actual, include_controls=True)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "compare",
            "structure",
            "--json",
            json.dumps(
                {
                    "expected": {"path": str(expected)},
                    "actual": {"path": str(actual)},
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["ok"] is False
    assert payload["command"] == "compare structure"
    assert payload["errorCode"] == "workbook_structure_differences_found"
    assert payload["details"]["summary"] == {"added": 1, "removed": 0, "changed": 0}
    assert payload["details"]["addedSheets"][0]["name"] == "Input Controls"


def test_compare_labeled_cells_command_returns_success_for_matching_cells(
    tmp_path: Path,
) -> None:
    expected = tmp_path / "expected.xlsx"
    actual = tmp_path / "actual.xlsx"
    _write_summary_workbook(expected, status="OK")
    _write_summary_workbook(actual, status="OK")
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "compare",
            "labeled-cells",
            "--json",
            json.dumps(
                {
                    "expected": {"path": str(expected), "sheet": "Summary"},
                    "actual": {"path": str(actual), "sheet": "Summary"},
                    "labelColumn": "A",
                    "valueColumn": "B",
                    "labels": ["Status"],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["command"] == "compare labeled-cells"
    assert payload["data"]["same"] is True
    assert payload["data"]["comparisons"] == [
        {
            "label": "Status",
            "status": "matched",
            "expected": {"cell": "B2", "value": "OK", "formula": None},
            "actual": {"cell": "B2", "value": "OK", "formula": None},
            "differences": [],
        }
    ]


def test_compare_labeled_cells_command_returns_difference_error(
    tmp_path: Path,
) -> None:
    expected = tmp_path / "expected.xlsx"
    actual = tmp_path / "actual.xlsx"
    _write_summary_workbook(expected, status="OK")
    _write_summary_workbook(actual, status="CHECK")
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "compare",
            "labeled-cells",
            "--json",
            json.dumps(
                {
                    "expected": {"path": str(expected), "sheet": "Summary"},
                    "actual": {"path": str(actual), "sheet": "Summary"},
                    "labelColumn": "A",
                    "valueColumn": "B",
                    "labels": ["Status"],
                }
            ),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["ok"] is False
    assert payload["command"] == "compare labeled-cells"
    assert payload["errorCode"] == "labeled_cell_differences_found"
    assert payload["details"]["summary"] == {"matched": 0, "changed": 1, "missing": 0}


def _write_structure_workbook(path: Path, *, include_controls: bool) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Status"
    if include_controls:
        controls = workbook.create_sheet("Input Controls")
        controls["A1"] = "Controls"
    workbook.save(path)


def _write_summary_workbook(path: Path, *, status: str) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Metric"
    summary["B1"] = "Value"
    summary["A2"] = "Status"
    summary["B2"] = status
    workbook.save(path)
