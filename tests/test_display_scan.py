from pathlib import Path

from openpyxl import Workbook

from zalos_sheetops.display.scan import scan_display_overflow


def test_scan_display_overflow_reports_narrow_numeric_and_date_cells(
    tmp_path: Path,
) -> None:
    path = tmp_path / "narrow-values.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.column_dimensions["A"].width = 4
    sheet.column_dimensions["B"].width = 14
    sheet["A1"] = 123456789
    sheet["B1"] = 123456789
    sheet["A2"] = "Long text that should clip instead of hash"
    sheet["A3"] = 44926
    sheet["A3"].number_format = "mm/dd/yyyy"
    workbook.save(path)

    result = scan_display_overflow(path)

    assert result["candidateCount"] == 2
    assert result["unknownCount"] == 0
    assert result["truncated"] is False
    assert result["candidates"] == [
        {
            "sheet": "Report",
            "cell": "A1",
            "valueType": "number",
            "displayText": "123456789",
            "columnWidthPixels": 33,
            "requiredWidthPixels": 70,
            "issue": "display_overflow_candidate",
            "confidence": "high",
            "reason": "numeric_display_too_narrow",
        },
        {
            "sheet": "Report",
            "cell": "A3",
            "valueType": "date",
            "displayText": "12/31/2022",
            "columnWidthPixels": 33,
            "requiredWidthPixels": 75,
            "issue": "display_overflow_candidate",
            "confidence": "medium",
            "reason": "date_display_too_narrow",
        },
    ]


def test_scan_display_overflow_can_scope_sheet_range_and_max_results(
    tmp_path: Path,
) -> None:
    path = tmp_path / "scoped-values.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.column_dimensions["A"].width = 3
    sheet["A1"] = 123456
    sheet["A2"] = 234567
    workbook.save(path)

    result = scan_display_overflow(
        path,
        sheet_name="Report",
        cell_range="A1:A2",
        max_results=1,
    )

    assert result["candidateCount"] == 2
    assert result["returnedCandidateCount"] == 1
    assert result["truncated"] is True
    assert result["scope"] == {"sheet": "Report", "range": "A1:A2"}
    assert [candidate["cell"] for candidate in result["candidates"]] == ["A1"]


def test_scan_display_overflow_honors_max_cells(
    tmp_path: Path,
) -> None:
    path = tmp_path / "bounded-values.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.column_dimensions["A"].width = 3
    for row in range(1, 6):
        sheet.cell(row=row, column=1, value=123456)
    workbook.save(path)

    result = scan_display_overflow(path, max_cells=3)

    assert result["scannedCellCount"] == 3
    assert result["candidateCount"] == 3
    assert result["stoppedEarly"] is True
    assert [candidate["cell"] for candidate in result["candidates"]] == [
        "A1",
        "A2",
        "A3",
    ]


def test_scan_display_overflow_reports_formula_without_cached_value_as_unknown(
    tmp_path: Path,
) -> None:
    path = tmp_path / "formula-no-cache.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "=1+1"
    workbook.save(path)

    result = scan_display_overflow(path)

    assert result["candidateCount"] == 0
    assert result["unknownCount"] == 1
    assert result["unknowns"] == [
        {
            "sheet": "Report",
            "cell": "A1",
            "reason": "formula_without_cached_value",
        }
    ]


def test_scan_display_overflow_accepts_upstream_xlsxwriter_fixture() -> None:
    path = Path("tests/fixtures/upstream/xlsxwriter/autofit01.xlsx")

    result = scan_display_overflow(path, sheet_name="Sheet1")

    assert result["scope"] == {"sheet": "Sheet1", "range": None}
    assert result["scannedCellCount"] > 0
    assert result["candidateCount"] == 0
