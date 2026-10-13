from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from zalos_sheetops.errors import SheetOpsError


def book(tmp_path: Path, rows: list[list[object]]) -> Path:
    path = tmp_path / "bulk.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_bulk_preserves_typed_values_identifiers_and_blank_rows(tmp_path: Path) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    path = book(
        tmp_path,
        [
            ["Account ID", "Amount", "Date", "Active"],
            ["000000000000000001", 12.5, datetime(2026, 6, 30), True],
            [None, None, None, None],
            ["000000000000000002", 0, datetime(2026, 7, 1), False],
        ],
    )
    result = read_bulk_table(path, "Data", 1)
    assert result.column_names == ["account id", "amount", "date", "active"]
    assert result.to_pylist() == [
        {
            "account id": "000000000000000001",
            "amount": 12.5,
            "date": datetime(2026, 6, 30),
            "active": True,
        },
        {"account id": None, "amount": None, "date": None, "active": None},
        {
            "account id": "000000000000000002",
            "amount": 0,
            "date": datetime(2026, 7, 1),
            "active": False,
        },
    ]


@pytest.mark.parametrize(
    "value,code",
    [
        ("#DIV/0!", "bulk_cell_error"),
        ("=1+1", "bulk_formulas_unsupported"),
        (9007199254740993, "bulk_precision_risk"),
        (-(2**63), "bulk_precision_risk"),
    ],
)
def test_bulk_rejects_lossy_values(tmp_path: Path, value: object, code: str) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    path = book(tmp_path, [["id", "amount"], ["001", value]])
    with pytest.raises(SheetOpsError) as exc:
        read_bulk_table(path, "Data", 1)
    assert exc.value.error_code == code


def test_string_override_cannot_hide_numeric_precision_loss(tmp_path: Path) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    path = book(tmp_path, [["id", "amount"], [9007199254740993, 1]])
    with pytest.raises(SheetOpsError) as exc:
        read_bulk_table(path, "Data", 1, column_types={"id": "string"})
    assert exc.value.error_code == "bulk_precision_risk"


def test_bulk_rejects_mixed_column_without_explicit_schema(tmp_path: Path) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    path = book(tmp_path, [["id", "amount"], ["001", 12.5], ["002", "pending"]])
    with pytest.raises(SheetOpsError) as exc:
        read_bulk_table(path, "Data", 1)
    assert exc.value.error_code == "bulk_type_conflict"
    table = read_bulk_table(path, "Data", 1, column_types={"amount": "string"})
    assert table.column("amount").to_pylist() == ["12.5", "pending"]


def test_bulk_preserves_explicit_header_position(tmp_path: Path) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    path = book(tmp_path, [[], ["Report"], [], ["id", "amount"], ["001", 12]])
    table = read_bulk_table(path, "Data", 4)
    assert table.to_pylist() == [{"id": "001", "amount": 12}]


@pytest.mark.parametrize(
    "rows,code",
    [
        ([["ID", "id"], [1, 2]], "invalid_headers"),
        ([["id", None, "amount"], [1, 2, 3]], "invalid_headers"),
        ([["id", "amount"], [1, 2, 3]], "bulk_width_mismatch"),
    ],
)
def test_bulk_rejects_ambiguous_structure(
    tmp_path: Path, rows: list, code: str
) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    with pytest.raises(SheetOpsError) as exc:
        read_bulk_table(book(tmp_path, rows), "Data", 1)
    assert exc.value.error_code == code


def test_hybrid_benchmark_uses_bulk_tables_and_exact_xml_cells(tmp_path: Path) -> None:
    from zalos_sheetops.benchmarks.public.adapters import read_case

    path = book(tmp_path, [["id", "amount"], ["001", 12.5]])
    case = {"path": str(path), "sheet": "Data", "operation": "table"}
    assert read_case(case, "sheetops-hybrid") == [["id", "amount"], ["001", "12.5"]]
    assert read_case(
        {**case, "operation": "cells", "range": "A2:B2"}, "sheetops-hybrid"
    ) == [["001", 12.5]]


def test_bulk_checks_types_beyond_initial_sample(tmp_path: Path) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    path = book(
        tmp_path,
        [["id", "amount"], *[[str(i), i] for i in range(1100)], ["late", "pending"]],
    )
    with pytest.raises(SheetOpsError) as exc:
        read_bulk_table(path, "Data", 1)
    assert exc.value.error_code == "bulk_type_conflict"


def test_bulk_1904_epoch_and_other_sheet_formulas(tmp_path: Path) -> None:
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    from zalos_sheetops.io.read_bulk import read_bulk_table

    path = tmp_path / "mac.xlsx"
    workbook = Workbook()
    workbook.epoch = CALENDAR_MAC_1904
    workbook.active.title = "Data"
    workbook.active.append(["date", "id"])
    workbook.active.append([datetime(2026, 6, 30, 12, 30), "001"])
    workbook.create_sheet("Summary")["A1"] = "=1+1"
    workbook.save(path)
    assert read_bulk_table(path, "Data", 1).to_pylist() == [
        {"date": datetime(2026, 6, 30, 12, 30), "id": "001"}
    ]


@pytest.mark.parametrize("schema", [{"missing": "string"}, {"id": "int"}, []])
def test_bulk_rejects_invalid_schema(tmp_path: Path, schema: object) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    with pytest.raises(SheetOpsError) as exc:
        read_bulk_table(
            book(tmp_path, [["id"], ["001"]]), "Data", 1, column_types=schema
        )
    assert exc.value.error_code == "invalid_bulk_schema"


def test_bulk_header_only_table_is_empty_with_named_columns(tmp_path: Path) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    table = read_bulk_table(book(tmp_path, [["id", "amount"]]), "Data", 1)
    assert table.column_names == ["id", "amount"]
    assert table.num_rows == 0


def test_bulk_rejects_literal_int64_minimum(tmp_path: Path) -> None:
    from zalos_sheetops.io.read_bulk import read_bulk_table

    path = book(tmp_path, [["id", "amount"], ["001", -123]])
    with ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
        b"<v>-123</v>", b"<v>-9223372036854775808</v>"
    )
    with ZipFile(path, "w") as target:
        for name, data in entries.items():
            target.writestr(name, data)
    with pytest.raises(SheetOpsError) as exc:
        read_bulk_table(path, "Data", 1)
    assert exc.value.error_code == "bulk_precision_risk"


def test_bulk_precision_guard_handles_arrow_signed_minimum(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import fastexcel
    import pyarrow as pa

    from zalos_sheetops.io.read_bulk import read_bulk_table

    # fastexcel currently infers floats; its supported Arrow integer output
    # must obey the same guard without signed abs() overflow.
    def integer_batch(self: object) -> tuple[pa.RecordBatch, None]:
        return pa.record_batch(
            [pa.array([-(2**63)], type=pa.int64())], names=["id"]
        ), None

    monkeypatch.setattr(fastexcel.ExcelSheet, "to_arrow_with_errors", integer_batch)
    with pytest.raises(SheetOpsError) as exc:
        read_bulk_table(book(tmp_path, [["id"], [1]]), "Data", 1)
    assert exc.value.error_code == "bulk_precision_risk"
