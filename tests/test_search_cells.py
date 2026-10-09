from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.search_cells import search_cells


def test_search_cells_paginates_matches_in_workbook_order(tmp_path: Path) -> None:
    path = tmp_path / "search.xlsx"
    workbook = Workbook()
    visible = workbook.active
    visible.title = "Visible"
    visible["A1"] = "Revenue North"
    visible["A2"] = "revenue south"
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Revenue Hidden"
    workbook.save(path)

    first_page = search_cells(path, "revenue", fields=["values"], limit=1)
    second_page = search_cells(
        path,
        "revenue",
        fields=["values"],
        limit=1,
        offset=1,
    )

    assert first_page["returnedCount"] == 1
    assert first_page["hasMore"] is True
    assert first_page["nextOffset"] == 1
    assert first_page["stoppedEarly"] is False
    assert first_page["matches"] == [
        {
            "sheet": "Visible",
            "cell": "A1",
            "matchedFields": ["values"],
            "value": "Revenue North",
        }
    ]
    assert second_page["returnedCount"] == 1
    assert second_page["hasMore"] is False
    assert second_page["nextOffset"] is None
    assert second_page["matches"][0]["cell"] == "A2"


def test_search_cells_uses_sparse_cell_coordinate(tmp_path: Path) -> None:
    path = tmp_path / "sparse-search.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["C5"] = "Needle"
    workbook.save(path)

    result = search_cells(path, "needle", fields=["values"])

    assert result["matches"][0]["cell"] == "C5"


def test_search_cells_can_search_resolved_style_metadata() -> None:
    path = Path("tests/fixtures/upstream/openpyxl/complex-styles.xlsx")

    result = search_cells(
        path,
        r"^#4BACC6$",
        fields=["fillColors"],
        sheet_name="Sheet1",
        cell_range="A8:G8",
        regex=True,
        limit=2,
    )

    assert result["returnedCount"] == 2
    assert result["hasMore"] is True
    assert [match["cell"] for match in result["matches"]] == ["A8", "C8"]
    assert result["matches"][0]["matchedFields"] == ["fillColors"]
    assert result["matches"][0]["fill"]["fgColor"]["resolvedRgb"] == "#4BACC6"


def test_search_cells_fill_colors_excludes_fill_pattern_names(
    tmp_path: Path,
) -> None:
    path = tmp_path / "fill-pattern.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="4BACC6")
    workbook.save(path)

    result = search_cells(path, "solid", fields=["fillColors"])

    assert result["matches"] == []


def test_search_cells_can_search_formulas_formats_and_cached_values(
    tmp_path: Path,
) -> None:
    path = tmp_path / "formula-search.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 10
    sheet["B1"] = "=A1*2"
    sheet["B1"].number_format = "#,##0.00"
    workbook.save(path)

    formula_result = search_cells(
        path,
        "A1*2",
        fields=["formulas"],
        sheet_name="Sheet",
    )
    format_result = search_cells(
        path,
        "#,##0.00",
        fields=["numberFormats"],
        sheet_name="Sheet",
    )
    cached_result = search_cells(
        Path("tests/fixtures/upstream/xlsxwriter/formula_results01.xlsx"),
        "2",
        fields=["cachedValues"],
        sheet_name="Sheet1",
        cell_range="A1:A1",
    )

    assert formula_result["matches"] == [
        {
            "sheet": "Sheet",
            "cell": "B1",
            "matchedFields": ["formulas"],
            "formula": "=A1*2",
        }
    ]
    assert format_result["matches"][0]["cell"] == "B1"
    assert format_result["matches"][0]["numberFormat"] == "#,##0.00"
    assert cached_result["matches"][0]["cachedValue"] == 2


def test_search_cells_reports_when_scan_bound_stops_the_search(
    tmp_path: Path,
) -> None:
    path = tmp_path / "bounded.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "first"
    sheet["A2"] = "second"
    workbook.save(path)

    result = search_cells(
        path,
        "missing",
        fields=["values"],
        max_cells=1,
    )

    assert result["returnedCount"] == 0
    assert result["scannedCellCount"] == 1
    assert result["stoppedEarly"] is True
    assert result["hasMore"] is None
    assert result["nextOffset"] is None


def test_search_cells_rejects_range_without_sheet(tmp_path: Path) -> None:
    path = tmp_path / "search.xlsx"
    workbook = Workbook()
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error_info:
        search_cells(path, "value", cell_range="A1:A2")

    assert error_info.value.error_code == "invalid_scope"
    assert error_info.value.details == {"field": "range", "requires": "sheet"}


def test_search_cells_rejects_invalid_regex(tmp_path: Path) -> None:
    path = tmp_path / "search.xlsx"
    workbook = Workbook()
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error_info:
        search_cells(path, "(", regex=True)

    assert error_info.value.error_code == "invalid_regex"
    assert (
        error_info.value.message == "Search pattern is not a valid regular expression"
    )
    assert error_info.value.details == {"query": "("}
