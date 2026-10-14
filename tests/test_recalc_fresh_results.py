"""Fresh recalculation must not reuse stale saved formula caches."""

import shutil
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from zalos_sheetops.recalc.libreoffice import recalc_workbook


def test_real_recalc_replaces_stale_cache_and_preserves_source(tmp_path: Path) -> None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        pytest.skip("LibreOffice is not available on this test host")
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.calculation = None
    workbook.active["A1"] = "=1+1"
    workbook.save(source)
    with ZipFile(source) as archive:
        parts = {member: archive.read(member) for member in archive.namelist()}
    # A valid minimal producer package has no application metadata or calc flags.
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/'
            'content-types"><Default Extension="rels" ContentType="application/'
            'vnd.openxmlformats-package.relationships+xml"/><Default '
            'Extension="xml" ContentType="application/xml"/><Override '
            'PartName="/xl/workbook.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>'
        ).encode(),
        "_rels/.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships"><Relationship Id="rId1" Type="http://schemas.'
            'openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="xl/workbook.xml"/></Relationships>'
        ).encode(),
        "xl/workbook.xml": (
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/'
            'main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships"><sheets><sheet name="Data" sheetId="1" '
            'r:id="rId1"/></sheets></workbook>'
        ).encode(),
        "xl/_rels/workbook.xml.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships"><Relationship Id="rId1" Type="http://schemas.'
            'openxmlformats.org/officeDocument/2006/relationships/worksheet" '
            'Target="worksheets/sheet1.xml"/></Relationships>'
        ).encode(),
        "xl/worksheets/sheet1.xml": (
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/'
            'main"><sheetData><row r="1"><c r="A1"><f>1+1</f><v>999</v></c>'
            "</row></sheetData></worksheet>"
        ).encode(),
    }
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    original = source.read_bytes()
    result = recalc_workbook(source, tmp_path / "output", None, executable)
    assert result["formulaEvaluation"] == "fresh"
    evaluated = load_workbook(result["recalculatedPath"], data_only=True)
    assert evaluated.active["A1"].value == 2
    assert source.read_bytes() == original


def test_engine_input_omits_formula_cache_but_preserves_literal_values(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 7
    workbook.active["B1"] = "=A1*2"
    workbook.save(source)
    original = source.read_bytes()
    import subprocess

    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.detect_libreoffice",
        lambda _: {"available": True, "executable": "calc"},
    )

    def engine(command: list[str], **kwargs: object) -> subprocess.CompletedProcess:
        staged = Path(command[4])
        with ZipFile(staged) as archive:
            xml = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        assert xml.find('.//m:c[@r="A1"]/m:v', ns).text == "7"
        assert xml.find('.//m:c[@r="B1"]/m:v', ns) is None
        destination = Path(command[6]) / staged.name
        destination.write_bytes(staged.read_bytes())
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr("zalos_sheetops.recalc.libreoffice.subprocess.run", engine)
    recalc_workbook(source, tmp_path / "output", None)
    assert source.read_bytes() == original
