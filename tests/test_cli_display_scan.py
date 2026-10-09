import json
from pathlib import Path

from openpyxl import Workbook

from zalos_sheetops.cli import app
from zalos_sheetops.testing import CliRunner


def test_display_scan_command_returns_success_when_no_candidates(
    tmp_path: Path,
) -> None:
    path = tmp_path / "wide-values.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.column_dimensions["A"].width = 14
    sheet["A1"] = 123456
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "display",
            "scan",
            "--json",
            json.dumps({"path": str(path), "sheet": "Report"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads(result.output)
    assert payload["ok"] is True
    assert payload["command"] == "display scan"
    assert payload["data"]["candidateCount"] == 0
    assert payload["warnings"] == []
    assert payload["artifacts"] == []


def test_display_scan_command_returns_nonzero_when_candidates_exist(
    tmp_path: Path,
) -> None:
    path = tmp_path / "narrow-values.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.column_dimensions["A"].width = 3
    sheet["A1"] = 123456
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "display",
            "scan",
            "--json",
            json.dumps({"path": str(path)}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    payload = json.loads(result.output)
    assert payload["ok"] is False
    assert payload["command"] == "display scan"
    assert payload["errorCode"] == "display_overflow_candidates_found"
    assert payload["message"] == (
        "Cells may display as hashes because their rendered value exceeds column width"
    )
    assert payload["details"]["candidateCount"] == 1
    assert payload["details"]["candidates"][0]["cell"] == "A1"


def test_display_scan_command_rejects_invalid_max_results(
    tmp_path: Path,
) -> None:
    path = tmp_path / "values.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "display",
            "scan",
            "--json",
            json.dumps({"path": str(path), "maxResults": 0}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "display scan",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "maxResults", "expected": "positive integer | null"},
        "warnings": [],
        "artifacts": [],
    }


def test_display_scan_command_rejects_invalid_include_hidden(
    tmp_path: Path,
) -> None:
    path = tmp_path / "values.xlsx"
    workbook = Workbook()
    workbook.save(path)
    runner = CliRunner()

    result = runner.invoke(
        app,
        [
            "display",
            "scan",
            "--json",
            json.dumps({"path": str(path), "includeHidden": "yes"}),
            "--raw-output",
        ],
    )

    assert result.exit_code == 1
    assert json.loads(result.output) == {
        "ok": False,
        "command": "display scan",
        "errorCode": "invalid_field",
        "message": "Field has invalid type",
        "details": {"field": "includeHidden", "expected": "boolean"},
        "warnings": [],
        "artifacts": [],
    }
