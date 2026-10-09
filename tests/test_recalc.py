import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.recalc.libreoffice import recalc_workbook


def _mock_detect_libreoffice(monkeypatch: pytest.MonkeyPatch, executable: str) -> None:
    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.detect_libreoffice",
        lambda requested: {
            "available": True,
            "executable": executable,
            "source": "explicit",
            "platform": "test",
        },
    )


def test_recalc_workbook_returns_stable_error_when_executable_missing(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    Workbook().save(path)

    with pytest.raises(SheetOpsError) as error:
        recalc_workbook(path, tmp_path / "out", None, executable="/missing/libreoffice")

    assert error.value.error_code == "recalc_executable_missing"
    assert error.value.details["executable"] == "/missing/libreoffice"


def test_recalc_workbook_returns_stable_error_when_executable_broken(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    Workbook().save(path)
    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.detect_libreoffice",
        lambda requested: {
            "available": False,
            "errorCode": "libreoffice_broken",
            "executable": "/usr/local/bin/soffice",
            "source": "path",
            "platform": "linux",
            "stderr": "No such file or directory",
        },
    )

    with pytest.raises(SheetOpsError) as error:
        recalc_workbook(path, tmp_path / "out", None)

    assert error.value.error_code == "recalc_executable_broken"
    assert error.value.message == "LibreOffice executable is broken"
    assert error.value.details["stderr"] == "No such file or directory"


def test_recalc_workbook_creates_artifact_with_fake_executable(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 1
    sheet["B1"] = "=A1*2"
    workbook.save(path)
    fake = tmp_path / "fake-office"
    fake.write_text(
        "#!/bin/sh\n"
        'if [ "${1:-}" = "--headless" ] && [ "${2:-}" = "--version" ]; then\n'
        "  echo 'LibreOffice fake'\n"
        "  exit 0\n"
        "fi\n"
        'cp "$4" "$6/source.xlsx"\n',
        encoding="utf-8",
    )
    fake.chmod(0o755)

    result = recalc_workbook(
        path, tmp_path / "out", "source-recalculated.xlsx", executable=str(fake)
    )

    assert Path(result["recalculatedPath"]).exists()
    assert result["formulaCount"] == 1
    assert result["errorTokenCount"] == 0


def test_recalc_workbook_does_not_remove_source_when_output_dir_matches_parent(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 1
    sheet["B1"] = "=A1*2"
    workbook.save(path)
    output_dir = tmp_path
    output_path = output_dir / "source-recalculated.xlsx"

    def fake_run(
        command: list[str],
        check: bool,
        capture_output: bool,
        text: bool,
    ) -> subprocess.CompletedProcess[str]:
        source_path = Path(command[4])
        converted_path = Path(command[6]) / source_path.name
        converted_path.write_bytes(source_path.read_bytes())
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.subprocess.run",
        fake_run,
    )
    _mock_detect_libreoffice(monkeypatch, sys.executable)

    result = recalc_workbook(
        path,
        output_dir,
        "source-recalculated.xlsx",
        executable=sys.executable,
    )

    assert path.exists()
    assert Path(result["recalculatedPath"]) == output_path
    assert output_path.exists()


def test_recalc_workbook_preserves_source_bytes_with_same_parent_output(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 1
    sheet["B1"] = "=A1*2"
    workbook.save(path)
    original_bytes = path.read_bytes()
    output_path = tmp_path / "source-recalculated.xlsx"

    def fake_run(
        command: list[str],
        check: bool,
        capture_output: bool,
        text: bool,
    ) -> subprocess.CompletedProcess[str]:
        source_path = Path(command[4])
        converted_path = Path(command[6]) / source_path.name
        recalculated = Workbook()
        recalculated_sheet = recalculated.active
        recalculated_sheet["A1"] = 10
        recalculated_sheet["B1"] = "=A1*4"
        recalculated.save(converted_path)
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.subprocess.run",
        fake_run,
    )
    _mock_detect_libreoffice(monkeypatch, sys.executable)

    result = recalc_workbook(
        path,
        tmp_path,
        "source-recalculated.xlsx",
        executable=sys.executable,
    )

    assert path.read_bytes() == original_bytes
    assert Path(result["recalculatedPath"]) == output_path
    assert output_path.exists()
    assert output_path.read_bytes() != original_bytes


def test_recalc_workbook_rejects_exact_source_path_collision(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 7
    sheet["B1"] = "=A1*3"
    workbook.save(path)
    original_bytes = path.read_bytes()

    def fake_run(
        command: list[str],
        check: bool,
        capture_output: bool,
        text: bool,
    ) -> subprocess.CompletedProcess[str]:
        raise AssertionError(
            "LibreOffice should not be invoked for a source-path collision"
        )

    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.subprocess.run",
        fake_run,
    )
    _mock_detect_libreoffice(monkeypatch, sys.executable)

    with pytest.raises(SheetOpsError) as error:
        recalc_workbook(path, tmp_path, "source.xlsx", executable=sys.executable)

    assert error.value.error_code == "invalid_path"
    assert path.read_bytes() == original_bytes


@pytest.mark.parametrize("filename", [str(Path("/tmp/escape.xlsx")), "../escape.xlsx"])
def test_recalc_workbook_rejects_unsafe_output_filenames(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    filename: str,
) -> None:
    path = tmp_path / "source.xlsx"
    Workbook().save(path)

    def fake_run(
        command: list[str],
        check: bool,
        capture_output: bool,
        text: bool,
    ) -> subprocess.CompletedProcess[str]:
        source_path = Path(command[4])
        converted_path = Path(command[6]) / source_path.name
        converted_path.write_bytes(source_path.read_bytes())
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.subprocess.run",
        fake_run,
    )
    _mock_detect_libreoffice(monkeypatch, sys.executable)

    with pytest.raises(SheetOpsError) as error:
        recalc_workbook(path, tmp_path / "out", filename, executable=sys.executable)

    assert error.value.error_code == "invalid_path"


def test_recalc_workbook_wraps_subprocess_errors(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    Workbook().save(path)

    def fake_run(
        command: list[str],
        check: bool,
        capture_output: bool,
        text: bool,
    ) -> subprocess.CompletedProcess[str]:
        raise OSError("permission denied")

    monkeypatch.setattr(
        "zalos_sheetops.recalc.libreoffice.subprocess.run",
        fake_run,
    )
    _mock_detect_libreoffice(monkeypatch, sys.executable)

    with pytest.raises(SheetOpsError) as error:
        recalc_workbook(path, tmp_path / "out", None, executable=sys.executable)

    assert error.value.error_code == "recalc_failed"
