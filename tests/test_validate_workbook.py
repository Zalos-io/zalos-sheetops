from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.validation.workbook import (
    validate_workbook,
    validate_workbook_rules,
)


def test_validate_workbook_reports_missing_columns(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    workbook.save(path)

    result = validate_workbook(
        path,
        required_sheets=["Data"],
        required_columns={"Data": ["Date", "Reference"]},
        unique_keys={},
    )

    assert result["valid"] is False
    assert result["issues"] == [
        {
            "code": "missing_required_columns",
            "sheet": "Data",
            "missing": ["Reference"],
        }
    ]


def test_validate_workbook_checks_table_contract_without_openpyxl_materialization(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Reference", "Amount"])
    sheet.append(["2026-06-01", "INV-1", "10"])
    workbook.save(path)

    def fail_if_openpyxl_loads_workbook(*args: object, **kwargs: object) -> None:
        raise AssertionError("table validation should use XML streaming")

    monkeypatch.setattr(
        "openpyxl.load_workbook",
        fail_if_openpyxl_loads_workbook,
    )

    result = validate_workbook(
        path,
        required_sheets=["Data"],
        required_columns={"Data": ["Date", "Reference", "Amount"]},
        unique_keys={"Data": ["Date", "Reference"]},
    )

    assert result == {
        "valid": True,
        "issues": [],
        "uniqueKeysChecked": {"Data": ["Date", "Reference"]},
    }


def test_validate_workbook_detects_duplicate_keys_without_openpyxl_materialization(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Reference", "Amount"])
    sheet.append(["2026-06-01", "INV-1", "10"])
    sheet.append(["2026-06-01", "INV-1", "20"])
    workbook.save(path)

    def fail_if_openpyxl_loads_workbook(*args: object, **kwargs: object) -> None:
        raise AssertionError("table validation should use XML streaming")

    monkeypatch.setattr(
        "openpyxl.load_workbook",
        fail_if_openpyxl_loads_workbook,
    )

    result = validate_workbook(
        path,
        required_sheets=["Data"],
        required_columns={"Data": ["Date", "Reference", "Amount"]},
        unique_keys={"Data": ["Date", "Reference"]},
    )

    assert result["valid"] is False
    assert result["issues"] == [
        {
            "code": "duplicate_unique_key",
            "sheet": "Data",
            "keyColumns": ["Date", "Reference"],
            "duplicates": [
                {
                    "key": {"Date": "2026-06-01", "Reference": "INV-1"},
                    "rows": [2, 3],
                }
            ],
        }
    ]


def test_validate_workbook_rules_accepts_summary_freeform_sheet(
    tmp_path: Path,
) -> None:
    path = tmp_path / "recon.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Chase Payment Tech Summary"
    summary["A4"] = "FX adjustment"
    summary["B4"] = "=-'Merchant'!L1"
    merchant = workbook.create_sheet("Merchant")
    merchant.append(["Merchant Order Number", "Deposit Amount in USD"])
    merchant.append(["P-001", 10])
    workbook.save(path)

    result = validate_workbook_rules(
        path,
        [
            {
                "name": "Summary",
                "kind": "summary",
                "requiredCells": [
                    {"cell": "A1", "equals": "Chase Payment Tech Summary"}
                ],
                "requiredFormulas": [{"cell": "B4", "contains": "Merchant"}],
            },
            {
                "name": "Merchant",
                "kind": "table",
                "requiredColumns": [
                    "Merchant Order Number",
                    "Deposit Amount in USD",
                ],
                "uniqueKeys": ["Merchant Order Number"],
            },
        ],
    )

    assert result["valid"] is True
    assert result["issues"] == []


def test_validate_workbook_rules_checks_table_only_without_openpyxl_materialization(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "recon.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Merchant"
    sheet.append(["Merchant Order Number", "Deposit Amount in USD"])
    sheet.append(["P-001", 10])
    workbook.save(path)

    def fail_if_openpyxl_loads_workbook(*args: object, **kwargs: object) -> None:
        raise AssertionError("table validation should use XML streaming")

    monkeypatch.setattr(
        "openpyxl.load_workbook",
        fail_if_openpyxl_loads_workbook,
    )

    result = validate_workbook_rules(
        path,
        [
            {
                "name": "Merchant",
                "kind": "table",
                "requiredColumns": [
                    "Merchant Order Number",
                    "Deposit Amount in USD",
                ],
                "uniqueKeys": ["Merchant Order Number"],
            }
        ],
    )

    assert result == {
        "valid": True,
        "issues": [],
        "sheetsChecked": ["Merchant"],
    }


def test_validate_workbook_rules_reports_missing_summary_formula(
    tmp_path: Path,
) -> None:
    path = tmp_path / "recon.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Chase Payment Tech Summary"
    summary["B4"] = 25
    workbook.save(path)

    result = validate_workbook_rules(
        path,
        [
            {
                "name": "Summary",
                "kind": "summary",
                "requiredFormulas": [{"cell": "B4", "contains": "Merchant"}],
            }
        ],
    )

    assert result["valid"] is False
    assert result["issues"] == [
        {
            "code": "missing_required_formula",
            "sheet": "Summary",
            "cell": "B4",
            "value": "25",
        }
    ]


def test_validate_workbook_detects_header_row_below_title_row(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Quarterly revenue report", None])
    sheet.append(["Date", "Reference"])
    sheet.append(["2026-06-01", "INV-1"])
    workbook.save(path)

    result = validate_workbook(
        path,
        required_sheets=["Data"],
        required_columns={"Data": ["Date", "Reference"]},
        unique_keys={},
    )

    assert result == {"valid": True, "issues": [], "uniqueKeysChecked": {}}


def test_validate_workbook_rejects_duplicate_canonical_headers(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Amount", " amount "])
    sheet.append(["10", "20"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        validate_workbook(
            path,
            required_sheets=["Data"],
            required_columns={"Data": ["Amount"]},
            unique_keys={},
        )

    assert error.value.error_code == "invalid_headers"
    assert error.value.details == {"duplicates": ["amount"], "blankColumns": []}


def test_validate_workbook_rejects_blank_canonical_headers(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "   "])
    sheet.append(["2026-06-01", "10"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        validate_workbook(
            path,
            required_sheets=["Data"],
            required_columns={"Data": ["Date"]},
            unique_keys={},
        )

    assert error.value.error_code == "invalid_headers"
    assert error.value.details == {"duplicates": [], "blankColumns": [2]}


def test_validate_workbook_rejects_ambiguous_required_column_header_rows(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Reference"])
    sheet.append(["Date", "Reference"])
    sheet.append(["2026-06-01", "INV-1"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        validate_workbook(
            path,
            required_sheets=["Data"],
            required_columns={"Data": ["Date", "Reference"]},
            unique_keys={},
        )

    assert error.value.error_code == "ambiguous_header_row"
    assert error.value.details == {
        "sheet": "Data",
        "headerRows": [1, 2],
        "requiredColumns": ["Date", "Reference"],
    }


def test_validate_workbook_rejects_ambiguous_required_columns_with_extra_cells(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Reference", "Notes"])
    sheet.append(["Date", "Reference"])
    sheet.append(["2026-06-01", "INV-1", "memo"])
    workbook.save(path)

    with pytest.raises(SheetOpsError) as error:
        validate_workbook(
            path,
            required_sheets=["Data"],
            required_columns={"Data": ["Date", "Reference"]},
            unique_keys={},
        )

    assert error.value.error_code == "ambiguous_header_row"
    assert error.value.details == {
        "sheet": "Data",
        "headerRows": [1, 2],
        "requiredColumns": ["Date", "Reference"],
    }


def test_validate_workbook_reports_duplicate_unique_keys(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Reference", "Amount"])
    sheet.append(["2026-06-01", "INV-1", "10"])
    sheet.append(["2026-06-01", "INV-1", "20"])
    workbook.save(path)

    result = validate_workbook(
        path,
        required_sheets=["Data"],
        required_columns={"Data": ["Date", "Reference", "Amount"]},
        unique_keys={"Data": ["Date", "Reference"]},
    )

    assert result["valid"] is False
    assert result["issues"] == [
        {
            "code": "duplicate_unique_key",
            "sheet": "Data",
            "keyColumns": ["Date", "Reference"],
            "duplicates": [
                {
                    "key": {"Date": "2026-06-01", "Reference": "INV-1"},
                    "rows": [2, 3],
                }
            ],
        }
    ]


def test_validate_workbook_reports_missing_unique_key_columns(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount"])
    sheet.append(["2026-06-01", "10"])
    workbook.save(path)

    result = validate_workbook(
        path,
        required_sheets=["Data"],
        required_columns={"Data": ["Date", "Amount"]},
        unique_keys={"Data": ["Date", "Reference"]},
    )

    assert result["valid"] is False
    assert result["issues"] == [
        {
            "code": "missing_unique_key_columns",
            "sheet": "Data",
            "missing": ["Reference"],
        }
    ]


def test_validate_workbook_rejects_invalid_workbook_content(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("not really an xlsx", encoding="utf-8")

    with pytest.raises(SheetOpsError) as error:
        validate_workbook(
            path,
            required_sheets=["Data"],
            required_columns={"Data": ["Date"]},
            unique_keys={},
        )

    assert error.value.error_code == "invalid_workbook"
    assert error.value.details == {"path": str(path)}


def test_validate_workbook_rejects_malformed_zip_valid_workbook(tmp_path: Path) -> None:
    path = tmp_path / "broken-xml.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types")

    with pytest.raises(SheetOpsError) as error:
        validate_workbook(
            path,
            required_sheets=["Data"],
            required_columns={"Data": ["Date"]},
            unique_keys={},
        )

    assert error.value.error_code == "invalid_workbook"
    assert error.value.details == {"path": str(path)}
