from pathlib import Path
from typing import TYPE_CHECKING

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.cellref import range_boundaries

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet
from zalos_sheetops.formulas.xml_scan import (
    XmlFormulaScanUnsupported,
    scan_xlsx_formulas_xml,
)

ERROR_TOKENS = {
    "#REF!",
    "#VALUE!",
    "#DIV/0!",
    "#NAME?",
    "#NULL!",
    "#NUM!",
    "#N/A",
    "#CALC!",
    "#SPILL!",
}
DEFAULT_INCLUDE = ["formulas", "errors"]
VALID_INCLUDE = set(DEFAULT_INCLUDE)
VALID_ENGINES = {"auto", "xml", "openpyxl"}


def scan_formulas(
    path: Path,
    sheet_name: str | None = None,
    cell_range: str | None = None,
    max_results: int | None = None,
    engine: str = "auto",
    include: list[str] | None = None,
    fail_fast: bool = False,
) -> dict[str, object]:
    _validate_max_results(max_results)
    _validate_engine(engine)
    include_items = _validate_include(include)
    _validate_fail_fast(fail_fast)
    if engine in {"auto", "xml"}:
        try:
            xml_result = scan_xlsx_formulas_xml(
                path,
                sheet_name,
                cell_range,
                max_results,
                ERROR_TOKENS,
                include_items,
                fail_fast,
            )
            return _build_scan_result(
                "xml",
                include_items,
                int(xml_result["formulaCount"]),
                int(xml_result["errorTokenCount"]),
                xml_result["formulas"],
                xml_result["errorTokens"],
                bool(xml_result["stoppedEarly"]),
                sheet_name,
                cell_range,
            )
        except XmlFormulaScanUnsupported as error:
            if engine == "xml":
                raise SheetOpsError(
                    "unsupported_formula_scan_engine",
                    "XML formula scan cannot handle this workbook feature",
                    error.details,
                ) from error

    return _scan_formulas_openpyxl(
        path,
        sheet_name,
        cell_range,
        max_results,
        include_items,
        fail_fast,
    )


def _scan_formulas_openpyxl(
    path: Path,
    sheet_name: str | None,
    cell_range: str | None,
    max_results: int | None,
    include: list[str],
    fail_fast: bool,
) -> dict[str, object]:
    formula_workbook = None
    value_workbook = None
    formulas = {"count": 0, "items": []}
    error_tokens = {"count": 0, "items": [], "stoppedEarly": False}
    try:
        if "formulas" in include:
            formula_workbook = _load_workbook(path, data_only=False)
            formulas = _scan_formula_cells(
                formula_workbook,
                sheet_name,
                cell_range,
                max_results,
            )
        if "errors" in include:
            value_workbook = _load_workbook(path, data_only=True)
            error_tokens = _scan_error_tokens(
                value_workbook,
                sheet_name,
                cell_range,
                max_results,
                fail_fast,
            )
    finally:
        if formula_workbook is not None:
            formula_workbook.close()
        if value_workbook is not None:
            value_workbook.close()

    return _build_scan_result(
        "openpyxl",
        include,
        int(formulas["count"]),
        int(error_tokens["count"]),
        formulas["items"],
        error_tokens["items"],
        bool(error_tokens["stoppedEarly"]),
        sheet_name,
        cell_range,
    )


def _build_scan_result(
    engine: str,
    include: list[str],
    formula_count: int,
    error_token_count: int,
    formulas: object,
    error_tokens: object,
    stopped_early: bool,
    sheet_name: str | None,
    cell_range: str | None,
) -> dict[str, object]:
    formula_items = list(formulas) if isinstance(formulas, list) else []
    error_token_items = list(error_tokens) if isinstance(error_tokens, list) else []
    returned_formula_count = len(formula_items)
    returned_error_token_count = len(error_token_items)
    return {
        "engine": engine,
        "include": include,
        "formulaCount": formula_count,
        "errorTokenCount": error_token_count,
        "returnedFormulaCount": returned_formula_count,
        "returnedErrorTokenCount": returned_error_token_count,
        "truncated": (
            not stopped_early
            and (
                returned_formula_count < formula_count
                or returned_error_token_count < error_token_count
            )
        ),
        "stoppedEarly": stopped_early,
        "scope": {"sheet": sheet_name, "range": cell_range},
        "formulas": formula_items,
        "errorTokens": error_token_items,
    }


def _scan_formula_cells(
    workbook: "Workbook",
    sheet_name: str | None,
    cell_range: str | None,
    max_results: int | None,
) -> dict[str, object]:
    count = 0
    formulas: list[dict[str, str]] = []
    for sheet in _resolve_sheets(workbook, sheet_name):
        for cell in _iter_cells(sheet, cell_range):
            value = cell.value
            if not isinstance(value, str) or not value.startswith("="):
                continue
            count += 1
            if max_results is None or len(formulas) < max_results:
                formulas.append(
                    {
                        "sheet": sheet.title,
                        "cell": cell.coordinate,
                        "formula": value,
                    }
                )
    return {"count": count, "items": formulas}


def _scan_error_tokens(
    workbook: "Workbook",
    sheet_name: str | None,
    cell_range: str | None,
    max_results: int | None,
    fail_fast: bool,
) -> dict[str, object]:
    count = 0
    error_tokens: list[dict[str, str]] = []
    for sheet in _resolve_sheets(workbook, sheet_name):
        for cell in _iter_cells(sheet, cell_range):
            value = cell.value
            if value not in ERROR_TOKENS:
                continue
            count += 1
            if max_results is None or len(error_tokens) < max_results:
                error_tokens.append(
                    {
                        "sheet": sheet.title,
                        "cell": cell.coordinate,
                        "value": value,
                    }
                )
            if fail_fast:
                return {"count": count, "items": error_tokens, "stoppedEarly": True}
    return {"count": count, "items": error_tokens, "stoppedEarly": False}


def _resolve_sheets(
    workbook: "Workbook",
    sheet_name: str | None,
) -> list["Worksheet"]:
    if sheet_name is None:
        return list(workbook.worksheets)
    if sheet_name not in workbook.sheetnames:
        raise SheetOpsError(
            "invalid_sheet",
            "Requested sheet does not exist",
            {"sheet": sheet_name},
        )
    return [workbook[sheet_name]]


def _iter_cells(sheet: "Worksheet", cell_range: str | None) -> object:
    if cell_range is None:
        for row in sheet.iter_rows():
            yield from row
        return

    min_col, min_row, max_col, max_row = _parse_cell_range(cell_range)
    for row in sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        yield from row


def _parse_cell_range(cell_range: str) -> tuple[int, int, int, int]:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except ValueError as error:
        raise SheetOpsError(
            "invalid_range",
            "Cell range must be a bounded A1-style range",
            {"range": cell_range},
        ) from error
    if None in {min_col, min_row, max_col, max_row}:
        raise SheetOpsError(
            "invalid_range",
            "Cell range must be a bounded A1-style range",
            {"range": cell_range},
        )
    return min_col, min_row, max_col, max_row


def _validate_max_results(max_results: int | None) -> None:
    if max_results is None:
        return
    if (
        isinstance(max_results, bool)
        or not isinstance(max_results, int)
        or max_results < 1
    ):
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "maxResults", "expected": "positive integer | null"},
        )


def _validate_engine(engine: str) -> None:
    if engine not in VALID_ENGINES:
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "engine", "expected": "auto | xml | openpyxl"},
        )


def _validate_include(include: list[str] | None) -> list[str]:
    if include is None:
        return list(DEFAULT_INCLUDE)
    if (
        not isinstance(include, list)
        or not include
        or any(
            not isinstance(item, str) or item not in VALID_INCLUDE for item in include
        )
    ):
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "include", "expected": "non-empty list[formulas | errors]"},
        )
    include_set = set(include)
    return [item for item in DEFAULT_INCLUDE if item in include_set]


def _validate_fail_fast(fail_fast: bool) -> None:
    if not isinstance(fail_fast, bool):
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "failFast", "expected": "boolean"},
        )


def _load_workbook(path: Path, data_only: bool) -> "Workbook":
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=data_only, read_only=True)
    except Exception as error:
        raise SheetOpsError(
            "invalid_workbook",
            "Workbook could not be opened",
            {"path": str(path)},
        ) from error
