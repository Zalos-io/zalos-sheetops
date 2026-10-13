import argparse
import hashlib
import importlib.metadata
import json
import platform
import random
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

from .fixtures import file_sha256, generate_readbench, prepare_public_inputs
from .rubric import ORACLE_ISSUES
from .scoring import score_outputs, summarize
from .tasks import reviewed_selection, select_tasks


def execute(case: dict, adapter: str, timeout: float) -> dict:
    start = time.perf_counter()
    try:
        process = subprocess.run(
            [
                sys.executable,
                "-m",
                "zalos_sheetops.benchmarks.public.worker",
                json.dumps(case),
                adapter,
            ],
            capture_output=True,
            text=True,
            timeout=timeout,
        )
    except subprocess.TimeoutExpired:
        return {"status": "timeout", "seconds": time.perf_counter() - start}
    try:
        result = json.loads(process.stdout)
    except (ValueError, TypeError):
        result = {"status": "error", "reason": process.stderr[-2000:]}
    result.update(
        seconds=time.perf_counter() - start,
        exitCode=process.returncode,
        outputBytes=len(process.stdout.encode("utf-8")),
    )
    if process.returncode:
        result["status"] = "error"
    return result


def build_cases(root: Path, sizes: list[int]) -> list[dict]:
    cases = []
    for size in sizes:
        file = root / "generated" / f"readbench-{size}.xlsx"
        if not file.exists():
            generate_readbench(file, size)
        common = {"path": str(file), "sheet": "data", "source": "excel-read-benchmark"}
        for name, row in [
            ("start", 2),
            ("middle", size // 2 + 1),
            ("end", max(2, size - 8)),
        ]:
            cases.append(
                {
                    **common,
                    "id": f"{size}-cells-{name}",
                    "operation": "cells",
                    "range": f"A{row}:L{min(row + 9, size + 1)}",
                }
            )
            cases.append(
                {
                    **common,
                    "id": f"{size}-scalar-cells-{name}",
                    "operation": "cells",
                    "range": f"A{row}:I{min(row + 9, size + 1)}",
                }
            )
        cases.append({**common, "id": f"{size}-table", "operation": "table"})
    manifest = json.loads(
        (root / "excelbench/manifest.json").read_text(encoding="utf-8")
    )
    from openpyxl import load_workbook

    for entry in manifest["files"]:
        file = root / "excelbench" / entry["path"]
        if not file.exists():
            continue
        book = load_workbook(file, read_only=True, data_only=False)
        try:
            for sheet in book:
                cases.append(
                    {
                        "id": f"fidelity-{entry['feature']}-{sheet.title}",
                        "path": str(file),
                        "sheet": sheet.title,
                        "source": "ExcelBench",
                        "operation": "cells",
                        "formulas": True,
                        "range": f"A1:J{min(100, sheet.max_row or 1)}",
                    }
                )
        finally:
            book.close()
    return cases


def compare_measurement(reference: dict, measured: dict) -> dict:
    result = measured.copy()
    if reference["status"] != "ok":
        result["status"] = "oracle_failure"
        result["oracleReason"] = reference
    elif result["status"] == "ok":
        matched = result["digest"] == reference["digest"]
        result["status"] = "pass" if matched else "incorrect"
        if not matched:
            result["sampleComparison"] = score_outputs(
                reference["sample"], result["sample"]
            )
            result["expectedDigest"] = reference["digest"]
    return result


def replay_case(
    root: Path,
    output: Path,
    sizes: list[int],
    case_id: str,
    adapter: str,
    timeout: float = 120,
) -> dict:
    cases = [case for case in build_cases(root, sizes) if case["id"] == case_id]
    if len(cases) != 1:
        return {"status": "case_not_found", "case": case_id}
    case = cases[0]
    reference = execute(case, "openpyxl", timeout)
    measured = compare_measurement(reference, execute(case, adapter, timeout))
    result = {
        "case": case_id,
        "caseDefinition": case,
        "adapter": adapter,
        "inputSha256": file_sha256(Path(case["path"])),
        **measured,
    }
    output.mkdir(parents=True, exist_ok=True)
    (output / "replay.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


def run_suite(
    root: Path, output: Path, sizes: list[int], repeats: int, timeout: float = 120
) -> dict:
    if repeats < 1:
        raise ValueError("repeats must be positive")
    output.mkdir(parents=True, exist_ok=True)
    cases = build_cases(root, sizes)
    adapters = [
        "openpyxl",
        "sheetops-api",
        "sheetops-cli",
        "calamine",
        "fastexcel",
        "sheetops-hybrid",
    ]
    inputs = {}
    for case in cases:
        path = Path(case["path"])
        if str(path) not in inputs:
            with path.open("rb") as source:
                inputs[str(path)] = hashlib.file_digest(source, "sha256").hexdigest()
    versions = {}
    for name in [
        "zalos-sheetops",
        "openpyxl",
        "fastexcel",
        "python-calamine",
        "XlsxWriter",
        "pyarrow",
    ]:
        try:
            versions[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            versions[name] = None
    metadata = {
        "schemaVersion": 1,
        "startedAt": datetime.now(timezone.utc).isoformat(),
        "python": sys.version,
        "platform": platform.platform(),
        "machine": platform.machine(),
        "versions": versions,
        "repeats": repeats,
        "timeoutSeconds": timeout,
        "seed": 20261009,
        "cases": cases,
        "adapters": adapters,
        "harnessSha256": {
            name: file_sha256(Path(__file__).with_name(name))
            for name in [
                "runner.py",
                "adapters.py",
                "worker.py",
                "scoring.py",
                "fixtures.py",
            ]
        },
        "sourceManifestSha256": file_sha256(root / "sources.json")
        if (root / "sources.json").exists()
        else None,
        "inputSha256": inputs,
        "cachePolicy": "fresh process; OS cache uncontrolled",
        "memoryMetric": "maximum RSS of any one worker/CLI process; not aggregate RSS",
        "oracle": "independent openpyxl reads; separate untimed reference per case",
    }
    (output / "metadata.json").write_text(
        json.dumps(metadata, indent=2), encoding="utf-8"
    )
    records = []
    with (output / "operations.jsonl").open("w", encoding="utf-8") as destination:
        for case in cases:
            reference = execute(case, "openpyxl", timeout)
            for repeat in range(repeats):
                order = adapters.copy()
                random.Random(20261009 + repeat).shuffle(order)
                for adapter in order:
                    measured = execute(case, adapter, timeout)
                    measured = compare_measurement(reference, measured)
                    record = {
                        "case": case["id"],
                        "adapter": adapter,
                        "repeat": repeat,
                        **measured,
                    }
                    records.append(record)
                    destination.write(json.dumps(record, ensure_ascii=False) + "\n")
                    destination.flush()
                    print(
                        case["id"],
                        adapter,
                        measured["status"],
                        round(measured["seconds"], 3),
                        flush=True,
                    )
    for path, expected in inputs.items():
        with Path(path).open("rb") as source:
            if hashlib.file_digest(source, "sha256").hexdigest() != expected:
                raise RuntimeError(f"Input mutated during benchmark: {path}")
    for name, expected in metadata["harnessSha256"].items():
        if file_sha256(Path(__file__).with_name(name)) != expected:
            raise RuntimeError(f"Benchmark source changed during measurements: {name}")
    result = {"summaries": summarize(records), "recordCount": len(records)}
    (output / "summary.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Public spreadsheet benchmark runner")
    parser.add_argument(
        "action", choices=["prepare", "operations", "select-tasks", "replay"]
    )
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--archive", type=Path)
    parser.add_argument("--sizes", default="25000,100000,400000")
    parser.add_argument("--repeats", type=int, default=3)
    parser.add_argument("--case")
    parser.add_argument(
        "--adapter",
        choices=[
            "openpyxl",
            "sheetops-api",
            "sheetops-cli",
            "calamine",
            "fastexcel",
            "sheetops-hybrid",
        ],
        default="sheetops-cli",
    )
    args = parser.parse_args()
    if args.action == "prepare":
        prepare_public_inputs(args.data, args.archive)
    elif args.action == "operations":
        run_suite(
            args.data,
            args.output,
            [int(n) for n in args.sizes.split(",")],
            args.repeats,
        )
    elif args.action == "replay":
        result = replay_case(
            args.data,
            args.output,
            [int(n) for n in args.sizes.split(",")],
            args.case,
            args.adapter,
        )
        print(json.dumps(result, ensure_ascii=False))
    else:
        dataset = args.data / "verified/spreadsheetbench_verified_400/dataset.json"
        pool = json.loads(dataset.read_text(encoding="utf-8"))
        initial = select_tasks(pool)
        tasks = reviewed_selection(pool, initial)
        args.output.mkdir(parents=True, exist_ok=True)
        (args.output / "tasks.json").write_text(
            json.dumps(tasks, indent=2), encoding="utf-8"
        )
        selection = {
            "seed": 20261009,
            "baseCount": len(initial),
            "taskCount": len(tasks),
            "baseTaskIds": [t["id"] for t in initial],
            "supplementalTaskIds": [t["id"] for t in tasks[len(initial) :]],
            "oracleIssues": {
                str(t["id"]): ORACLE_ISSUES[str(t["id"])]
                for t in initial
                if str(t["id"]) in ORACLE_ISSUES
            },
            "replacementPolicy": "same seed and instruction type; original "
            "attempts retained; no selection by model performance",
        }
        (args.output / "selection.json").write_text(
            json.dumps(selection, indent=2), encoding="utf-8"
        )
        print(json.dumps(selection))


if __name__ == "__main__":
    main()
