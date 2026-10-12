"""Feature correctness contracts; timings are not cross-library performance rankings."""

import argparse
import csv
import json
import subprocess
import sys
import time
from pathlib import Path

from openpyxl import load_workbook


def cli_record(case: str, command: list[str], payload: dict, expected: dict) -> dict:
    start = time.perf_counter()
    process = subprocess.run(
        [
            sys.executable,
            "-m",
            "zalos_sheetops.cli",
            *command,
            "--json",
            json.dumps(payload),
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )
    elapsed = time.perf_counter() - start
    try:
        response = json.loads(process.stdout)
    except ValueError:
        return {"case": case, "status": "error", "seconds": elapsed}
    actual = response.get("data", response)
    differences = []
    for path, wanted in expected.items():
        value = actual
        for part in path.split("."):
            value = value[int(part)] if isinstance(value, list) else value.get(part)
        if value != wanted:
            differences.append({"property": path, "expected": wanted, "actual": value})
    return {
        "case": case,
        "status": "pass" if not differences else "incorrect",
        "seconds": elapsed,
        "exitCode": process.returncode,
        "differences": differences,
        "command": command,
        "payload": payload,
    }


def run_contracts(root: Path, output: Path) -> list[dict]:
    records = []
    fixture_root = root / "excelbench"
    manifest = json.loads((fixture_root / "manifest.json").read_text())
    for entry in manifest["files"]:
        path = fixture_root / entry["path"]
        if not path.exists():
            continue
        book = load_workbook(path)
        expected = {"workbook.sheetCount": len(book.sheetnames)}
        for index, sheet in enumerate(book):
            prefix = f"sheets.{index}"
            expected.update(
                {
                    f"{prefix}.name": sheet.title,
                    f"{prefix}.state": sheet.sheet_state,
                    f"{prefix}.maxRow": sheet.max_row,
                    f"{prefix}.maxColumn": sheet.max_column,
                    f"{prefix}.mergedRanges": sorted(
                        str(r) for r in sheet.merged_cells
                    ),
                    f"{prefix}.formulaCount": sum(
                        c.data_type == "f" for row in sheet for c in row
                    ),
                }
            )
        records.append(
            cli_record(
                f"metadata-{entry['feature']}",
                ["inspect"],
                {"path": str(path), "engine": "xml"},
                expected,
            )
        )
        if entry["feature"] == "number_formats":
            for case in entry["test_cases"]:
                address = f"B{case['row']}"
                stored = book[book.sheetnames[0]][address].number_format
                record = cli_record(
                    "number_formats",
                    ["styles", "read"],
                    {"path": str(path), "sheet": book.sheetnames[0], "range": address},
                    {"cells.0.style.numberFormat": stored},
                )
                if stored != case["expected"]["number_format"]:
                    record["fixtureManifestDifference"] = {
                        "manifest": case["expected"]["number_format"],
                        "stored": stored,
                    }
                records.append(record)
        if entry["feature"] == "hyperlinks":
            checks = []
            for case in entry["test_cases"]:
                wanted = case["expected"]["hyperlink"]
                checks.append(
                    {k: wanted[k] for k in ["cell", "target", "display", "tooltip"]}
                )
                if wanted.get("internal"):
                    checks[-1]["target"] = (
                        "#" + book["hyperlinks"][wanted["cell"]].hyperlink.location
                    )
            records.append(
                cli_record(
                    "hyperlinks",
                    ["hyperlinks", "assert"],
                    {"path": str(path), "sheet": "hyperlinks", "checks": checks},
                    {"valid": True, "checkedCount": len(checks), "failureCount": 0},
                )
            )
        book.close()
    output.mkdir(parents=True, exist_ok=True)
    left, right = output / "reconcile-left.csv", output / "reconcile-right.csv"
    rows = [
        ["reference", "amount"],
        ["001", "0.30"],
        ["002", "100.01"],
        ["003", "2.00"],
    ]
    for path, values in [
        (left, rows),
        (right, [rows[0], rows[1], ["002", "100.00"], ["004", "2.00"]]),
    ]:
        with path.open("w", newline="") as destination:
            csv.writer(destination).writerows(values)
    payload = {
        "left": {"path": str(left)},
        "right": {"path": str(right)},
        "leftKeyColumns": ["reference"],
        "rightKeyColumns": ["reference"],
        "leftAmountColumn": "amount",
        "rightAmountColumn": "amount",
        "amountTolerance": "0.00",
    }
    records.append(
        cli_record(
            "decimal-reconciliation",
            ["reconcile"],
            payload,
            {
                "summary": {
                    "matched": 1,
                    "leftUnmatched": 1,
                    "rightUnmatched": 1,
                    "variance": 1,
                },
                "variance.0.difference": "0.01",
                "matched.0.key.reference": "001",
            },
        )
    )
    (output / "contracts.jsonl").write_text(
        "".join(json.dumps(r) + "\n" for r in records)
    )
    return records


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    for record in run_contracts(args.data, args.output):
        print(record["case"], record["status"], flush=True)


if __name__ == "__main__":
    main()
