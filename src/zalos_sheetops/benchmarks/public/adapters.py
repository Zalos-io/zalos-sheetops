import json
import subprocess
import sys
from pathlib import Path

from .scoring import scalar


def openpyxl_cells(case: dict) -> list[list[object]]:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries

    bounds = range_boundaries(case["range"])
    options = dict(
        min_col=bounds[0], min_row=bounds[1], max_col=bounds[2], max_row=bounds[3]
    )
    formula = load_workbook(case["path"], read_only=True, data_only=False)
    cached = load_workbook(case["path"], read_only=True, data_only=True)
    try:
        rows = []
        for fr, vr in zip(
            formula[case["sheet"]].iter_rows(**options),
            cached[case["sheet"]].iter_rows(**options),
            strict=True,
        ):
            row = []
            for f, v in zip(fr, vr, strict=True):
                text = (
                    f.value if f.data_type == "f" and isinstance(f.value, str) else None
                )
                row.append(
                    [scalar(v.value), text, scalar(v.value)]
                    if case.get("formulas")
                    else scalar(v.value)
                )
            rows.append(row)
        return rows
    finally:
        formula.close()
        cached.close()


def sheetops_cells(case: dict, cli: bool = False) -> list[list[object]]:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    include = (
        ["values", "formulas", "cachedValues"] if case.get("formulas") else ["values"]
    )
    if cli:
        request = {
            "path": case["path"],
            "sheet": case["sheet"],
            "range": case["range"],
            "include": include,
            "engine": "xml",
        }
        result = subprocess.run(
            [
                str(Path(sys.executable).with_name("sheetops")),
                "cells",
                "read",
                "--json",
                json.dumps(request),
            ],
            capture_output=True,
            check=True,
            text=True,
        )
        response = json.loads(result.stdout)
        if not response.get("ok"):
            raise RuntimeError(response)
        cells = response["data"]["cells"]
    else:
        from zalos_sheetops.io.read_cells import read_cells

        cells = read_cells(
            Path(case["path"]), case["sheet"], case["range"], include, engine="xml"
        )["cells"]
    by_address = {c["address"]: c for c in cells}
    left, top, right, bottom = range_boundaries(case["range"])
    rows = []
    for r in range(top, bottom + 1):
        row = []
        for c in range(left, right + 1):
            item = by_address.get(f"{get_column_letter(c)}{r}", {})
            row.append(
                [item.get("value"), item.get("formula"), item.get("cachedValue")]
                if case.get("formulas")
                else item.get("value")
            )
        rows.append(row)
    return rows


def openpyxl_table(case: dict) -> list[list[object]]:
    from openpyxl import load_workbook

    book = load_workbook(case["path"], read_only=True, data_only=True)
    try:
        rows = list(book[case["sheet"]].iter_rows(values_only=True))
        return [["" if v is None else str(v) for v in row] for row in rows]
    finally:
        book.close()


def sheetops_table(case: dict, cli: bool = False) -> list[list[object]]:
    from zalos_sheetops.io.read_table import read_table

    if cli:
        request = {
            "path": case["path"],
            "sheet": case["sheet"],
            "headerStrategy": "row",
            "headerRow": 1,
        }
        process = subprocess.run(
            [
                str(Path(sys.executable).with_name("sheetops")),
                "read-table",
                "--json",
                json.dumps(request),
            ],
            capture_output=True,
            text=True,
            check=True,
        )
        response = json.loads(process.stdout)
        if not response.get("ok"):
            raise RuntimeError(response)
        result = response["data"]
    else:
        result = read_table(Path(case["path"]), case["sheet"], "row", 1)
    if result["truncated"]:
        raise RuntimeError("Full table was truncated")
    return [result["headers"], *result["rows"]]


def table_text(value: object) -> str:
    from datetime import date, datetime, time

    if value is None:
        return ""
    if isinstance(value, date) and not isinstance(value, datetime):
        return str(datetime.combine(value, time()))
    if type(value) is float and value.is_integer():
        return str(int(value))
    return str(value)


def external_table(case: dict, adapter: str) -> list[list[object]]:
    if adapter == "calamine":
        from python_calamine import CalamineWorkbook

        with CalamineWorkbook.from_path(case["path"]) as book:
            data = book.get_sheet_by_name(case["sheet"]).to_python(
                skip_empty_area=False
            )
    else:
        import fastexcel

        sheet = fastexcel.read_excel(case["path"]).load_sheet(
            case["sheet"], header_row=0
        )
        data = [
            sheet.to_arrow().column_names,
            *[list(r.values()) for r in sheet.to_arrow().to_pylist()],
        ]
    return [[table_text(v) for v in row] for row in data]


def external_cells(case: dict, adapter: str) -> list[list[object]]:
    if case.get("formulas"):
        raise NotImplementedError(
            "Adapter exposes values, not formula text/cache pairs"
        )
    from openpyxl.utils.cell import range_boundaries

    left, top, right, bottom = range_boundaries(case["range"])
    if adapter == "calamine":
        from python_calamine import CalamineWorkbook

        with CalamineWorkbook.from_path(case["path"]) as book:
            sheet = book.get_sheet_by_name(case["sheet"])
            data = sheet.to_python(skip_empty_area=False)
            # skip_empty_area=False includes leading blank rows and columns.
            offset_row = offset_col = 0
    else:
        import fastexcel

        sheet = fastexcel.read_excel(case["path"]).load_sheet(
            case["sheet"],
            header_row=None,
            skip_rows=top - 1,
            n_rows=bottom - top + 1,
            dtype_coercion="strict",
        )
        data = sheet.to_arrow().to_pylist()
        data = [list(row.values()) for row in data]
        offset_row, offset_col = top - 1, 0
    result = []
    for r in range(top - 1, bottom):
        source = data[r - offset_row] if 0 <= r - offset_row < len(data) else []
        result.append(
            [
                scalar(source[c - offset_col])
                if 0 <= c - offset_col < len(source)
                else None
                for c in range(left - 1, right)
            ]
        )
    return result


def read_case(case: dict, adapter: str) -> list[list[object]]:
    if case["operation"] == "table":
        if adapter == "openpyxl":
            return openpyxl_table(case)
        if adapter in {"sheetops-api", "sheetops-cli"}:
            return sheetops_table(case, cli=adapter == "sheetops-cli")
        if adapter in {"fastexcel", "calamine"}:
            return external_table(case, adapter)
        raise NotImplementedError("No equivalent full-table adapter")
    if adapter == "openpyxl":
        return openpyxl_cells(case)
    if adapter in {"sheetops-api", "sheetops-cli"}:
        return sheetops_cells(case, cli=adapter == "sheetops-cli")
    if adapter in {"calamine", "fastexcel"}:
        return external_cells(case, adapter)
    raise ValueError(f"Unknown adapter: {adapter}")
