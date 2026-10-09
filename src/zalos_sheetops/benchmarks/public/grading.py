import json
import posixpath
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from .scoring import canonical, scalar


def evaluated_blank_caches(path: Path) -> set[tuple[str, str]]:
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    relationship = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    blanks = set()
    with ZipFile(path) as archive:
        targets = {
            r.get("Id"): r.get("Target")
            for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        }
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in workbook.find(f"{namespace}sheets"):
            target = targets[sheet.get(relationship)]
            member = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
            tree = ET.fromstring(archive.read(member))
            for cell in tree.iter(f"{namespace}c"):
                cache = cell.find(f"{namespace}v")
                if (
                    cell.get("t") == "str"
                    and cell.find(f"{namespace}f") is not None
                    and cache is not None
                    and not cache.text
                ):
                    blanks.add((sheet.get("name"), cell.get("r")))
    return blanks


def recalculate(source: Path, destination: Path, soffice: Path) -> Path:
    destination.mkdir(parents=True, exist_ok=True)
    profile = destination / "lo-profile"
    process = subprocess.run(
        [
            str(soffice),
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(destination),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=90,
    )
    target = destination / source.name
    if process.returncode or not target.exists():
        raise RuntimeError("LibreOffice could not recalculate the evaluation copy")
    return target


def answer_cells(task: dict, sheetnames: list[str]) -> set[tuple[str, str]]:
    from openpyxl.utils.cell import get_column_letter

    position = task.get("answer_position", "")
    if not position:
        return set()
    if "!" in position:
        sheet, cell_range = position.rsplit("!", 1)
        sheet = sheet.strip("'\"")
    else:
        cell_range = position
        authored = task.get("answer_sheet")
        sheet = str(authored).strip("'\"") if authored else sheetnames[0]
    if sheet not in sheetnames:
        raise ValueError(f"Ambiguous answer sheet: {sheet}")
    left, top, right, bottom = range_boundaries(cell_range)
    if None in (left, top, right, bottom):
        raise ValueError("Answer position must be bounded")
    return {
        (sheet, f"{get_column_letter(c)}{r}")
        for r in range(top, bottom + 1)
        for c in range(left, right + 1)
    }


def cell_signature(cell: object) -> list:
    value = getattr(cell, "value", None)
    if getattr(cell, "data_type", None) == "f" and not isinstance(value, str):
        value = {
            "type": type(value).__name__,
            "attributes": dict(value),
            "text": getattr(value, "text", None),
        }
    # Style ids are workbook-local: compare authored properties, never ids.
    properties = []
    for name in ["font", "fill", "alignment", "border", "protection"]:
        item = getattr(cell, name, None)
        properties.append(str(item) if item is not None else None)
    return [
        scalar(value),
        getattr(cell, "number_format", None),
        properties,
        str(getattr(cell, "hyperlink", None)),
        str(getattr(cell, "comment", None)),
    ]


def structure_differences(expected: object, actual: object) -> list[dict]:
    differences = []
    expected_names = sorted(str(item) for item in expected.defined_names.values())
    actual_names = sorted(str(item) for item in actual.defined_names.values())
    if expected_names != actual_names:
        differences.append({"kind": "defined_names"})
    for sheet in expected:
        if sheet.title not in actual.sheetnames:
            continue
        output = actual[sheet.title]
        properties = {
            "merged_ranges": (
                sorted(str(r) for r in sheet.merged_cells),
                sorted(str(r) for r in output.merged_cells),
            ),
            "sheet_state": (sheet.sheet_state, output.sheet_state),
            "sheet_view": (
                sheet.sheet_view.showGridLines,
                output.sheet_view.showGridLines,
            ),
            "tables": (
                sorted(str(t) for t in sheet.tables.values()),
                sorted(str(t) for t in output.tables.values()),
            ),
            "local_defined_names": (
                sorted(str(t) for t in sheet.defined_names.values()),
                sorted(str(t) for t in output.defined_names.values()),
            ),
        }
        for kind, (wanted, observed) in properties.items():
            if wanted != observed:
                differences.append({"sheet": sheet.title, "kind": kind})
    return differences


def add_failure(counts: dict, samples: list[dict], kind: str, detail: dict) -> None:
    counts[kind] += 1
    if len(samples) < 20:
        samples.append(detail)


def grade_workbook(
    initial: Path,
    golden: Path,
    actual: Path,
    task: dict,
    soffice: Path | None = None,
    evaluation_dir: Path | None = None,
    evaluated: tuple[Path, Path] | None = None,
) -> dict:
    if not actual.exists():
        return {"status": "missing_output"}
    cached_golden, cached_actual = (
        evaluated if evaluated is not None else (golden, actual)
    )
    if soffice is not None and evaluated is None:
        if evaluation_dir is None:
            raise ValueError("evaluation_dir is required for recalculation")
        try:
            cached_golden = recalculate(golden, evaluation_dir / "golden", soffice)
            cached_actual = recalculate(actual, evaluation_dir / "actual", soffice)
        except (RuntimeError, subprocess.TimeoutExpired):
            return {"status": "recalculation_failure"}
    books = []
    try:
        for path, data_only in [
            (initial, False),
            (golden, False),
            (actual, False),
            (cached_golden, True),
            (cached_actual, True),
        ]:
            books.append(load_workbook(path, data_only=data_only))
        before, expected_raw, actual_raw, expected, result = books
        known_blanks = evaluated_blank_caches(cached_golden)
        structure = structure_differences(expected_raw, actual_raw)
        try:
            allowed = answer_cells(task, expected_raw.sheetnames)
        except (ValueError, TypeError):
            return {"status": "ambiguous_task_metadata"}
        if not allowed:
            return {"status": "ambiguous_task_metadata"}
        style_requested = any(
            word in task.get("instruction", "").lower()
            for word in [
                "font",
                "bold",
                "fill",
                "color",
                "border",
                "alignment",
                "formatting",
            ]
        )
        differences, preservation, caches = [], [], []
        counts = {
            "values": 0,
            "styles": 0,
            "formulas": 0,
            "preservation": 0,
            "structure": len(structure),
        }
        if expected.sheetnames != result.sheetnames:
            counts["structure"] += 1
            differences.append(
                {
                    "kind": "sheet_names",
                    "expected": expected.sheetnames,
                    "actual": result.sheetnames,
                }
            )
        for name in expected.sheetnames:
            if name not in result.sheetnames:
                continue
            exp_sheet, out_sheet = expected[name], result[name]
            old = before[name] if name in before.sheetnames else None
            raw_expected = expected_raw[name]
            raw_actual = actual_raw[name]
            max_row = max(exp_sheet.max_row, out_sheet.max_row)
            max_col = max(exp_sheet.max_column, out_sheet.max_column)
            for row in range(1, max_row + 1):
                for column in range(1, max_col + 1):
                    target, output = (
                        exp_sheet.cell(row, column),
                        out_sheet.cell(row, column),
                    )
                    address = target.coordinate
                    in_answer = (name, address) in allowed
                    style_changed = (
                        old is None
                        or cell_signature(old[address])[1:3]
                        != cell_signature(raw_expected[address])[1:3]
                    )
                    if (
                        style_requested
                        and (in_answer or style_changed)
                        and cell_signature(raw_expected[address])[1:3]
                        != cell_signature(raw_actual[address])[1:3]
                    ):
                        add_failure(
                            counts,
                            differences,
                            "styles",
                            {"sheet": name, "cell": address, "kind": "required_style"},
                        )
                    expected_formula = raw_expected[address].data_type == "f"
                    if (
                        in_answer
                        and expected_formula
                        and "formula" in task.get("instruction", "").lower()
                        and raw_actual[address].data_type != "f"
                    ):
                        add_failure(
                            counts,
                            differences,
                            "formulas",
                            {"sheet": name, "cell": address, "kind": "missing_formula"},
                        )
                    if (
                        in_answer
                        and expected_formula
                        and target.value is None
                        and (name, address) not in known_blanks
                    ):
                        caches.append({"sheet": name, "cell": address})
                    elif in_answer and canonical(target.value) != canonical(
                        output.value
                    ):
                        add_failure(
                            counts,
                            differences,
                            "values",
                            {
                                "sheet": name,
                                "cell": address,
                                "expected": scalar(target.value),
                                "actual": scalar(output.value),
                            },
                        )
                    if old is not None and (name, address) not in allowed:
                        unchanged = cell_signature(old[address]) == cell_signature(
                            raw_expected[address]
                        )
                        if unchanged and cell_signature(old[address]) != cell_signature(
                            raw_actual[address]
                        ):
                            add_failure(
                                counts,
                                preservation,
                                "preservation",
                                {"sheet": name, "cell": address},
                            )
            if "calibri" in task.get("instruction", "").lower():
                for row in raw_expected:
                    for cell in row:
                        if (
                            cell.has_style
                            and raw_actual[cell.coordinate].font.name != cell.font.name
                        ):
                            add_failure(
                                counts,
                                differences,
                                "styles",
                                {
                                    "sheet": name,
                                    "cell": cell.coordinate,
                                    "kind": "required_font",
                                },
                            )
        if caches:
            return {"status": "oracle_cache_missing", "cells": caches[:20]}
        status = "pass" if not any(counts.values()) else "incorrect"
        return {
            "status": status,
            "valueDifferences": differences,
            "preservationFailures": preservation,
            "structureDifferences": structure,
            "failureCounts": counts,
            "checks": {
                "answerValues": counts["values"] == 0,
                "explicitStyles": counts["styles"] == 0,
                "formulaRequirement": counts["formulas"] == 0,
                "preservation": counts["preservation"] == 0,
                "structure": counts["structure"] == 0,
            },
            "rubric": (
                "golden answer-region values; explicit style requests; "
                "names, merges, tables, visibility and unchanged cells preserved"
            ),
        }
    except (
        ValueError,
        KeyError,
        OSError,
        json.JSONDecodeError,
        BadZipFile,
        InvalidFileException,
        ParseError,
    ) as error:
        status = "oracle_failure" if len(books) < 2 else "invalid_output"
        return {"status": status, "reason": type(error).__name__}
    finally:
        for book in books:
            book.close()
