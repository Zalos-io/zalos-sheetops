"""Independent output checks, including actual external formula evaluation."""

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer

from zalos_sheetops.benchmarks.public.grading import (
    cell_signature,
    recalculate,
    structure_differences,
)
from zalos_sheetops.benchmarks.public.native import formula_caches
from zalos_sheetops.compare.edits import sheet_structure


def grade_case(case: dict, workspace: Path, soffice: Path | None = None) -> dict:
    path = workspace / case["output"]
    expected = Path(case["golden"]) / case["output"]
    failures = []
    result = {
        "correct": False,
        "failures": failures,
        "formulaEvaluation": "not_required",
    }
    if not path.exists():
        failures.append("missing_output")
        return result
    if case["kind"] == "csv":
        with (
            path.open(encoding="utf-8-sig", newline="") as actual_file,
            expected.open(newline="") as golden_file,
        ):
            if list(csv.reader(actual_file)) != list(csv.reader(golden_file)):
                failures.append("csv_values_or_schema")
    else:
        actual, golden = load_workbook(path), load_workbook(expected)
        if actual.sheetnames != golden.sheetnames:
            failures.append("sheet_names")
        if case["kind"] == "edit":
            failures.extend(str(item) for item in structure_differences(golden, actual))
            for sheet in golden:
                if sheet.title not in actual:
                    continue
                other = actual[sheet.title]
                if sheet_structure(sheet) != sheet_structure(other):
                    failures.append(f"sheet_features:{sheet.title}")
                for row in sheet:
                    for cell in row:
                        if cell_signature(cell) != cell_signature(
                            other[cell.coordinate]
                        ):
                            failures.append(f"cell:{sheet.title}!{cell.coordinate}")
                if (other.max_row, other.max_column) != (
                    sheet.max_row,
                    sheet.max_column,
                ):
                    failures.append(f"dimensions:{sheet.title}")
        else:
            for sheet in golden:
                if sheet.title not in actual:
                    continue
                other = actual[sheet.title]
                if (other.max_row, other.max_column) != (
                    sheet.max_row,
                    sheet.max_column,
                ):
                    failures.append(f"dimensions:{sheet.title}")
                for row in sheet:
                    for cell in row:
                        candidate = other[cell.coordinate]
                        location = f"{sheet.title}!{cell.coordinate}"
                        if cell.data_type != "f" and candidate.value != cell.value:
                            failures.append(f"data:{location}")
                        if isinstance(cell.value, datetime):
                            if not candidate.is_date:
                                failures.append(f"date_type:{location}")
                            if candidate.number_format != cell.number_format:
                                failures.append(f"date_format:{location}")
        for location in case["formulaChecks"]:
            sheet, address = location.split("!")
            if sheet not in actual or actual[sheet][address].data_type != "f":
                failures.append(f"missing_formula:{location}")
            elif case["kind"] == "report":
                tokens = Tokenizer(actual[sheet][address].value).items
                references = [t for t in tokens if t.subtype == "RANGE"]
                functions = {t.value.upper() for t in tokens if t.type == "FUNC"}
                required = "COUNTA(" if address == "B3" else "SUM("
                if not references or (sheet == "Summary" and required not in functions):
                    failures.append(f"formula_relationship:{location}")
        actual.close()
        golden.close()
        result["formulaEvaluation"] = "unavailable"
        if soffice:
            evaluated = recalculate(
                path.resolve(), (workspace.parent / "evaluation").resolve(), soffice
            )
            caches = formula_caches(evaluated)
            result["formulaCaches"] = caches
            evaluated_book = load_workbook(evaluated, data_only=True)
            for location, expected_value in case["formulaChecks"].items():
                sheet, address = location.split("!")
                value = (
                    evaluated_book[sheet][address].value
                    if sheet in evaluated_book
                    else None
                )
                if (
                    not isinstance(value, (int, float))
                    or abs(value - expected_value) > 1e-9
                ):
                    failures.append(f"formula_result:{location}")
            evaluated_book.close()
            if caches["missingCaches"] or caches["errorCaches"]:
                failures.append("formula_cache_errors")
            result["formulaEvaluation"] = "evaluated"
            if case["kind"] == "report" and not failures:
                failures.extend(check_changed_inputs(path, workspace, soffice))
        else:
            failures.append("formula_evaluation_unavailable")
    result["correct"] = not failures
    return result


def check_changed_inputs(path: Path, workspace: Path, soffice: Path) -> list[str]:
    """Execute changed inputs to reject formulas that only fit the original answers."""
    probe = workspace.parent / "relationship-probe.xlsx"
    book = load_workbook(path)
    data = book["Data"]
    expected = {}
    total = 0
    for row in range(2, data.max_row + 1):
        amount = row * 2
        data.cell(row, 3).value = amount
        total += amount
        expected[f"Data!D{row}"] = amount * 0.15
    data["A2"] = None
    expected.update(
        {
            "Summary!B2": total,
            "Summary!B3": data.max_row - 2,
            "Summary!B4": total * 0.15,
        }
    )
    book.save(probe)
    book.close()
    evaluated = recalculate(
        probe.resolve(),
        (workspace.parent / "relationship-evaluation").resolve(),
        soffice,
    )
    calculated = load_workbook(evaluated, data_only=True)
    failures = []
    for location, value in expected.items():
        sheet, address = location.split("!")
        actual = calculated[sheet][address].value
        if not isinstance(actual, (int, float)) or abs(actual - value) > 1e-7:
            failures.append(f"formula_relationship:{location}")
    calculated.close()
    return failures
