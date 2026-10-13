"""Deterministic synthetic workflows inspired by tabular finance operations."""

import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


def write_csv(
    path: Path, rows: list[list], encoding: str = "utf-8", delimiter: str = ","
) -> None:
    with path.open("w", encoding=encoding, newline="") as stream:
        csv.writer(stream, delimiter=delimiter).writerows(rows)


def build_cases(root: Path, detail_rows: int = 100_000) -> list[dict]:
    cases = []
    for name, builder in [
        ("csv_ingestion", _csv),
        ("reconciliation", _reconcile),
        ("formula_report", _report),
        ("targeted_edit", _edit),
        ("je_extraction", _je),
        ("bulk_analysis", _bulk),
    ]:
        folder = root / name
        inputs, golden = folder / "inputs", folder / "golden"
        inputs.mkdir(parents=True, exist_ok=True)
        golden.mkdir(exist_ok=True)
        case = builder(inputs, golden, detail_rows)
        cases.append({"id": name, "inputs": str(inputs), "golden": str(golden), **case})
    (root / "cases.json").write_text(json.dumps(cases, indent=2), encoding="utf-8")
    return cases


def _bulk(inputs: Path, golden: Path, detail_rows: int) -> dict:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Transactions")
    sheet.append(["identifier", "region", "amount", "date", "active"])
    totals: dict[str, list[int | str]] = {}
    for index in range(1, detail_rows + 1):
        region = ("east", "north", "west")[index % 3]
        identifier = f"{index:018d}"
        cents = index % 1000
        active = index % 2 == 0
        date = datetime(2026, 6, 30) if active else datetime(2026, 7, 1)
        sheet.append([identifier, region, cents / 100, date, active])
        if active:
            total = totals.setdefault(region, [0, 0, identifier])
            total[0] += 1
            total[1] += cents
    workbook.save(inputs / "transactions.xlsx")
    write_csv(
        golden / "output.csv",
        [
            ["region", "count", "total", "first_id"],
            *[
                [region, count, f"{cents / 100:.2f}", first]
                for region, (count, cents, first) in sorted(totals.items())
            ],
        ],
    )
    return {
        "output": "output.csv",
        "kind": "csv",
        "detailRows": detail_rows,
        "instruction": (
            "Analyze the complete Transactions table in transactions.xlsx. "
            "Select active=true rows dated 30 June 2026. Group by region and write "
            "UTF-8 output.csv sorted by region, columns exactly region,count,total,"
            "first_id. Count selected rows, sum amounts with two decimals, and "
            "use the smallest identifier string per region, preserving leading zeros. "
            "Verify total input and selected row counts and total amount. "
            "Keep the source unchanged; do not dump all input rows into chat."
        ),
    }


def _csv(inputs: Path, golden: Path, _: int) -> dict:
    rows = [["Identifier", "Amount", "Date", "Note"]]
    expected = [["identifier", "amount", "date"]]
    for i in range(1, 201):
        identifier = f"{i:018d}"
        rows.append(
            [
                identifier,
                "1,234.50" if i % 2 else "(25.50)",
                "30/06/2026",
                "quoted\nmultiline" if i == 1 else "unchanged",
            ]
        )
        expected.append([identifier, "1234.50" if i % 2 else "-25.50", "2026-06-30"])
    write_csv(inputs / "source.csv", rows, "utf-16", "\t")
    write_csv(golden / "output.csv", expected)
    return {
        "output": "output.csv",
        "kind": "csv",
        "instruction": "Normalize source.csv (UTF-16 with BOM, tab "
        "delimiter) to UTF-8 output.csv. "
        "Columns exactly identifier,amount,date, in input order. Preserve "
        "identifier strings "
        "and leading zeros. Parse comma thousands and parentheses negatives, "
        "emit amounts "
        "with two decimals; dates are DD/MM/YYYY and must become YYYY-MM-DD. "
        "Ignore Note. "
        "Verify row count and total amount; report them in the final answer.",
    }


def _reconcile(inputs: Path, golden: Path, _: int) -> dict:
    left = [
        ["reference", "amount"],
        ["0001", "12.50"],
        ["0002", "20.00"],
        ["0003", "5.00"],
    ]
    right = [
        ["reference", "amount"],
        ["0001", "12.50"],
        ["0002", "19.00"],
        ["0004", "8.00"],
    ]
    write_csv(inputs / "bank.csv", left, "utf-16", ";")
    write_csv(inputs / "ledger.csv", right)
    write_csv(
        golden / "output.csv",
        [
            ["reference", "status", "bank", "ledger", "difference"],
            ["0001", "matched", "12.50", "12.50", "0.00"],
            ["0002", "variance", "20.00", "19.00", "1.00"],
            ["0003", "bank_only", "5.00", "", "5.00"],
            ["0004", "ledger_only", "", "8.00", "-8.00"],
        ],
    )
    return {
        "output": "output.csv",
        "kind": "csv",
        "instruction": "Reconcile bank.csv (UTF-16, semicolon) against "
        "ledger.csv (UTF-8, comma), "
        "using unique reference and 0.01 amount tolerance. Write UTF-8 "
        "output.csv sorted "
        "by reference, columns reference,status,bank,ledger,difference. Preserve IDs. "
        "Statuses: matched,variance,bank_only,ledger_only. Missing amounts are blank; "
        "difference is bank minus ledger with missing side treated as zero "
        "for difference "
        "only. Amounts two decimals. Verify populations and control totals.",
    }


def _report(inputs: Path, golden: Path, _: int) -> dict:
    count = 1200
    rows = [["identifier", "date", "amount"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["identifier", "date", "amount", "reserve"])
    checks = {}
    for i in range(count):
        identifier = f"{i + 1:018d}"
        amount = (100, 250.5, -25.5)[i % 3]
        rows.append([identifier, "2026-06-30", f"{amount:.2f}"])
        row = i + 2
        ws.append([identifier, datetime(2026, 6, 30), amount, f"=C{row}*0.15"])
        ws.cell(row, 2).number_format = "yyyy-mm-dd"
        checks[f"Data!D{row}"] = amount * 0.15
    write_csv(inputs / "sales.csv", rows)
    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Total", f"=SUM(Data!C2:C{count + 1})"])
    summary.append(["Count", f"=COUNTA(Data!A2:A{count + 1})"])
    summary.append(["Reserve", f"=SUM(Data!D2:D{count + 1})"])
    checks.update({"Summary!B2": 130000, "Summary!B3": count, "Summary!B4": 19500})
    wb.save(golden / "output.xlsx")
    return {
        "output": "output.xlsx",
        "kind": "report",
        "formulaChecks": checks,
        "instruction": (
            "Create output.xlsx from sales.csv. Exactly Data and Summary sheets. "
            "Data columns: identifier,date,amount,reserve. Copy all source rows; "
            "identifiers stay strings, dates must be real Excel dates displayed "
            "yyyy-mm-dd, amounts numeric. Every reserve cell is a live formula "
            "multiplying that row's amount by 0.15. Summary A1:B4: Metric/Value; "
            "Total/<SUM of Data amounts>; Count/<COUNTA of Data identifiers>; "
            "Reserve/<SUM of Data reserves>. B2:B4 must be live formulas. "
            "Evaluate all formulas and check against independent totals. "
            "Keep input files unchanged."
        ),
    }


def _edit(inputs: Path, golden: Path, _: int) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.append(["Reserve", 0.1])
    ws.append(["Report date", datetime(2026, 6, 30)])
    ws["B1"].number_format = "0.0%"
    ws["B2"].number_format = "yyyy-mm-dd"
    ws["D5"] = "  retain this whitespace  "
    ws["D5"].comment = Comment("Reviewer note", "Synthetic")
    ws["D5"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["D6"] = "Reference"
    ws["D6"].hyperlink = "https://example.com/reference"
    ws.freeze_panes = "A3"
    dv = DataValidation(type="decimal", operator="between", formula1=0, formula2=1)
    ws.add_data_validation(dv)
    dv.add("B1")
    wb.create_sheet("Summary")["B2"] = "=Inputs!B1*1000"
    wb.create_sheet("Archive")["A1"] = "do not change"
    wb["Archive"].sheet_state = "hidden"
    wb.save(inputs / "input.xlsx")
    ws["B1"] = 0.15
    ws["B2"] = datetime(2026, 7, 31)
    wb.save(golden / "output.xlsx")
    return {
        "output": "output.xlsx",
        "kind": "edit",
        "formulaChecks": {"Summary!B2": 150},
        "instruction": "Create output.xlsx from input.xlsx, changing only Inputs!B1 to "
        "15% (numeric 0.15) and Inputs!B2 to 31 July 2026 (real Excel date). Preserve "
        "all other cells, formulas, formats, comments, hyperlinks, validation, hidden "
        "sheets and layout. Evaluate the Summary formula and independently verify its "
        "result. Preserve the original input; describe changes and checks.",
    }


def _je(inputs: Path, golden: Path, detail_rows: int) -> dict:
    path = inputs / "input.xlsx"
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Summary")
    for _ in range(12):
        ws.append([])
    ws.append(["account", "description", "debit", "credit"])
    entries = [
        ["000101", "Monthly", 120, 0],
        ["000102", "Annual", 80, 0],
        ["000103", "AR", 0, 200],
        ["000101", "Reserve", 12, 0],
        ["000102", "Reserve", 8, 0],
        ["000103", "Reserve AR", 0, 20],
    ]
    for entry in entries:
        ws.append(entry)
    data = wb.create_sheet("Detail")
    data.append([f"column_{i}" for i in range(10)])
    for row in range(1, detail_rows + 1):
        data.append([row, f"ID-{row:09d}"] + [row % 101] * 8)
    wb.save(path)
    write_csv(
        golden / "output.csv",
        [
            ["account", "description", "debit", "credit", "source"],
            *[
                [
                    *entry[:2],
                    f"{entry[2]:.2f}",
                    f"{entry[3]:.2f}",
                    f"Summary!A{row}:D{row}",
                ]
                for row, entry in enumerate(entries, 14)
            ],
        ],
    )
    return {
        "output": "output.csv",
        "kind": "csv",
        "detailRows": detail_rows,
        "instruction": "Extract the JE lines from input.xlsx Summary sheet to UTF-8 "
        "output.csv, columns exactly account,description,debit,credit,source. Preserve "
        "account strings; amounts two decimals; source is the exact A:D range for each "
        "row, e.g. Summary!A14:D14. Keep source order. Check debit and credit totals "
        "and report balance. Do not copy unrelated Detail rows.",
    }
