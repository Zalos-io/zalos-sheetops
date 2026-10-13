"""Fresh-process library comparisons; separate from model and tool-call latency."""

import argparse
import csv
import json
import platform
import resource
import subprocess
import sys
import time
from io import StringIO
from itertools import islice
from pathlib import Path


def csv_rows(path: Path, engine: str, limit: int | None) -> list[list]:
    if engine == "stdlib":
        with path.open(encoding="utf-16", newline="") as stream:
            reader = csv.reader(stream, delimiter="\t")
            next(reader)
            return list(islice(reader, limit)) if limit else list(reader)
    if engine == "polars":
        import polars as pl

        frame = pl.read_csv(
            StringIO(path.read_text(encoding="utf-16")),
            separator="\t",
            infer_schema=False,
        )
        return list((frame.head(limit) if limit else frame).iter_rows())
    from zalos_sheetops.io.read_csv_table import read_csv_table

    table = read_csv_table(path, encoding="auto", delimiter="\t", max_rows=limit)
    return [[row[name] for name in table["headers"]] for row in table["rows"]]


def xlsx_rows(path: Path, engine: str, deep: bool) -> list[list]:
    sheet, start, end = ("Detail", 99_991, 100_001) if deep else ("Summary", 14, 19)
    width = 10 if deep else 4
    if engine == "openpyxl":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        rows = list(
            wb[sheet].iter_rows(
                min_row=start, max_row=end, max_col=width, values_only=True
            )
        )
        wb.close()
        return rows
    if engine == "fastexcel":
        import fastexcel

        data = fastexcel.read_excel(path).load_sheet(
            sheet,
            header_row=0 if deep else 12,
            skip_rows=start - 2 if deep else None,
            n_rows=end - start + 1,
        )
        return [list(row.values()) for row in data.to_arrow().to_pylist()]
    from zalos_sheetops.io.read_cells import read_cells

    last = "J" if deep else "D"
    result = read_cells(path, sheet, f"A{start}:{last}{end}", ["values"], engine="xml")
    by_address = {cell["address"]: cell["value"] for cell in result["cells"]}
    return [
        [by_address.get(f"{chr(65 + column)}{row}") for column in range(width)]
        for row in range(start, end + 1)
    ]


def worker(path: Path, scenario: str, engine: str) -> dict:
    started = time.perf_counter()
    if scenario.startswith("csv"):
        rows = csv_rows(path, engine, 50 if scenario == "csv_preview" else None)
        signature = {"rows": len(rows), "first": rows[0], "last": rows[-1]}
        expected_count = 50 if scenario == "csv_preview" else 100_000
        correct = len(rows) == expected_count and rows[0][0] == "000000000000000001"
        correct = correct and rows[-1][0] == f"{expected_count:018d}"
        correct = correct and all(
            list(row)[1:] == ["12.50", "2026-06-30"] for row in rows
        )
    else:
        rows = xlsx_rows(path, engine, scenario == "deep_range")
        signature = {"rows": len(rows), "first": rows[0], "last": rows[-1]}
        if scenario == "je_range":
            expected = [
                ["000101", "Monthly", 120, 0],
                ["000102", "Annual", 80, 0],
                ["000103", "AR", 0, 200],
                ["000101", "Reserve", 12, 0],
                ["000102", "Reserve", 8, 0],
                ["000103", "Reserve AR", 0, 20],
            ]
        else:
            expected = [
                [row, f"ID-{row:09d}"] + [row % 101] * 8
                for row in range(99_990, 100_001)
            ]
        correct = [list(row) for row in rows] == expected
    seconds = time.perf_counter() - started
    rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return {
        "scenario": scenario,
        "engine": engine,
        "correct": correct,
        "seconds": seconds,
        "peakRssBytes": rss if sys.platform == "darwin" else rss * 1024,
        "signature": signature,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--scenario")
    parser.add_argument("--engine")
    args = parser.parse_args()
    csv_path = args.root / "operations.csv"
    xlsx_path = args.root / "fixtures" / "je_extraction" / "inputs" / "input.xlsx"
    if args.scenario:
        print(
            json.dumps(
                worker(
                    csv_path if args.scenario.startswith("csv") else xlsx_path,
                    args.scenario,
                    args.engine,
                )
            )
        )
        return
    with csv_path.open("w", encoding="utf-16", newline="") as stream:
        writer = csv.writer(stream, delimiter="\t")
        writer.writerow(["identifier", "amount", "date"])
        writer.writerows(
            [f"{i:018d}", "12.50", "2026-06-30"] for i in range(1, 100_001)
        )
    records = []
    for repeat in range(3):
        for scenario in ["csv_preview", "csv_full", "je_range", "deep_range"]:
            engines = (
                ["stdlib", "polars", "sheetops"]
                if scenario.startswith("csv")
                else ["openpyxl", "fastexcel", "sheetops"]
            )
            for engine in engines[repeat:] + engines[:repeat]:
                result = subprocess.run(
                    [
                        sys.executable,
                        "-m",
                        __spec__.name,
                        "--root",
                        str(args.root),
                        "--scenario",
                        scenario,
                        "--engine",
                        engine,
                    ],
                    capture_output=True,
                    text=True,
                )
                record = (
                    json.loads(result.stdout)
                    if result.returncode == 0
                    else {
                        "scenario": scenario,
                        "engine": engine,
                        "correct": False,
                        "error": result.stderr[-2000:],
                    }
                )
                records.append({"repeat": repeat, **record})
    output = {
        "platform": platform.platform(),
        "python": platform.python_version(),
        "scope": "fresh child process per operation; warm OS cache; import included",
        "records": records,
    }
    (args.root / "operations.json").write_text(json.dumps(output, indent=2))
    print(
        json.dumps(
            {"runs": len(records), "correct": sum(r["correct"] for r in records)}
        )
    )


if __name__ == "__main__":
    main()
