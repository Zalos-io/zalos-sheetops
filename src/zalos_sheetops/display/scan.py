from datetime import date, datetime, time
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from zalos_sheetops.display.widths import (
    column_width_to_pixels,
    required_cell_width_pixels,
)
from zalos_sheetops.display.xlsx_metadata import (
    WorksheetMetadata,
    read_column_widths,
    read_worksheet_metadata,
)
from zalos_sheetops.errors import SheetOpsError


def scan_display_overflow(
    path: Path,
    sheet_name: str | None = None,
    cell_range: str | None = None,
    max_results: int | None = None,
    max_cells: int | None = None,
    include_hidden: bool = False,
) -> dict[str, object]:
    bounds = _parse_cell_range(cell_range) if cell_range is not None else None
    metadata = _select_worksheets(
        read_worksheet_metadata(path),
        sheet_name,
        include_hidden,
    )
    value_workbook = _load_workbook(path, data_only=True)
    formula_workbook = _load_workbook(path, data_only=False)
    try:
        candidates: list[dict[str, object]] = []
        unknowns: list[dict[str, object]] = []
        candidate_count = 0
        unknown_count = 0
        scanned_cell_count = 0
        stopped_early = False
        for worksheet_metadata in metadata:
            value_worksheet = value_workbook[worksheet_metadata.name]
            formula_worksheet = formula_workbook[worksheet_metadata.name]
            column_widths = read_column_widths(path, worksheet_metadata.path)
            result = _scan_worksheet(
                worksheet_metadata.name,
                value_worksheet,
                formula_worksheet,
                column_widths,
                bounds,
                max_results,
                max_cells,
                scanned_cell_count,
                candidates,
                unknowns,
            )
            candidate_count += int(result["candidateCount"])
            unknown_count += int(result["unknownCount"])
            scanned_cell_count = int(result["scannedCellCount"])
            stopped_early = bool(result["stoppedEarly"])
            if stopped_early:
                break
    finally:
        value_workbook.close()
        formula_workbook.close()

    return {
        "scope": {"sheet": sheet_name, "range": cell_range},
        "candidateCount": candidate_count,
        "returnedCandidateCount": len(candidates),
        "unknownCount": unknown_count,
        "returnedUnknownCount": len(unknowns),
        "scannedCellCount": scanned_cell_count,
        "truncated": candidate_count > len(candidates) or unknown_count > len(unknowns),
        "stoppedEarly": stopped_early,
        "candidates": candidates,
        "unknowns": unknowns,
    }


def _scan_worksheet(
    sheet_name: str,
    value_worksheet: ReadOnlyWorksheet,
    formula_worksheet: ReadOnlyWorksheet,
    column_widths: dict[int, float],
    bounds: tuple[int, int, int, int] | None,
    max_results: int | None,
    max_cells: int | None,
    scanned_cell_count: int,
    candidates: list[dict[str, object]],
    unknowns: list[dict[str, object]],
) -> dict[str, object]:
    candidate_count = 0
    unknown_count = 0
    min_col, min_row, max_col, max_row = _iteration_bounds(bounds)
    value_rows = value_worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    )
    formula_rows = formula_worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    )
    for value_row, formula_row in zip_longest(value_rows, formula_rows, fillvalue=()):
        for value_cell, formula_cell in zip_longest(
            value_row,
            formula_row,
            fillvalue=EmptyCell(),
        ):
            if _is_empty_value_cell(value_cell) and not _is_formula_without_cache(
                formula_cell,
                value_cell,
            ):
                continue
            if max_cells is not None and scanned_cell_count >= max_cells:
                return {
                    "candidateCount": candidate_count,
                    "unknownCount": unknown_count,
                    "scannedCellCount": scanned_cell_count,
                    "stoppedEarly": True,
                }
            scanned_cell_count += 1
            unknown = _unknown_for_formula_without_cache(
                sheet_name, formula_cell, value_cell
            )
            if unknown is not None:
                unknown_count += 1
                if max_results is None or len(unknowns) < max_results:
                    unknowns.append(unknown)
                continue
            if not isinstance(value_cell, ReadOnlyCell):
                continue
            candidate = _candidate_for_cell(sheet_name, value_cell, column_widths)
            if candidate is None:
                continue
            candidate_count += 1
            if max_results is None or len(candidates) < max_results:
                candidates.append(candidate)
    return {
        "candidateCount": candidate_count,
        "unknownCount": unknown_count,
        "scannedCellCount": scanned_cell_count,
        "stoppedEarly": False,
    }


def _is_empty_value_cell(cell: ReadOnlyCell | EmptyCell) -> bool:
    return isinstance(cell, EmptyCell) or cell.value is None


def _is_formula_without_cache(
    formula_cell: ReadOnlyCell | EmptyCell,
    value_cell: ReadOnlyCell | EmptyCell,
) -> bool:
    if not isinstance(formula_cell, ReadOnlyCell):
        return False
    if not isinstance(formula_cell.value, str) or not formula_cell.value.startswith(
        "="
    ):
        return False
    return isinstance(value_cell, EmptyCell) or value_cell.value is None


def _unknown_for_formula_without_cache(
    sheet_name: str,
    formula_cell: ReadOnlyCell | EmptyCell,
    value_cell: ReadOnlyCell | EmptyCell,
) -> dict[str, object] | None:
    if not _is_formula_without_cache(formula_cell, value_cell):
        return None
    if not isinstance(formula_cell, ReadOnlyCell):
        return None
    return {
        "sheet": sheet_name,
        "cell": formula_cell.coordinate,
        "reason": "formula_without_cached_value",
    }


def _candidate_for_cell(
    sheet_name: str,
    cell: ReadOnlyCell,
    column_widths: dict[int, float],
) -> dict[str, object] | None:
    value_type = _display_value_type(cell)
    if value_type is None:
        return None
    display_text = _display_text(cell.value, value_type)
    required_width = required_cell_width_pixels(display_text)
    column_width = column_width_to_pixels(column_widths.get(cell.column))
    if required_width <= column_width:
        return None
    return {
        "sheet": sheet_name,
        "cell": cell.coordinate,
        "valueType": value_type,
        "displayText": display_text,
        "columnWidthPixels": column_width,
        "requiredWidthPixels": required_width,
        "issue": "display_overflow_candidate",
        "confidence": _confidence_for_value_type(value_type),
        "reason": _reason_for_value_type(value_type),
    }


def _display_value_type(cell: ReadOnlyCell) -> str | None:
    value = cell.value
    if isinstance(value, bool):
        return None
    if isinstance(value, datetime | date | time):
        return "date"
    if isinstance(value, int | float):
        if cell.is_date:
            return "date"
        return "number"
    return None


def _display_text(value: object, value_type: str) -> str:
    if value_type == "date":
        if isinstance(value, datetime):
            return value.strftime("%m/%d/%Y")
        if isinstance(value, date):
            return value.strftime("%m/%d/%Y")
        if isinstance(value, time):
            return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _confidence_for_value_type(value_type: str) -> str:
    if value_type == "number":
        return "high"
    return "medium"


def _reason_for_value_type(value_type: str) -> str:
    if value_type == "number":
        return "numeric_display_too_narrow"
    return f"{value_type}_display_too_narrow"


def _load_workbook(path: Path, data_only: bool) -> Workbook:
    try:
        return load_workbook(path, data_only=data_only, read_only=True)
    except Exception as error:
        raise SheetOpsError(
            "invalid_workbook",
            "Workbook could not be opened",
            {"path": str(path)},
        ) from error


def _select_worksheets(
    worksheets: list[WorksheetMetadata],
    sheet_name: str | None,
    include_hidden: bool,
) -> list[WorksheetMetadata]:
    if sheet_name is not None:
        for worksheet in worksheets:
            if worksheet.name == sheet_name:
                if worksheet.hidden and not include_hidden:
                    return []
                return [worksheet]
        raise SheetOpsError(
            "invalid_sheet",
            "Requested sheet does not exist",
            {"sheet": sheet_name},
        )
    return [
        worksheet for worksheet in worksheets if include_hidden or not worksheet.hidden
    ]


def _parse_cell_range(cell_range: str) -> tuple[int, int, int, int]:
    try:
        boundaries = range_boundaries(cell_range)
    except ValueError as error:
        raise SheetOpsError(
            "invalid_range",
            "Cell range is not valid A1 notation",
            {"range": cell_range},
        ) from error
    if any(value is None for value in boundaries):
        raise SheetOpsError(
            "invalid_range",
            "Cell range is not valid A1 notation",
            {"range": cell_range},
        )
    return boundaries


def _iteration_bounds(
    bounds: tuple[int, int, int, int] | None,
) -> tuple[int | None, int | None, int | None, int | None]:
    if bounds is None:
        return None, None, None, None
    min_col, min_row, max_col, max_row = bounds
    return min_col, min_row, max_col, max_row
