"""Runtime-owned edit validation; agent reports never authorize delivery."""

import hashlib
import json
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from zalos_sheetops.compare.edits import (
    cached_formula_issues,
    empty_string_results,
    formula_issues,
    parse_ranges,
    verify_edits,
)
from zalos_sheetops.delivery_cache import refreshed_formula_bytes


class DeliveryPolicy:
    """An immutable input, frozen edit contract, and byte-bound approval.

    The owner must keep this object and its temporary directory outside the
    agent's sandbox and check ``is_approved`` on the bytes actually served.
    Scope correctness relative to the user's request remains a separate check.
    """

    def __init__(self, original: bytes, *, engine: str | None = None) -> None:
        self._temporary = tempfile.TemporaryDirectory(prefix="sheetops-delivery-")
        self._root = Path(self._temporary.name)
        self._original = self._root / "original.xlsx"
        self._original.write_bytes(original)
        self._scope: dict | None = None
        self._receipt: str | None = None
        self._engine = engine
        self.validated_bytes = b""

    def __enter__(self) -> "DeliveryPolicy":
        return self

    def __exit__(self, *args: object) -> None:
        self.close()

    def close(self) -> None:
        self._temporary.cleanup()

    def declare(self, allowed_ranges: object) -> None:
        parse_ranges(allowed_ranges)
        declaration = json.loads(json.dumps(allowed_ranges, sort_keys=True))
        if self._scope is not None and declaration != self._scope:
            raise ValueError("Edit scope is frozen and cannot be widened")
        self._scope = declaration

    def is_approved(self, output: bytes) -> bool:
        return (
            self._receipt is not None
            and self._receipt == hashlib.sha256(output).hexdigest()
        )

    def _block(self, code: str, message: str, **details: object) -> dict:
        self._receipt = None
        self.validated_bytes = b""
        return {
            "action": "block",
            "errorCode": code,
            "message": message,
            "details": details,
            "repairable": code in {"invalid_output", "invalid_formula_results"},
        }

    def _evaluate(self, actual: Path) -> Path:
        destination = self._root / "evaluated"
        destination.mkdir(exist_ok=True)
        produced = destination / actual.name
        produced.unlink(missing_ok=True)
        process = subprocess.run(
            [
                self._engine,
                "-env:UserInstallation=" + (self._root / "profile").as_uri(),
                "--headless",
                "--calc",
                "--convert-to",
                "xlsx:Calc MS Excel 2007 XML",
                "--outdir",
                str(destination),
                str(actual),
            ],
            capture_output=True,
            timeout=60,
            check=False,
        )
        if process.returncode != 0 or not produced.is_file():
            raise ValueError("Calculation engine did not produce a workbook")
        return produced

    def _validate_formulas(self, actual: Path, evaluated: Path) -> list[dict]:
        formulas = load_workbook(actual, data_only=False)
        evaluated_formulas = load_workbook(evaluated, data_only=False)
        cached = load_workbook(evaluated, data_only=True)
        try:
            blank = empty_string_results(evaluated)
            issues = []
            for sheet in formulas:
                errors = (
                    [
                        cell
                        for cell in cached[sheet.title]._cells.values()
                        if cell.data_type == "e"
                    ]
                    if sheet.title in cached.sheetnames
                    else []
                )
                for cell in sheet._cells.values():
                    if cell.data_type == "f":
                        issues.extend(
                            formula_issues(sheet.title, cell, evaluated_formulas)
                        )
                        issues.extend(
                            cached_formula_issues(
                                sheet.title, cell, cached, blank, errors
                            )
                        )
            return issues
        finally:
            formulas.close()
            evaluated_formulas.close()
            cached.close()

    def validate(self, output: bytes) -> dict:
        self._receipt = None
        self.validated_bytes = b""
        if self._scope is None:
            return self._block(
                "missing_contract",
                "Declare the edit scope before modifying the workbook.",
            )
        actual = self._root / "candidate.xlsx"
        actual.write_bytes(output)
        try:
            workbook = load_workbook(actual, data_only=False)
            try:
                has_formulas = any(
                    cell.data_type == "f"
                    for sheet in workbook
                    for cell in sheet._cells.values()
                )
            finally:
                workbook.close()
            evaluated = None
            if has_formulas:
                if not self._engine:
                    return self._block(
                        "calculation_engine_unavailable",
                        "Formula delivery requires an available calculation engine.",
                    )
                evaluated = self._evaluate(actual)
                issues = self._validate_formulas(actual, evaluated)
                if issues:
                    return self._block(
                        "invalid_formula_results",
                        "Formula evaluation rejected the workbook.",
                        issues=issues[:50],
                    )
            result = verify_edits(
                self._original, actual, self._scope, evaluated_path=evaluated
            )
            if not result["verified"]:
                return self._block(
                    "invalid_output",
                    "Workbook changes violate the frozen edit scope.",
                    validation=result,
                )
        except Exception as error:
            return self._block(
                "validation_failed",
                "Workbook validation failed.",
                errorType=type(error).__name__,
            )
        try:
            self.validated_bytes = (
                refreshed_formula_bytes(output, evaluated.read_bytes())
                if evaluated
                else output
            )
        except Exception as error:
            return self._block(
                "cache_refresh_failed",
                "Could not safely refresh formula caches.",
                errorType=type(error).__name__,
            )
        self._receipt = hashlib.sha256(self.validated_bytes).hexdigest()
        return {
            "action": "allow",
            "sha256": self._receipt,
            "formulaEvaluation": "engine" if evaluated else "not_required",
        }
