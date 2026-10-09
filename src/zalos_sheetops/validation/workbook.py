from pathlib import Path
from typing import TYPE_CHECKING

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.xlsx_xml import RowData, XlsxWorkbook
from zalos_sheetops.normalize.text import canonical_header

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet


def validate_workbook(
    path: Path,
    required_sheets: list[str],
    required_columns: dict[str, list[str]],
    unique_keys: dict[str, list[str]],
) -> dict[str, object]:
    with XlsxWorkbook(path) as workbook:
        sheet_names = set(workbook.sheet_names())
        issues: list[dict[str, object]] = []

        for sheet_name in required_sheets:
            if sheet_name not in sheet_names:
                issues.append({"code": "missing_required_sheet", "sheet": sheet_name})

        for sheet_name in sorted(set(required_columns) | set(unique_keys)):
            if sheet_name not in sheet_names:
                continue
            columns = required_columns.get(sheet_name, [])
            key_columns = unique_keys.get(sheet_name, [])
            _validate_table_sheet_xml(
                workbook, sheet_name, columns, key_columns, issues
            )

    return {"valid": not issues, "issues": issues, "uniqueKeysChecked": unique_keys}


def validate_workbook_rules(
    path: Path,
    sheet_rules: list[dict[str, object]],
) -> dict[str, object]:
    with XlsxWorkbook(path) as xml_workbook:
        sheet_names = set(xml_workbook.sheet_names())
        workbook: Workbook | None = None
        issues: list[dict[str, object]] = []
        checked_sheets: list[str] = []

        for rule in sheet_rules:
            sheet_name = str(rule.get("name", ""))
            kind = str(rule.get("kind", "table")).lower()
            checked_sheets.append(sheet_name)
            if sheet_name not in sheet_names:
                issues.append({"code": "missing_required_sheet", "sheet": sheet_name})
                continue
            if kind == "table":
                _validate_table_rule_xml(xml_workbook, sheet_name, rule, issues)
                continue
            if kind in {"summary", "freeform"}:
                if workbook is None:
                    workbook = _load_formula_workbook(path)
                sheet = workbook[sheet_name]
                _validate_freeform_rule(sheet, rule, issues)
                continue
            raise SheetOpsError(
                "invalid_sheet_validation_kind",
                "Sheet validation kind is not supported",
                {"sheet": sheet_name, "kind": kind},
            )

    return {"valid": not issues, "issues": issues, "sheetsChecked": checked_sheets}


def _load_formula_workbook(path: Path) -> "Workbook":
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=False, read_only=False)
    except Exception as error:
        raise SheetOpsError(
            "invalid_workbook",
            "Workbook could not be opened",
            {"path": str(path)},
        ) from error


def _validate_table_sheet_xml(
    workbook: XlsxWorkbook,
    sheet_name: str,
    columns: list[str],
    key_columns: list[str],
    issues: list[dict[str, object]],
) -> None:
    columns_to_find = _unique_strings(columns + key_columns)
    header_row, row_numbers_by_key = _scan_table_sheet_xml(
        workbook,
        sheet_name,
        columns_to_find,
        key_columns,
    )
    headers = _normalize_headers(header_row.values)
    header_set = set(headers)
    missing = [
        column for column in columns if canonical_header(column) not in header_set
    ]
    if missing:
        issues.append(
            {
                "code": "missing_required_columns",
                "sheet": sheet_name,
                "missing": missing,
            }
        )
    missing_key_columns = [
        column for column in key_columns if canonical_header(column) not in header_set
    ]
    if missing_key_columns:
        issues.append(
            {
                "code": "missing_unique_key_columns",
                "sheet": sheet_name,
                "missing": missing_key_columns,
            }
        )
        return
    if key_columns:
        if row_numbers_by_key is None:
            row_numbers_by_key = _collect_unique_key_rows_xml(
                workbook,
                sheet_name,
                header_row.number,
                headers,
                key_columns,
            )
        duplicate_keys = _duplicate_unique_key_payloads(
            row_numbers_by_key,
            key_columns,
        )
        if duplicate_keys:
            issues.append(
                {
                    "code": "duplicate_unique_key",
                    "sheet": sheet_name,
                    "keyColumns": key_columns,
                    "duplicates": duplicate_keys,
                }
            )


def _validate_table_rule_xml(
    workbook: XlsxWorkbook,
    sheet_name: str,
    rule: dict[str, object],
    issues: list[dict[str, object]],
) -> None:
    columns = _string_list_rule(rule, "requiredColumns")
    key_columns = _string_list_rule(rule, "uniqueKeys")
    if not columns and not key_columns:
        return
    _validate_table_sheet_xml(workbook, sheet_name, columns, key_columns, issues)


def _scan_table_sheet_xml(
    workbook: XlsxWorkbook,
    sheet_name: str,
    required_columns: list[str],
    key_columns: list[str],
) -> tuple[RowData, dict[str | tuple[str, ...], int | list[int]] | None]:
    expected = {canonical_header(column) for column in required_columns}
    canonical_key_columns = [canonical_header(column) for column in key_columns]
    candidate_rows: list[int] = []
    best_row: RowData | None = None
    best_score = -1
    selected_row: RowData | None = None
    selected_header_indexes: dict[str, int] = {}
    row_numbers_by_key: dict[str | tuple[str, ...], int | list[int]] | None = (
        {} if key_columns else None
    )

    for row in workbook.iter_rows(sheet_name):
        non_blank = set()
        for value in row.values:
            if value is None:
                continue
            header = canonical_header(value)
            if header:
                non_blank.add(header)
        if not non_blank:
            continue

        overlap = 0
        for header in expected:
            if header in non_blank:
                overlap += 1
        if expected:
            if overlap == len(expected):
                candidate_rows.append(row.number)
                if selected_row is None:
                    selected_row = row
                    selected_header_indexes = {}
                    for index, value in enumerate(row.values):
                        header = canonical_header(value)
                        if header:
                            selected_header_indexes[header] = index
            if (
                selected_row is not None
                and row.number > selected_row.number
                and row_numbers_by_key is not None
                and _has_key_columns(selected_header_indexes, canonical_key_columns)
            ):
                _record_unique_key_row(
                    row.values,
                    row.number,
                    selected_header_indexes,
                    canonical_key_columns,
                    row_numbers_by_key,
                )
            if overlap == 0:
                continue

        score = overlap * 100 + len(non_blank)
        if score > best_score:
            best_score = score
            best_row = row

    if len(candidate_rows) > 1:
        raise SheetOpsError(
            "ambiguous_header_row",
            "Multiple header rows match the required columns",
            {
                "sheet": sheet_name,
                "headerRows": candidate_rows,
                "requiredColumns": required_columns,
            },
        )
    if selected_row is not None:
        return selected_row, row_numbers_by_key
    if best_row is None:
        raise SheetOpsError(
            "header_not_found",
            "Unable to detect header row",
            {"sheet": sheet_name, "headerStrategy": "detect"},
        )
    return best_row, None


def _collect_unique_key_rows_xml(
    workbook: XlsxWorkbook,
    sheet_name: str,
    header_row_number: int,
    headers: list[str],
    key_columns: list[str],
) -> dict[str | tuple[str, ...], int | list[int]]:
    header_indexes = {header: index for index, header in enumerate(headers) if header}
    canonical_key_columns = [canonical_header(column) for column in key_columns]
    row_numbers_by_key: dict[str | tuple[str, ...], int | list[int]] = {}

    for row in workbook.iter_rows(sheet_name):
        if row.number <= header_row_number:
            continue
        _record_unique_key_row(
            row.values,
            row.number,
            header_indexes,
            canonical_key_columns,
            row_numbers_by_key,
        )

    return row_numbers_by_key


def _has_key_columns(
    header_indexes: dict[str, int],
    canonical_key_columns: list[str],
) -> bool:
    return all(column in header_indexes for column in canonical_key_columns)


def _record_unique_key_row(
    row: list[object],
    row_number: int,
    header_indexes: dict[str, int],
    canonical_key_columns: list[str],
    row_numbers_by_key: dict[str | tuple[str, ...], int | list[int]],
) -> None:
    for value in row:
        if value is None:
            continue
        if type(value) is str:
            if value.strip():
                break
        else:
            # Numbers and booleans always render as non-blank text.
            break
    else:
        return
    key: str | tuple[str, ...]
    if len(canonical_key_columns) == 1:
        key = _cell_text(row, header_indexes[canonical_key_columns[0]])
    else:
        key = tuple(
            _cell_text(row, header_indexes[column]) for column in canonical_key_columns
        )
    existing = row_numbers_by_key.get(key)
    if existing is None:
        # Store the first occurrence as a bare int; duplicate keys are rare
        # and a list per unique key doubles scan memory.
        row_numbers_by_key[key] = row_number
    elif type(existing) is int:
        row_numbers_by_key[key] = [existing, row_number]
    else:
        existing.append(row_number)


def _duplicate_unique_key_payloads(
    row_numbers_by_key: dict[str | tuple[str, ...], int | list[int]],
    key_columns: list[str],
) -> list[dict[str, object]]:
    duplicates = []
    single_key = len(key_columns) == 1
    for key, row_numbers in row_numbers_by_key.items():
        if type(row_numbers) is int:
            continue
        key_values = (key,) if single_key else key
        duplicates.append(
            {
                "key": {
                    original_column: key_values[index]
                    for index, original_column in enumerate(key_columns)
                },
                "rows": row_numbers,
            }
        )
    return duplicates


def _validate_freeform_rule(
    sheet: "Worksheet",
    rule: dict[str, object],
    issues: list[dict[str, object]],
) -> None:
    for required_cell in _dict_list_rule(rule, "requiredCells"):
        cell_ref = str(required_cell.get("cell", ""))
        if "equals" not in required_cell:
            continue
        expected = _value_text(required_cell.get("equals"))
        actual = _value_text(sheet[cell_ref].value)
        if actual != expected:
            issues.append(
                {
                    "code": "cell_value_mismatch",
                    "sheet": sheet.title,
                    "cell": cell_ref,
                    "expected": expected,
                    "actual": actual,
                }
            )

    for required_formula in _dict_list_rule(rule, "requiredFormulas"):
        cell_ref = str(required_formula.get("cell", ""))
        value = sheet[cell_ref].value
        if not isinstance(value, str) or not value.startswith("="):
            issues.append(
                {
                    "code": "missing_required_formula",
                    "sheet": sheet.title,
                    "cell": cell_ref,
                    "value": _value_text(value),
                }
            )
            continue
        contains = required_formula.get("contains")
        if contains is not None and str(contains) not in value:
            issues.append(
                {
                    "code": "formula_mismatch",
                    "sheet": sheet.title,
                    "cell": cell_ref,
                    "contains": str(contains),
                    "formula": value,
                }
            )


def _string_list_rule(rule: dict[str, object], field_name: str) -> list[str]:
    value = rule.get(field_name, [])
    if not isinstance(value, list):
        return []
    return [str(item) for item in value if isinstance(item, str)]


def _dict_list_rule(
    rule: dict[str, object],
    field_name: str,
) -> list[dict[str, object]]:
    value = rule.get(field_name, [])
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _value_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _normalize_headers(header_row: list[object]) -> list[str]:
    headers = [canonical_header(value) for value in header_row]
    duplicates = _duplicate_headers(headers)
    blank_columns = [index + 1 for index, header in enumerate(headers) if not header]
    if duplicates or blank_columns:
        raise SheetOpsError(
            "invalid_headers",
            "Header row contains duplicate or blank columns",
            {"duplicates": duplicates, "blankColumns": blank_columns},
        )
    return headers


def _duplicate_headers(headers: list[str]) -> list[str]:
    seen: set[str] = set()
    duplicates = []
    for header in headers:
        if not header:
            continue
        if header in seen and header not in duplicates:
            duplicates.append(header)
            continue
        seen.add(header)
    return duplicates


def _cell_text(row: list[object], index: int) -> str:
    if index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    return str(value)


def _unique_strings(values: list[str]) -> list[str]:
    seen: set[str] = set()
    unique = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        unique.append(value)
    return unique
