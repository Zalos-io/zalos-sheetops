"""Transplant evaluated caches without rewriting unrelated OOXML parts."""

import io
import re
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

from zalos_sheetops.io.xlsx_xml import MAIN_NS, XlsxWorkbook
from zalos_sheetops.recalc.fresh_input import without_formula_caches

_CELL = re.compile(rb"<(?P<p>[\w.-]+:)?c\b[^>]*(?:/>|>.*?</(?:[\w.-]+:)?c>)", re.S)
_REF = re.compile(rb'\br=["\']([^"\']+)["\']')
_TYPE = re.compile(rb'\s+t=["\'][^"\']*["\']')
_VALUE = re.compile(
    rb"<(?:[\w.-]+:)?(?:v|is)\b[^>]*(?:/>|>.*?</(?:[\w.-]+:)?(?:v|is)>)", re.S
)


def fresh_calculation_input(source: Path, destination: Path) -> None:
    """Clear only formula caches in a disposable copy before calculation."""

    with XlsxWorkbook(source) as workbook:
        worksheet_parts = {sheet.path for sheet in workbook.worksheets}
    with ZipFile(source) as original, ZipFile(destination, "w") as staged:
        staged.comment = original.comment
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename in worksheet_parts:
                data = without_formula_caches(data)
            staged.writestr(item, data)


def _result_cell(original: bytes, value: str, kind: str) -> bytes:
    opening, _, tail = original.partition(b">")
    prefix = re.match(rb"<([\w.-]+:)?c\b", opening).group(1) or b""
    opening = _TYPE.sub(b"", opening.rstrip(b"/"))
    opening += b' t="' + kind.encode() + b'">'
    tail = _VALUE.sub(b"", tail)
    if not tail:
        tail = b"</" + prefix + b"c>"
    cache = b"<" + prefix + b"v>" + escape(value).encode() + b"</" + prefix + b"v>"
    close = tail.rfind(b"</")
    return opening + tail[:close] + cache + tail[close:]


def _refresh_sheet(xml: bytes, values: dict[str, tuple[str, str]]) -> bytes:
    pending = dict(values)

    def replace(match: re.Match) -> bytes:
        cell = match.group()
        reference = _REF.search(cell.split(b">", 1)[0])
        coordinate = reference.group(1).decode() if reference else ""
        if coordinate not in pending:
            return cell
        value, kind = pending.pop(coordinate)
        return _result_cell(cell, value, kind)

    result = _CELL.sub(replace, xml)
    # Missing spill cells require new cell nodes in their existing rows. Refuse
    # unfamiliar layouts instead of rewriting a worksheet with lost extensions.
    by_row: dict[int, list[tuple[str, tuple[str, str]]]] = {}
    for coordinate, value in pending.items():
        row = int(re.search(r"\d+$", coordinate).group())
        by_row.setdefault(row, []).append((coordinate, value))
    for row, cells in by_row.items():
        pattern = re.compile(
            rb'<(?P<p>[\w.-]+:)?row\b(?P<a>[^>]*\br=["\']'
            + str(row).encode()
            + rb'["\'][^>]*)(?:/>|>(?P<body>.*?)</(?:[\w.-]+:)?row>)',
            re.S,
        )
        match = pattern.search(result)
        if match is None:
            raise ValueError("Evaluated spill requires a missing worksheet row")
        prefix = match.group("p") or b""
        body = match.group("body") or b""
        for coordinate, (value, kind) in cells:
            raw = b"<" + prefix + b'c r="' + coordinate.encode() + b'"/>'
            body += _result_cell(raw, value, kind)
        # Cell order matters to streaming readers; sort only this result row.
        nodes = list(_CELL.finditer(body))
        if nodes and b"".join(node.group() for node in nodes) != body:
            raise ValueError("Cannot safely insert a spill into this row layout")
        nodes.sort(
            key=lambda node: range_boundaries(
                _REF.search(node.group()).group(1).decode()
            )[0]
        )
        replacement = b"<" + prefix + b"row" + match.group("a") + b">"
        replacement += (
            b"".join(node.group() for node in nodes) + b"</" + prefix + b"row>"
        )
        result = result[: match.start()] + replacement + result[match.end() :]
    return result


def refreshed_formula_bytes(actual: bytes, evaluated: bytes) -> bytes:
    """Return original package bytes with only formula result caches refreshed."""
    formula_book = load_workbook(io.BytesIO(actual), data_only=False)
    value_book = load_workbook(io.BytesIO(evaluated), data_only=True)
    replacements = {}
    try:
        if formula_book.epoch != value_book.epoch:
            raise ValueError("Calculation changed the workbook date system")
        with (
            XlsxWorkbook(io.BytesIO(actual)) as original_xml,
            XlsxWorkbook(io.BytesIO(evaluated)) as evaluated_xml,
        ):
            with (
                ZipFile(io.BytesIO(actual)) as original_zip,
                ZipFile(io.BytesIO(evaluated)) as evaluated_zip,
            ):
                for sheet in formula_book:
                    coordinates = set()
                    for cell in sheet._cells.values():
                        if cell.data_type != "f":
                            continue
                        coordinates.add(cell.coordinate)
                        if isinstance(cell.value, ArrayFormula):
                            left, top, right, bottom = range_boundaries(cell.value.ref)
                            for row in sheet.iter_rows(
                                min_row=top, max_row=bottom, min_col=left, max_col=right
                            ):
                                coordinates.update(item.coordinate for item in row)
                    if not coordinates:
                        continue
                    source_part = original_xml.worksheet(sheet.title).path
                    evaluated_part = evaluated_xml.worksheet(sheet.title).path
                    tree = ET.fromstring(evaluated_zip.read(evaluated_part))
                    raw_cells = {
                        cell.get("r"): cell for cell in tree.iter(f"{{{MAIN_NS}}}c")
                    }
                    values = {}
                    for coordinate in coordinates:
                        cell = raw_cells.get(coordinate)
                        if cell is None:
                            raise ValueError("Evaluated formula result cell is missing")
                        kind = cell.get("t", "n")
                        if kind in {"s", "inlineStr", "str"}:
                            value = value_book[sheet.title][coordinate].value
                            values[coordinate] = (
                                "" if value is None else str(value),
                                "str",
                            )
                        else:
                            node = cell.find(f"{{{MAIN_NS}}}v")
                            if node is None or node.text is None:
                                raise ValueError(
                                    "Evaluated formula result value is missing"
                                )
                            values[coordinate] = (node.text, kind)
                    replacements[source_part] = _refresh_sheet(
                        original_zip.read(source_part), values
                    )
                output = io.BytesIO()
                with ZipFile(output, "w") as destination:
                    destination.comment = original_zip.comment
                    for item in original_zip.infolist():
                        destination.writestr(
                            item,
                            replacements.get(
                                item.filename, original_zip.read(item.filename)
                            ),
                        )
                return output.getvalue()
    finally:
        formula_book.close()
        value_book.close()
