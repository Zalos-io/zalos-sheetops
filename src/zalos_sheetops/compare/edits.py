"""Check edit scope and results in a separately evaluated formula copy.

This is a verifier, not a formula evaluator. The caller must evaluate the copy
in Excel or LibreOffice; a saved cache alone does not establish freshness.
"""

from copy import copy
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.etree.ElementTree import tostring
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.xlsx_xml import MAIN_NS, XlsxWorkbook

FUTURE_FUNCTIONS = {
    "TEXTJOIN",
    "CONCAT",
    "FILTER",
    "UNIQUE",
    "SORT",
    "SORTBY",
    "XLOOKUP",
    "XMATCH",
    "LET",
    "LAMBDA",
    "SEQUENCE",
    "TEXTSPLIT",
    "HSTACK",
    "VSTACK",
    "TAKE",
    "DROP",
    "CHOOSECOLS",
    "CHOOSEROWS",
    "BYROW",
    "BYCOL",
    "MAP",
    "REDUCE",
    "SCAN",
    "MAKEARRAY",
    "RANDARRAY",
    "TOCOL",
    "TOROW",
    "WRAPROWS",
    "WRAPCOLS",
    "EXPAND",
}
CELL_STYLE_FIELDS = (
    "font",
    "fill",
    "border",
    "alignment",
    "protection",
    "number_format",
)
DYNAMIC_FUNCTIONS = {
    "FILTER",
    "UNIQUE",
    "SORT",
    "SORTBY",
    "SEQUENCE",
    "TEXTSPLIT",
    "HSTACK",
    "VSTACK",
    "TAKE",
    "DROP",
    "CHOOSECOLS",
    "CHOOSEROWS",
    "BYROW",
    "BYCOL",
    "MAP",
    "REDUCE",
    "SCAN",
    "MAKEARRAY",
    "RANDARRAY",
    "TOCOL",
    "TOROW",
    "WRAPROWS",
    "WRAPCOLS",
    "EXPAND",
}


def parse_ranges(ranges: object) -> dict[str, list[tuple[int, int, int, int]]]:
    if not isinstance(ranges, dict):
        raise SheetOpsError(
            "invalid_field",
            "allowedRanges must map sheets to bounded ranges",
            {"field": "allowedRanges"},
        )
    parsed = {}
    for sheet, areas in ranges.items():
        if not isinstance(sheet, str) or not isinstance(areas, list):
            raise SheetOpsError(
                "invalid_field", "Invalid allowedRanges", {"field": "allowedRanges"}
            )
        bounds = []
        for area in areas:
            try:
                box = range_boundaries(area)
                if None in box or box[0] > box[2] or box[1] > box[3] or min(box) < 1:
                    raise ValueError(area)
            except (TypeError, ValueError) as error:
                raise SheetOpsError(
                    "invalid_range",
                    "Edit scope must use bounded A1 ranges",
                    {"sheet": sheet, "range": area},
                ) from error
            bounds.append(box)
        parsed[sheet] = bounds
    return parsed


def in_scope(sheet: str, cell: object, scopes: dict) -> bool:
    return any(
        left <= cell.column <= right and top <= cell.row <= bottom
        for left, top, right, bottom in scopes.get(sheet, [])
    )


def formula_text(cell: object) -> str | None:
    if cell is None or cell.data_type != "f":
        return None
    return cell.value if isinstance(cell.value, str) else cell.value.text


def cell_signature(cell: object) -> tuple:
    if cell is None:
        return (None,)
    value = formula_text(cell) or cell.value
    if cell.data_type == "f" and not isinstance(cell.value, str):
        value = (value, cell.value.ref)
    if value is None and not cell.has_style and not cell.comment and not cell.hyperlink:
        return (None,)
    note = (cell.comment.text, cell.comment.author) if cell.comment else None
    link = (
        (
            cell.hyperlink.target,
            cell.hyperlink.location,
            cell.hyperlink.tooltip,
            cell.hyperlink.display,
        )
        if cell.hyperlink
        else None
    )
    return (
        value,
        cell.data_type,
        *(copy(getattr(cell, field)) for field in CELL_STYLE_FIELDS),
        note,
        link,
    )


def sheet_structure(sheet: object) -> tuple:
    return (
        sheet.sheet_state,
        tuple(str(area) for area in sheet.merged_cells.ranges),
        str(sheet.freeze_panes),
        xml_signature(sheet.auto_filter),
        xml_signature(sheet.data_validations),
        xml_signature(sheet.sheet_properties),
        xml_signature(sheet.sheet_format),
        xml_signature(sheet.protection),
        xml_signature(sheet.print_options),
        xml_signature(sheet.page_margins),
        xml_signature(sheet.page_setup),
        tuple(
            (
                str(key.sqref),
                tuple(
                    xml_signature(rule) for rule in sheet.conditional_formatting[key]
                ),
            )
            for key in sheet.conditional_formatting
        ),
        str(sheet.print_area),
        tuple(
            (key, xml_signature(value)) for key, value in sheet.row_dimensions.items()
        ),
        tuple(
            (key, xml_signature(value))
            for key, value in sheet.column_dimensions.items()
        ),
        tuple((table.name, xml_signature(table)) for table in sheet.tables.values()),
    )


def xml_signature(value: object) -> bytes:
    return tostring(value.to_tree())


def normalized_formula(text: str | None) -> str:
    # Prefixes describe function storage, not string literals or sheet names.
    tokens = Tokenizer(text or "=").items
    parts = []
    index = 0
    while index < len(tokens):
        token = tokens[index]
        value = token.value
        if token.type == "FUNC":
            value = value.replace("_xlfn.", "").replace("_xlws.", "").upper()
            # LibreOffice serializes Excel's boolean constants as zero-arg
            # functions. Keep quoted text and reference operands untouched.
            if (
                value in {"TRUE(", "FALSE("}
                and index + 1 < len(tokens)
                and tokens[index + 1].subtype == "CLOSE"
            ):
                parts.append(value[:-1])
                index += 2
                continue
        parts.append(value)
        index += 1
    return "".join(parts)


def protected_parts(path: Path) -> dict[str, bytes]:
    with ZipFile(path) as archive:
        return {
            name: archive.read(name)
            for name in archive.namelist()
            if name.startswith(
                (
                    "xl/media/",
                    "xl/charts/",
                    "xl/drawings/",
                    "xl/embeddings/",
                    "xl/activeX/",
                    "customXml/",
                )
            )
            or "vbaProject" in name
        }


def empty_string_results(path: Path) -> set[tuple[str, str]]:
    results = set()
    with XlsxWorkbook(path) as book:
        for sheet in book.worksheets:
            with book.archive.open(sheet.path) as stream:
                for _, cell in ET.iterparse(stream, events=("end",)):
                    if cell.tag != f"{{{MAIN_NS}}}c":
                        continue
                    value = cell.find(f"{{{MAIN_NS}}}v")
                    if (
                        cell.get("t") == "str"
                        and value is not None
                        and value.text in (None, "")
                    ):
                        results.add((sheet.name, cell.get("r")))
                    cell.clear()
    return results


def formula_issues(sheet: str, cell: object, evaluated: object) -> list[dict]:
    formula = formula_text(cell)
    reasons = []
    try:
        functions = [
            token.value[:-1]
            for token in Tokenizer(formula).items
            if token.type == "FUNC" and token.subtype == "OPEN"
        ]
    except TokenizerError as error:
        return [
            {
                "sheet": sheet,
                "cell": cell.coordinate,
                "reason": "invalid_formula_syntax",
                "message": str(error),
            }
        ]
    if any(name.upper() in FUTURE_FUNCTIONS for name in functions):
        reasons.append("missing_function_prefix")
    if any(
        name.upper().split(".")[-1] in DYNAMIC_FUNCTIONS for name in functions
    ) and not isinstance(cell.value, ArrayFormula):
        reasons.append("array_storage_required")
    if evaluated is None:
        reasons.append("evaluation_required")
    elif sheet not in evaluated.sheetnames:
        reasons.append("evaluation_sheet_missing")
    else:
        evaluated_cell = evaluated[sheet]._cells.get((cell.row, cell.column))
        other_formula = formula_text(evaluated_cell)
        if normalized_formula(formula) != normalized_formula(other_formula):
            reasons.append("formula_mismatch")
    return [
        {"sheet": sheet, "cell": cell.coordinate, "reason": reason}
        for reason in reasons
    ]


def verify_edits(
    original: Path,
    actual: Path,
    allowed_ranges: object,
    *,
    evaluated_path: Path | None = None,
    structural_sheets: list[str] | None = None,
) -> dict:
    scopes = parse_ranges(allowed_ranges)
    structural = set(structural_sheets or [])
    if not structural <= scopes.keys():
        raise SheetOpsError(
            "invalid_field",
            "Structural sheets must have an explicit edit scope",
            {"field": "structuralSheets"},
        )
    opened = []
    try:
        for path in [original, actual, *([evaluated_path] if evaluated_path else [])]:
            opened.append(
                load_workbook(
                    path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm"
                )
            )
        before, after = opened[:2]
        if not scopes.keys() <= set(before.sheetnames):
            raise SheetOpsError(
                "invalid_sheet",
                "Edit scope contains an unknown original sheet",
                {"sheets": sorted(scopes.keys() - set(before.sheetnames))},
            )
        evaluated = opened[2] if evaluated_path else None
        cached = None
        blank_results = set()
        if evaluated_path:
            cached = load_workbook(evaluated_path, data_only=True)
            opened.append(cached)
            blank_results = empty_string_results(evaluated_path)
        unrelated, structure, issues = [], [], []
        if before.sheetnames != after.sheetnames:
            structure.append("sheet_names_or_order")
        if before.epoch != after.epoch:
            structure.append("date_system")
        if str(before.defined_names) != str(after.defined_names):
            structure.append("defined_names")
        if protected_parts(original) != protected_parts(actual):
            structure.append("embedded_or_drawing_parts")
        for name in dict.fromkeys([*before.sheetnames, *after.sheetnames]):
            if name not in before.sheetnames or name not in after.sheetnames:
                continue
            old, new = before[name], after[name]
            if name not in structural and sheet_structure(old) != sheet_structure(new):
                structure.append(f"sheet_structure:{name}")
            for address in sorted(old._cells.keys() | new._cells.keys()):
                initial, current = old._cells.get(address), new._cells.get(address)
                representative = current if current is not None else initial
                if not in_scope(name, representative, scopes) and cell_signature(
                    initial
                ) != cell_signature(current):
                    unrelated.append({"sheet": name, "cell": representative.coordinate})
                if formula_text(current) is None or formula_text(
                    current
                ) == formula_text(initial):
                    continue
                cell_issues = formula_issues(name, current, evaluated)
                if cached is not None and not cell_issues:
                    result = cached[name][current.coordinate]
                    if result.data_type == "e":
                        cell_issues.append(
                            {
                                "sheet": name,
                                "cell": current.coordinate,
                                "reason": "formula_error",
                                "value": result.value,
                            }
                        )
                    elif (
                        result.value is None
                        and (name, current.coordinate) not in blank_results
                    ):
                        cell_issues.append(
                            {
                                "sheet": name,
                                "cell": current.coordinate,
                                "reason": "missing_cached_result",
                            }
                        )
                issues.extend(cell_issues)
        return {
            "verified": not (unrelated or structure or issues),
            "unrelatedChangeCount": len(unrelated),
            "unrelatedChanges": unrelated[:50],
            "structureChanges": structure,
            "formulaIssueCount": len(issues),
            "formulaIssues": issues[:50],
            "formulaValidation": "caller_evaluated_copy"
            if evaluated_path
            else "not_evaluated",
        }
    except (OSError, ValueError, KeyError, BadZipFile, ET.ParseError) as error:
        raise SheetOpsError(
            "invalid_workbook", "Workbook could not be verified", {"error": str(error)}
        ) from error
    finally:
        for book in opened:
            book.close()
