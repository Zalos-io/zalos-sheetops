import colorsys
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.styles.alignment import Alignment
from openpyxl.styles.borders import Border, Side
from openpyxl.styles.colors import COLOR_INDEX, Color
from openpyxl.styles.fills import PatternFill
from openpyxl.styles.fonts import Font
from openpyxl.styles.protection import Protection
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from zalos_sheetops.errors import SheetOpsError

MAX_CELLS = 10000
DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
THEME_INDEX_BY_NAME = {
    "lt1": 0,
    "dk1": 1,
    "lt2": 2,
    "dk2": 3,
    "accent1": 4,
    "accent2": 5,
    "accent3": 6,
    "accent4": 7,
    "accent5": 8,
    "accent6": 9,
    "hlink": 10,
    "folHlink": 11,
}


def read_styles(
    path: Path,
    sheet_name: str,
    cell_range: str,
    max_cells: int = MAX_CELLS,
) -> dict[str, object]:
    min_col, min_row, max_col, max_row = _parse_cell_range(cell_range)
    cell_count = (max_row - min_row + 1) * (max_col - min_col + 1)
    if cell_count > max_cells:
        raise SheetOpsError(
            "range_too_large",
            "Requested range is too large",
            {"range": cell_range, "cellCount": cell_count, "maxCells": max_cells},
        )

    workbook = _load_workbook(path)
    try:
        sheet = _resolve_sheet(workbook, sheet_name)
        theme_palette = _theme_palette(workbook.loaded_theme)
        cells = _read_style_range(
            sheet,
            min_row,
            max_row,
            min_col,
            max_col,
            theme_palette,
        )
    finally:
        workbook.close()

    return {
        "sheet": sheet_name,
        "range": cell_range,
        "rowCount": max_row - min_row + 1,
        "columnCount": max_col - min_col + 1,
        "cellCount": cell_count,
        "cells": cells,
    }


def _read_style_range(
    sheet: ReadOnlyWorksheet,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    theme_palette: dict[int, str],
) -> list[dict[str, object]]:
    cells = []
    rows = sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    )
    for row_offset, row in enumerate(rows):
        row_number = min_row + row_offset
        for column_offset, cell in enumerate(row):
            column_number = min_col + column_offset
            cells.append(
                _cell_payload(
                    cell,
                    row_number,
                    column_number,
                    theme_palette,
                )
            )
    return cells


def _cell_payload(
    cell: ReadOnlyCell | EmptyCell,
    row_number: int,
    column_number: int,
    theme_palette: dict[int, str],
) -> dict[str, object]:
    address = f"{get_column_letter(column_number)}{row_number}"
    return {
        "address": address,
        "style": _style_payload(cell, theme_palette),
    }


def _style_payload(
    cell: ReadOnlyCell | EmptyCell,
    theme_palette: dict[int, str],
) -> dict[str, object]:
    if isinstance(cell, EmptyCell):
        return _default_style_payload(theme_palette)
    return {
        "hasStyle": bool(cell.has_style),
        "numberFormat": cell.number_format,
        "font": _font_payload(cell.font, theme_palette),
        "fill": _fill_payload(cell.fill, theme_palette),
        "alignment": _alignment_payload(cell.alignment),
        "border": _border_payload(cell.border, theme_palette),
        "protection": _protection_payload(cell.protection),
    }


def _default_style_payload(theme_palette: dict[int, str]) -> dict[str, object]:
    return {
        "hasStyle": False,
        "numberFormat": "General",
        "font": _font_payload(Font(), theme_palette),
        "fill": _fill_payload(PatternFill(), theme_palette),
        "alignment": _alignment_payload(Alignment()),
        "border": _border_payload(Border(), theme_palette),
        "protection": _protection_payload(Protection()),
    }


def _font_payload(font: Font, theme_palette: dict[int, str]) -> dict[str, object]:
    return {
        "name": font.name,
        "size": float(font.sz) if font.sz is not None else None,
        "bold": bool(font.bold),
        "italic": bool(font.italic),
        "underline": font.underline,
        "strike": bool(font.strike),
        "color": _color_payload(font.color, theme_palette),
    }


def _fill_payload(
    fill: PatternFill,
    theme_palette: dict[int, str],
) -> dict[str, object]:
    if fill.fill_type is None:
        return {"type": None, "fgColor": None, "bgColor": None}
    return {
        "type": fill.fill_type,
        "fgColor": _color_payload(fill.fgColor, theme_palette),
        "bgColor": _color_payload(fill.bgColor, theme_palette),
    }


def _alignment_payload(alignment: Alignment) -> dict[str, object]:
    return {
        "horizontal": alignment.horizontal,
        "vertical": alignment.vertical,
        "wrapText": bool(alignment.wrap_text),
        "shrinkToFit": bool(alignment.shrink_to_fit),
        "textRotation": alignment.textRotation,
    }


def _border_payload(
    border: Border,
    theme_palette: dict[int, str],
) -> dict[str, object]:
    return {
        "left": _side_payload(border.left, theme_palette),
        "right": _side_payload(border.right, theme_palette),
        "top": _side_payload(border.top, theme_palette),
        "bottom": _side_payload(border.bottom, theme_palette),
    }


def _side_payload(
    side: Side | None,
    theme_palette: dict[int, str],
) -> dict[str, object]:
    if side is None:
        return {"style": None, "color": None}
    return {
        "style": side.style,
        "color": _color_payload(side.color, theme_palette),
    }


def _protection_payload(protection: Protection) -> dict[str, object]:
    return {"locked": bool(protection.locked), "hidden": bool(protection.hidden)}


def _color_payload(
    color: Color | None,
    theme_palette: dict[int, str],
) -> dict[str, object] | None:
    if color is None or color.type is None:
        return None
    if color.type == "rgb":
        payload = _rgb_color_payload(color.rgb)
        if payload is not None:
            payload["resolvedRgb"] = payload.get("rgb")
        return payload
    if color.type == "theme":
        payload = {
            "type": "theme",
            "theme": color.theme,
            "resolvedRgb": _resolve_theme_color(color, theme_palette),
        }
        if color.tint != 0:
            payload["tint"] = color.tint
        return payload
    if color.type == "indexed":
        return {
            "type": "indexed",
            "indexed": color.indexed,
            "resolvedRgb": _resolve_indexed_color(color),
        }
    if color.type == "auto":
        return {
            "type": "auto",
            "auto": bool(color.auto),
            "resolvedRgb": None,
        }
    return {
        "type": str(color.type),
        "value": color.value,
        "resolvedRgb": None,
    }


def _rgb_color_payload(rgb: str | None) -> dict[str, object] | None:
    if rgb is None:
        return None
    color = rgb.upper()
    if color == "00000000":
        return None
    if len(color) == 8:
        color = color[2:]
    if len(color) != 6:
        return {"type": "rgb", "rgb": rgb}
    return {"type": "rgb", "rgb": f"#{color}"}


def _theme_palette(theme_xml: bytes | None) -> dict[int, str]:
    if theme_xml is None:
        return {}
    try:
        root = ET.fromstring(theme_xml)
    except ET.ParseError:
        return {}
    scheme = root.find(f".//{{{DRAWINGML_NS}}}clrScheme")
    if scheme is None:
        return {}
    palette: dict[int, str] = {}
    for slot in scheme:
        name = _local_name(slot.tag)
        index = THEME_INDEX_BY_NAME.get(name)
        if index is None:
            continue
        color = _theme_slot_color(slot)
        if color is not None:
            palette[index] = color
    return palette


def _theme_slot_color(slot: ET.Element) -> str | None:
    rgb = slot.find(f"{{{DRAWINGML_NS}}}srgbClr")
    if rgb is not None:
        return _normalize_hex6(rgb.attrib.get("val"))
    system = slot.find(f"{{{DRAWINGML_NS}}}sysClr")
    if system is not None:
        return _normalize_hex6(system.attrib.get("lastClr"))
    return None


def _resolve_theme_color(
    color: Color,
    theme_palette: dict[int, str],
) -> str | None:
    if not isinstance(color.theme, int):
        return None
    base = theme_palette.get(color.theme)
    if base is None:
        return None
    return f"#{_apply_tint(base, float(color.tint or 0.0))}"


def _resolve_indexed_color(color: Color) -> str | None:
    if not isinstance(color.indexed, int):
        return None
    if color.indexed < 0 or color.indexed >= len(COLOR_INDEX):
        return None
    return _argb_to_rgb(COLOR_INDEX[color.indexed])


def _apply_tint(hex6: str, tint: float) -> str:
    if tint == 0:
        return hex6
    red = int(hex6[0:2], 16) / 255.0
    green = int(hex6[2:4], 16) / 255.0
    blue = int(hex6[4:6], 16) / 255.0
    hue, luminance, saturation = colorsys.rgb_to_hls(red, green, blue)
    if tint < 0:
        luminance *= 1.0 + tint
    else:
        luminance += (1.0 - luminance) * tint
    luminance = max(0.0, min(1.0, luminance))
    red, green, blue = colorsys.hls_to_rgb(hue, luminance, saturation)
    return "".join(f"{_color_byte(channel):02X}" for channel in (red, green, blue))


def _color_byte(channel: float) -> int:
    return max(0, min(255, round(channel * 255)))


def _argb_to_rgb(value: str) -> str | None:
    normalized = value.upper()
    if len(normalized) == 8:
        normalized = normalized[2:]
    if _normalize_hex6(normalized) is None:
        return None
    return f"#{normalized}"


def _normalize_hex6(value: str | None) -> str | None:
    if value is None:
        return None
    normalized = value.upper()
    if len(normalized) != 6:
        return None
    if any(character not in "0123456789ABCDEF" for character in normalized):
        return None
    return normalized


def _local_name(tag: str) -> str:
    if "}" not in tag:
        return tag
    return tag.rsplit("}", 1)[1]


def _parse_cell_range(cell_range: str) -> tuple[int, int, int, int]:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except ValueError as error:
        raise SheetOpsError(
            "invalid_range",
            "Cell range must be a bounded A1-style range",
            {"range": cell_range},
        ) from error
    if None in {min_col, min_row, max_col, max_row}:
        raise SheetOpsError(
            "invalid_range",
            "Cell range must be a bounded A1-style range",
            {"range": cell_range},
        )
    return min_col, min_row, max_col, max_row


def _resolve_sheet(workbook: Workbook, sheet_name: str) -> ReadOnlyWorksheet:
    if sheet_name not in workbook.sheetnames:
        raise SheetOpsError(
            "invalid_sheet",
            "Requested sheet does not exist",
            {"sheet": sheet_name},
        )
    return workbook[sheet_name]


def _load_workbook(path: Path) -> Workbook:
    try:
        return load_workbook(path, data_only=False, read_only=True)
    except Exception as error:
        raise SheetOpsError(
            "invalid_workbook",
            "Workbook could not be opened",
            {"path": str(path)},
        ) from error
