from pathlib import Path
from typing import TYPE_CHECKING

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.cellref import coordinate_to_tuple

if TYPE_CHECKING:
    from openpyxl.cell.cell import Cell
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet
from zalos_sheetops.styles.xml_assertions import (
    assert_styles_xml,
    can_assert_styles_xml,
)


def assert_styles(
    path: Path,
    sheet_name: str,
    checks: list[dict[str, object]],
) -> dict[str, object]:
    _validate_checks(checks)
    if can_assert_styles_xml(checks):
        return assert_styles_xml(path, sheet_name, checks)

    workbook = _load_workbook(path)
    try:
        sheet = _resolve_sheet(workbook, sheet_name)
        failures = []
        for check in checks:
            cell = _resolve_cell(sheet, check)
            failures.extend(_check_cell(cell, check))
    finally:
        workbook.close()

    return {
        "valid": not failures,
        "sheet": sheet_name,
        "checkedCount": len(checks),
        "failureCount": len(failures),
        "failures": failures,
    }


def _validate_checks(checks: list[dict[str, object]]) -> None:
    if not checks:
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "checks", "expected": "non-empty list[styleCheck]"},
        )
    for index, check in enumerate(checks):
        if not isinstance(check, dict):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {"field": f"checks[{index}]", "expected": "styleCheck"},
            )
        cell = check.get("cell")
        if not isinstance(cell, str):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {"field": f"checks[{index}].cell", "expected": "A1-style cell"},
            )
        try:
            coordinate_to_tuple(cell)
        except ValueError as error:
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {"field": f"checks[{index}].cell", "expected": "A1-style cell"},
            ) from error
        fill = check.get("fill")
        if fill is not None and not isinstance(fill, str):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {"field": f"checks[{index}].fill", "expected": "hex color string"},
            )
        if fill is not None:
            _normalize_color(fill, f"checks[{index}].fill")
        comment_contains = check.get("commentContains")
        if comment_contains is not None and not isinstance(comment_contains, str):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {
                    "field": f"checks[{index}].commentContains",
                    "expected": "string",
                },
            )
        formula = check.get("formula")
        if formula is not None and not isinstance(formula, bool):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {"field": f"checks[{index}].formula", "expected": "boolean"},
            )
        bold = check.get("bold")
        if bold is not None and not isinstance(bold, bool):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {"field": f"checks[{index}].bold", "expected": "boolean"},
            )
        italic = check.get("italic")
        if italic is not None and not isinstance(italic, bool):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {"field": f"checks[{index}].italic", "expected": "boolean"},
            )
        if (
            fill is None
            and comment_contains is None
            and formula is None
            and bold is None
            and italic is None
        ):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {
                    "field": f"checks[{index}]",
                    "expected": (
                        "styleCheck with fill, commentContains, formula, bold, "
                        "or italic"
                    ),
                },
            )


def _check_cell(cell: "Cell", check: dict[str, object]) -> list[dict[str, object]]:
    failures = []
    fill = check.get("fill")
    if isinstance(fill, str):
        expected_fill = _normalize_color(fill, "fill")
        actual_fill = _cell_fill(cell)
        if actual_fill != expected_fill:
            failures.append(
                {
                    "cell": cell.coordinate,
                    "code": "fill_mismatch",
                    "expected": expected_fill,
                    "actual": actual_fill,
                }
            )

    comment_contains = check.get("commentContains")
    if isinstance(comment_contains, str):
        actual_comment = cell.comment.text if cell.comment is not None else None
        if actual_comment is None or comment_contains not in actual_comment:
            failures.append(
                {
                    "cell": cell.coordinate,
                    "code": "comment_mismatch",
                    "expected": comment_contains,
                    "actual": actual_comment,
                }
            )

    expected_formula = check.get("formula")
    if isinstance(expected_formula, bool):
        actual_formula = isinstance(cell.value, str) and cell.value.startswith("=")
        if actual_formula != expected_formula:
            failures.append(
                {
                    "cell": cell.coordinate,
                    "code": "formula_mismatch",
                    "expected": expected_formula,
                    "actual": actual_formula,
                }
            )

    expected_bold = check.get("bold")
    if isinstance(expected_bold, bool):
        actual_bold = bool(cell.font.bold)
        if actual_bold != expected_bold:
            failures.append(
                {
                    "cell": cell.coordinate,
                    "code": "bold_mismatch",
                    "expected": expected_bold,
                    "actual": actual_bold,
                }
            )

    expected_italic = check.get("italic")
    if isinstance(expected_italic, bool):
        actual_italic = bool(cell.font.italic)
        if actual_italic != expected_italic:
            failures.append(
                {
                    "cell": cell.coordinate,
                    "code": "italic_mismatch",
                    "expected": expected_italic,
                    "actual": actual_italic,
                }
            )
    return failures


def _cell_fill(cell: "Cell") -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type is None:
        return None
    color = fill.fgColor
    if color is None or color.type != "rgb" or color.rgb is None:
        return None
    return _normalize_color(str(color.rgb), "fill")


def _normalize_color(value: str, field_name: str) -> str:
    color = value.strip().upper()
    if color.startswith("#"):
        color = color[1:]
    if len(color) == 8:
        color = color[2:]
    if len(color) != 6 or any(char not in "0123456789ABCDEF" for char in color):
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": field_name, "expected": "hex color string"},
        )
    return f"#{color}"


def _resolve_cell(sheet: "Worksheet", check: dict[str, object]) -> "Cell":
    return sheet[str(check["cell"])]


def _resolve_sheet(workbook: "Workbook", sheet_name: str) -> "Worksheet":
    if sheet_name not in workbook.sheetnames:
        raise SheetOpsError(
            "invalid_sheet",
            "Requested sheet does not exist",
            {"sheet": sheet_name},
        )
    return workbook[sheet_name]


def _load_workbook(path: Path) -> "Workbook":
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=False, read_only=False)
    except Exception as error:
        raise SheetOpsError(
            "invalid_workbook",
            "Workbook could not be opened",
            {"path": str(path)},
        ) from error
