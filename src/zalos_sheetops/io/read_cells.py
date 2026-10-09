from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.cellref import get_column_letter, range_boundaries
from zalos_sheetops.io.xlsx_xml import read_cells_xml

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

INCLUDE_FIELDS = {
    "values",
    "formulas",
    "cachedValues",
    "numberFormats",
    "dataTypes",
}
INCLUDE_EXPECTED = "list[values | formulas | cachedValues | numberFormats | dataTypes]"
MAX_CELLS = 10000


def read_cells(
    path: Path,
    sheet_name: str,
    cell_range: str,
    include: list[str] | None = None,
    max_cells: int = MAX_CELLS,
    engine: str = "openpyxl",
) -> dict[str, object]:
    include_fields = _normalize_include(include)
    min_col, min_row, max_col, max_row = _parse_cell_range(cell_range)
    cell_count = (max_row - min_row + 1) * (max_col - min_col + 1)
    if cell_count > max_cells:
        raise SheetOpsError(
            "range_too_large",
            "Requested range is too large",
            {"range": cell_range, "cellCount": cell_count, "maxCells": max_cells},
        )
    if engine not in {"auto", "xml", "openpyxl"}:
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "engine", "expected": "auto | xml | openpyxl"},
        )
    if engine in {"auto", "xml"}:
        return read_cells_xml(path, sheet_name, cell_range, include_fields)

    formula_workbook = _load_workbook(path, data_only=False)
    value_workbook = None
    try:
        formula_sheet = _resolve_sheet(formula_workbook, sheet_name)
        if "values" in include_fields or "cachedValues" in include_fields:
            value_workbook = _load_workbook(path, data_only=True)
            value_sheet = _resolve_sheet(value_workbook, sheet_name)
        else:
            value_sheet = None

        cells = _read_cell_range(
            formula_sheet,
            value_sheet,
            min_row,
            max_row,
            min_col,
            max_col,
            include_fields,
        )
    finally:
        _close_workbook(formula_workbook)
        if value_workbook is not None:
            _close_workbook(value_workbook)

    return {
        "sheet": sheet_name,
        "range": cell_range,
        "rowCount": max_row - min_row + 1,
        "columnCount": max_col - min_col + 1,
        "cellCount": cell_count,
        "returnedCellCount": len(cells),
        "cells": cells,
    }


def _read_cell_range(
    formula_sheet: "Worksheet",
    value_sheet: "Worksheet | None",
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    include_fields: list[str],
) -> list[dict[str, object]]:
    value_rows = None
    if value_sheet is not None:
        value_rows = value_sheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        )

    cells = []
    formula_rows = formula_sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    )
    for row_offset, formula_row in enumerate(formula_rows):
        value_row = next(value_rows) if value_rows is not None else None
        row_number = min_row + row_offset
        for column_offset, formula_cell in enumerate(formula_row):
            column_number = min_col + column_offset
            value_cell = value_row[column_offset] if value_row is not None else None
            # A cell with no stored content has value None in both the
            # formula view and the cached-value view; omit it entirely.
            if (
                getattr(formula_cell, "value", None) is None
                and getattr(value_cell, "value", None) is None
            ):
                continue
            cells.append(
                _cell_payload(
                    formula_cell,
                    value_cell,
                    row_number,
                    column_number,
                    include_fields,
                )
            )
    return cells


def _cell_payload(
    formula_cell: object,
    value_cell: object | None,
    row_number: int,
    column_number: int,
    include_fields: list[str],
) -> dict[str, object]:
    value = getattr(value_cell, "value", None)
    formula_value = getattr(formula_cell, "value", None)
    payload: dict[str, object] = {
        "address": f"{get_column_letter(column_number)}{row_number}",
    }
    if "values" in include_fields:
        payload["value"] = _json_value(value)
    if (
        "formulas" in include_fields
        and isinstance(formula_value, str)
        and formula_value.startswith("=")
    ):
        payload["formula"] = formula_value
    if "cachedValues" in include_fields:
        payload["cachedValue"] = _json_value(value)
    if "numberFormats" in include_fields:
        payload["numberFormat"] = (
            getattr(formula_cell, "number_format", None) or "General"
        )
    if "dataTypes" in include_fields:
        payload["dataType"] = getattr(formula_cell, "data_type", None) or "n"
    return payload


def _normalize_include(include: list[str] | None) -> list[str]:
    include_fields = include or ["values", "formulas"]
    invalid = [field for field in include_fields if field not in INCLUDE_FIELDS]
    if not include_fields or invalid:
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "include", "expected": INCLUDE_EXPECTED, "invalid": invalid},
        )
    return include_fields


def _parse_cell_range(cell_range: str) -> tuple[int, int, int, int]:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except ValueError as error:
        raise SheetOpsError(
            "invalid_range",
            "Cell range must be a bounded A1-style range",
            {"range": cell_range},
        ) from error
    if None in {min_col, min_row, max_col, max_row}:
        raise SheetOpsError(
            "invalid_range",
            "Cell range must be a bounded A1-style range",
            {"range": cell_range},
        )
    return min_col, min_row, max_col, max_row


def _resolve_sheet(workbook: "Workbook", sheet_name: str) -> "Worksheet":
    if sheet_name not in workbook.sheetnames:
        raise SheetOpsError(
            "invalid_sheet",
            "Requested sheet does not exist",
            {"sheet": sheet_name},
        )
    return workbook[sheet_name]


def _load_workbook(path: Path, data_only: bool) -> "Workbook":
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=data_only, read_only=True)
    except Exception as error:
        raise SheetOpsError(
            "invalid_workbook",
            "Workbook could not be opened",
            {"path": str(path)},
        ) from error


def _close_workbook(workbook: "Workbook") -> None:
    workbook.close()


def _json_value(value: object) -> object:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value
