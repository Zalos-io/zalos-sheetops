import re
from collections.abc import Callable
from datetime import date, datetime, time
from decimal import Decimal
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.styles.read import _color_payload, _theme_palette

SEARCH_FIELDS = [
    "values",
    "formulas",
    "cachedValues",
    "numberFormats",
    "fontColors",
    "fillColors",
]
SEARCH_FIELDS_EXPECTED = (
    "non-empty list[values | formulas | cachedValues | "
    "numberFormats | fontColors | fillColors]"
)
DEFAULT_FIELDS = ["values", "formulas"]
DEFAULT_LIMIT = 100
DEFAULT_MAX_CELLS = 1_000_000


def search_cells(
    path: Path,
    query: str,
    *,
    fields: list[str] | None = None,
    sheet_name: str | None = None,
    cell_range: str | None = None,
    regex: bool = False,
    case_sensitive: bool = False,
    include_hidden: bool = False,
    limit: int = DEFAULT_LIMIT,
    offset: int = 0,
    max_cells: int = DEFAULT_MAX_CELLS,
) -> dict[str, object]:
    search_fields = _normalize_fields(fields)
    _validate_search_request(
        query,
        sheet_name,
        cell_range,
        regex,
        case_sensitive,
        include_hidden,
        limit,
        offset,
        max_cells,
    )
    matcher = _build_matcher(query, regex, case_sensitive)
    bounds = _parse_cell_range(cell_range) if cell_range is not None else None
    needs_formula_workbook = bool(set(search_fields) - {"values", "cachedValues"})

    formula_workbook = _load_workbook(path, data_only=not needs_formula_workbook)
    value_workbook = None
    try:
        if needs_formula_workbook and {"values", "cachedValues"} & set(search_fields):
            value_workbook = _load_workbook(path, data_only=True)
        sheet_names = _select_sheet_names(
            formula_workbook,
            sheet_name,
            include_hidden,
        )
        result = _search_workbooks(
            formula_workbook,
            value_workbook,
            sheet_names,
            bounds,
            search_fields,
            matcher,
            limit,
            offset,
            max_cells,
            primary_data_only=not needs_formula_workbook,
        )
    finally:
        formula_workbook.close()
        if value_workbook is not None:
            value_workbook.close()

    stopped_early = bool(result["stoppedEarly"])
    has_more: bool | None = bool(result["hasMore"])
    if stopped_early and not has_more:
        has_more = None
    matches = result["matches"]
    returned_count = len(matches) if isinstance(matches, list) else 0
    return {
        "scope": {"sheet": sheet_name, "range": cell_range},
        "query": query,
        "fields": search_fields,
        "regex": regex,
        "caseSensitive": case_sensitive,
        "includeHidden": include_hidden,
        "limit": limit,
        "offset": offset,
        "returnedCount": returned_count,
        "hasMore": has_more,
        "nextOffset": offset + returned_count if has_more is True else None,
        "scannedCellCount": result["scannedCellCount"],
        "stoppedEarly": stopped_early,
        "matches": matches,
    }


def _search_workbooks(
    formula_workbook: Workbook,
    value_workbook: Workbook | None,
    sheet_names: list[str],
    bounds: tuple[int, int, int, int] | None,
    fields: list[str],
    matcher: Callable[[str], bool],
    limit: int,
    offset: int,
    max_cells: int,
    *,
    primary_data_only: bool,
) -> dict[str, object]:
    matches: list[dict[str, object]] = []
    matched_count = 0
    scanned_cell_count = 0
    stopped_early = False
    has_more = False
    theme_palette = _theme_palette(formula_workbook.loaded_theme)

    for sheet_name in sheet_names:
        formula_sheet = formula_workbook[sheet_name]
        value_sheet = value_workbook[sheet_name] if value_workbook is not None else None
        for (
            row_number,
            column_number,
            formula_cell,
            value_cell,
        ) in _iter_sheet_cells(formula_sheet, value_sheet, bounds):
            if scanned_cell_count >= max_cells:
                stopped_early = True
                break
            scanned_cell_count += 1
            if isinstance(formula_cell, EmptyCell) and isinstance(
                value_cell, EmptyCell
            ):
                continue
            field_values, payload = _searchable_cell_payload(
                formula_cell,
                value_cell,
                fields,
                theme_palette,
                primary_data_only,
            )
            matched_fields = [
                field
                for field in fields
                if any(matcher(text) for text in field_values.get(field, []))
            ]
            if not matched_fields:
                continue
            if matched_count < offset:
                matched_count += 1
                continue
            if len(matches) >= limit:
                has_more = True
                break
            matched_count += 1
            matches.append(
                {
                    "sheet": sheet_name,
                    "cell": f"{get_column_letter(column_number)}{row_number}",
                    "matchedFields": matched_fields,
                    **payload,
                }
            )
        if stopped_early or has_more:
            break

    return {
        "matches": matches,
        "scannedCellCount": scanned_cell_count,
        "stoppedEarly": stopped_early,
        "hasMore": has_more,
    }


def _iter_sheet_cells(
    formula_sheet: ReadOnlyWorksheet,
    value_sheet: ReadOnlyWorksheet | None,
    bounds: tuple[int, int, int, int] | None,
) -> object:
    min_col, min_row, max_col, max_row = _iteration_bounds(bounds)
    formula_rows = formula_sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    )
    value_rows = (
        value_sheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        )
        if value_sheet is not None
        else ()
    )
    start_row = min_row or 1
    start_column = min_col or 1
    for row_offset, (formula_row, value_row) in enumerate(
        zip_longest(formula_rows, value_rows, fillvalue=()),
    ):
        for column_offset, (formula_cell, value_cell) in enumerate(
            zip_longest(formula_row, value_row, fillvalue=EmptyCell()),
        ):
            coordinate_cell = (
                formula_cell
                if isinstance(formula_cell, ReadOnlyCell)
                else value_cell
                if isinstance(value_cell, ReadOnlyCell)
                else None
            )
            row_number = (
                coordinate_cell.row
                if coordinate_cell is not None
                else start_row + row_offset
            )
            column_number = (
                coordinate_cell.column
                if coordinate_cell is not None
                else start_column + column_offset
            )
            yield (
                row_number,
                column_number,
                formula_cell,
                value_cell,
            )


def _searchable_cell_payload(
    formula_cell: ReadOnlyCell | EmptyCell,
    value_cell: ReadOnlyCell | EmptyCell,
    fields: list[str],
    theme_palette: dict[int, str],
    primary_data_only: bool,
) -> tuple[dict[str, list[str]], dict[str, object]]:
    field_values: dict[str, list[str]] = {}
    payload: dict[str, object] = {}
    primary_value = (
        formula_cell.value if isinstance(formula_cell, ReadOnlyCell) else None
    )
    formula_value = None if primary_data_only else primary_value
    if primary_data_only:
        cached_value = primary_value
    else:
        cached_value = (
            value_cell.value if isinstance(value_cell, ReadOnlyCell) else None
        )

    if "values" in fields:
        json_value = _json_value(cached_value)
        payload["value"] = json_value
        field_values["values"] = _value_search_text(json_value)
    if "formulas" in fields:
        formula = (
            formula_value
            if isinstance(formula_value, str) and formula_value.startswith("=")
            else None
        )
        payload["formula"] = formula
        field_values["formulas"] = [formula] if formula is not None else []
    if "cachedValues" in fields:
        json_value = _json_value(cached_value)
        payload["cachedValue"] = json_value
        field_values["cachedValues"] = _value_search_text(json_value)
    if "numberFormats" in fields:
        number_format = (
            formula_cell.number_format
            if isinstance(formula_cell, ReadOnlyCell)
            else "General"
        )
        payload["numberFormat"] = number_format
        field_values["numberFormats"] = [number_format]
    if "fontColors" in fields:
        font_color = (
            _color_payload(formula_cell.font.color, theme_palette)
            if isinstance(formula_cell, ReadOnlyCell)
            else None
        )
        payload["fontColor"] = font_color
        field_values["fontColors"] = _color_search_text(font_color)
    if "fillColors" in fields:
        fill = _fill_payload(formula_cell, theme_palette)
        payload["fill"] = fill
        field_values["fillColors"] = _fill_search_text(fill)
    return field_values, payload


def _fill_payload(
    cell: ReadOnlyCell | EmptyCell,
    theme_palette: dict[int, str],
) -> dict[str, object] | None:
    if isinstance(cell, EmptyCell) or cell.fill.fill_type is None:
        return None
    return {
        "type": cell.fill.fill_type,
        "fgColor": _color_payload(cell.fill.fgColor, theme_palette),
        "bgColor": _color_payload(cell.fill.bgColor, theme_palette),
    }


def _fill_search_text(fill: dict[str, object] | None) -> list[str]:
    if fill is None:
        return []
    values = []
    for field in ("fgColor", "bgColor"):
        color = fill.get(field)
        if isinstance(color, dict):
            values.extend(_color_search_text(color))
    return values


def _color_search_text(color: dict[str, object] | None) -> list[str]:
    if color is None:
        return []
    values = []
    rgb = color.get("rgb")
    if isinstance(rgb, str):
        values.append(rgb)
    resolved = color.get("resolvedRgb")
    if isinstance(resolved, str) and resolved not in values:
        values.append(resolved)
    theme = color.get("theme")
    if isinstance(theme, int):
        values.append(f"theme:{theme}")
    indexed = color.get("indexed")
    if isinstance(indexed, int):
        values.append(f"indexed:{indexed}")
    return values


def _value_search_text(value: object) -> list[str]:
    if value is None:
        return []
    if isinstance(value, bool):
        return ["true" if value else "false"]
    return [str(value)]


def _json_value(value: object) -> object:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _select_sheet_names(
    workbook: Workbook,
    sheet_name: str | None,
    include_hidden: bool,
) -> list[str]:
    if sheet_name is not None:
        if sheet_name not in workbook.sheetnames:
            raise SheetOpsError(
                "invalid_sheet",
                "Requested sheet does not exist",
                {"sheet": sheet_name},
            )
        return [sheet_name]
    return [
        name
        for name in workbook.sheetnames
        if include_hidden or workbook[name].sheet_state == "visible"
    ]


def _build_matcher(
    query: str,
    regex: bool,
    case_sensitive: bool,
) -> Callable[[str], bool]:
    if regex:
        flags = 0 if case_sensitive else re.IGNORECASE
        try:
            pattern = re.compile(query, flags)
        except re.error as error:
            raise SheetOpsError(
                "invalid_regex",
                "Search pattern is not a valid regular expression",
                {"query": query},
            ) from error
        return lambda value: pattern.search(value) is not None
    if case_sensitive:
        return lambda value: query in value
    normalized_query = query.casefold()
    return lambda value: normalized_query in value.casefold()


def _normalize_fields(fields: list[str] | None) -> list[str]:
    if fields is None:
        return list(DEFAULT_FIELDS)
    if (
        not isinstance(fields, list)
        or not fields
        or any(
            not isinstance(field, str) or field not in SEARCH_FIELDS for field in fields
        )
    ):
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "fields", "expected": SEARCH_FIELDS_EXPECTED},
        )
    requested = set(fields)
    return [field for field in SEARCH_FIELDS if field in requested]


def _validate_search_request(
    query: str,
    sheet_name: str | None,
    cell_range: str | None,
    regex: bool,
    case_sensitive: bool,
    include_hidden: bool,
    limit: int,
    offset: int,
    max_cells: int,
) -> None:
    if not isinstance(query, str) or query == "":
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "query", "expected": "non-empty string"},
        )
    if sheet_name is not None and not isinstance(sheet_name, str):
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "sheet", "expected": "string | null"},
        )
    if cell_range is not None and not isinstance(cell_range, str):
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "range", "expected": "A1-style range | null"},
        )
    if cell_range is not None and sheet_name is None:
        raise SheetOpsError(
            "invalid_scope",
            "Cell range requires an explicit sheet",
            {"field": "range", "requires": "sheet"},
        )
    for field_name, value in (
        ("regex", regex),
        ("caseSensitive", case_sensitive),
        ("includeHidden", include_hidden),
    ):
        if not isinstance(value, bool):
            raise SheetOpsError(
                "invalid_field",
                "Field has invalid type",
                {"field": field_name, "expected": "boolean"},
            )
    _require_positive_integer(limit, "limit")
    _require_non_negative_integer(offset, "offset")
    _require_positive_integer(max_cells, "maxCells")


def _require_positive_integer(value: int, field_name: str) -> None:
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": field_name, "expected": "positive integer"},
        )


def _require_non_negative_integer(value: int, field_name: str) -> None:
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": field_name, "expected": "non-negative integer"},
        )


def _parse_cell_range(cell_range: str) -> tuple[int, int, int, int]:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except ValueError as error:
        raise SheetOpsError(
            "invalid_range",
            "Cell range must be a bounded A1-style range",
            {"range": cell_range},
        ) from error
    if None in {min_col, min_row, max_col, max_row}:
        raise SheetOpsError(
            "invalid_range",
            "Cell range must be a bounded A1-style range",
            {"range": cell_range},
        )
    return min_col, min_row, max_col, max_row


def _iteration_bounds(
    bounds: tuple[int, int, int, int] | None,
) -> tuple[int | None, int | None, int | None, int | None]:
    if bounds is None:
        return None, None, None, None
    min_col, min_row, max_col, max_row = bounds
    return min_col, min_row, max_col, max_row


def _load_workbook(path: Path, data_only: bool) -> Workbook:
    try:
        return load_workbook(path, data_only=data_only, read_only=True)
    except Exception as error:
        raise SheetOpsError(
            "invalid_workbook",
            "Workbook could not be opened",
            {"path": str(path)},
        ) from error
