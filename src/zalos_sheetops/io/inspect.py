from pathlib import Path
from typing import TYPE_CHECKING

from zalos_sheetops.errors import SheetOpsError
from zalos_sheetops.io.xlsx_xml import read_workbook_xml_summary

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet


def inspect_workbook(
    path: Path,
    engine: str = "auto",
    max_formula_cells: int = 20,
    max_merged_ranges: int = 1000,
    max_graphics: int = 1000,
    max_defined_names: int = 1000,
) -> dict[str, object]:
    if engine not in {"auto", "xml", "openpyxl"}:
        raise SheetOpsError(
            "invalid_field",
            "Field has invalid type",
            {"field": "engine", "expected": "auto | xml | openpyxl"},
        )
    if engine in {"auto", "xml"}:
        return read_workbook_xml_summary(
            path,
            max_formula_cells=max_formula_cells,
            max_merged_ranges=max_merged_ranges,
            max_graphics=max_graphics,
            max_defined_names=max_defined_names,
        )

    workbook = _load_inspection_workbook(path)
    sheets = [_inspect_sheet(sheet) for sheet in workbook.worksheets]
    return {
        "workbook": {
            "path": str(path),
            "sheetCount": len(workbook.worksheets),
            "formulaCount": _count_workbook_formulas(workbook),
        },
        "sheets": sheets,
    }


def _load_inspection_workbook(path: Path) -> "Workbook":
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=False, read_only=False)
    except Exception as error:
        raise SheetOpsError(
            "invalid_workbook",
            "Workbook could not be opened",
            {"path": str(path)},
        ) from error


def _inspect_sheet(sheet: "Worksheet") -> dict[str, object]:
    formula_cells = _formula_cells(sheet)
    return {
        "name": sheet.title,
        "maxRow": sheet.max_row,
        "maxColumn": sheet.max_column,
        "hidden": sheet.sheet_state != "visible",
        "formulaCount": len(formula_cells),
        "formulaCells": formula_cells,
        "dimensions": sheet.calculate_dimension(),
        "mergedRangeCount": len(sheet.merged_cells.ranges),
        "hyperlinkCount": _count_hyperlinks(sheet),
    }


def _count_workbook_formulas(workbook: "Workbook") -> int:
    return sum(_count_sheet_formulas(sheet) for sheet in workbook.worksheets)


def _count_sheet_formulas(sheet: "Worksheet") -> int:
    return len(_formula_cells(sheet))


def _formula_cells(sheet: "Worksheet") -> list[str]:
    cells = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cells.append(cell.coordinate)
    return cells


def _count_hyperlinks(sheet: "Worksheet") -> int:
    count = 0
    for row in sheet.iter_rows():
        for cell in row:
            if cell.hyperlink is not None:
                count += 1
    return count
